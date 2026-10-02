"""Reader for commissioning and GTM events in ``База по скважинам.xlsx``."""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gdis_loader.domain.dates import parse_partial_date
from gdis_loader.domain.models import (
    DiagnosticIssue,
    IssueSeverity,
    PartialDate,
    SourceLocation,
    WellEvent,
    WellEventKind,
    WellHistory,
    WellHistoryParseResult,
    WellStage,
)
from gdis_loader.matching.normalize import normalize_field_name, normalize_plast, normalize_well_id
from gdis_loader.readers.common import (
    as_text,
    coordinate,
    display_value,
    normalized_header,
    openpyxl_read_warnings,
    sheet_location,
)

FIELD_SHEETS: dict[str, str] = {
    "БЛУ": "Береговое",
    "ХЛУ": "Хадырьяхинское",
    "ПЛУ": "Пырейное",
}
_GTM_HEADER_RE = re.compile(r"^ГТМ\s*(\d+)$", re.IGNORECASE)
_EVENT_KIND_MAP = {
    "ЗБС": WellEventKind.ZBS,
    "ПВЛГ": WellEventKind.PVLG,
    "ПНЛГ": WellEventKind.PNLG,
    "ПВР": WellEventKind.PVR,
    "КЛК": WellEventKind.KLK,
}
_BLOCK_HEADER_NAMES = ("Дата", "Пласт", "Отчет", "Заканчивание", "Примечание")


@dataclass(frozen=True, slots=True)
class WellHistorySchema:
    """Detected two-row header positions for the base well and GTM blocks."""

    well_column: int
    launch_date_column: int
    initial_plast_column: int
    initial_block_columns: tuple[int, ...]
    gtm_blocks: tuple[tuple[int, int, int], ...]


def find_well_history_schema(worksheet: Any) -> WellHistorySchema | None:
    """Find and validate the documented two-row header and all GTM blocks."""

    if worksheet.max_row < 2:  # type: ignore[attr-defined]
        return None
    first = next(
        worksheet.iter_rows(min_row=1, max_row=1, values_only=True),  # type: ignore[attr-defined]
        (),
    )
    second = next(
        worksheet.iter_rows(min_row=2, max_row=2, values_only=True),  # type: ignore[attr-defined]
        (),
    )
    width = max(len(first), len(second))
    combined = [
        normalized_header(" ".join(str(value) for value in (first[i], second[i]) if value))
        for i in range(width)
    ]
    required = tuple(
        normalized_header(value)
        for value in (
            "№ куста",
            "№ скв.",
            "дата внс",
            "пласт",
            "информация",
            "заканчивание",
            "примечание",
        )
    )
    if any(token not in combined for token in required):
        return None
    locations = {
        "well": next(
            index for index, header in enumerate(combined) if header == normalized_header("№ скв.")
        ),
        "launch": next(index for index, header in enumerate(combined) if header == "дата внс"),
        "plast": next(index for index, header in enumerate(combined) if header == "пласт"),
    }
    if any(index is None for index in locations.values()):
        return None

    gtm_blocks: list[tuple[int, int, int]] = []
    for index, value in enumerate(second):
        if not isinstance(value, str):
            continue
        gtm_match = _GTM_HEADER_RE.fullmatch(normalized_header(value))
        if gtm_match is None:
            continue
        if index + 5 >= width:
            return None
        block_headers = tuple(normalized_header(second[index + offset]) for offset in range(1, 6))
        if block_headers != tuple(normalized_header(name) for name in _BLOCK_HEADER_NAMES):
            return None
        block_number = int(gtm_match.group(1))
        gtm_blocks.append((index, index + 1, block_number))

    if not gtm_blocks:
        return None
    return WellHistorySchema(
        well_column=int(locations["well"]),
        launch_date_column=int(locations["launch"]),
        initial_plast_column=int(locations["plast"]),
        initial_block_columns=tuple(range(min(7, width))),
        gtm_blocks=tuple(gtm_blocks),
    )


def _event_kind(label: str | None) -> WellEventKind:
    return _EVENT_KIND_MAP.get((label or "").strip().upper(), WellEventKind.OTHER)


def _cell_location(source_file: str, sheet: str, row: int, column: int) -> SourceLocation:
    return SourceLocation(source_file, sheet, coordinate(row, column))


def _value(row: tuple[Any, ...], column: int) -> Any:
    return row[column] if column < len(row) else None


def _parse_cell_date(cell: Any) -> PartialDate | None:
    return parse_partial_date(cell.value, cell.number_format)


def read_well_history(
    path: Path,
    target_field: str,
) -> WellHistoryParseResult:
    """Read only the worksheet mapped to the currently selected target field."""

    target_key = normalize_field_name(target_field)
    sheet_name = next(
        (
            sheet
            for sheet, field_name in FIELD_SHEETS.items()
            if normalize_field_name(field_name) == target_key
        ),
        None,
    )
    if sheet_name is None:
        raise ValueError(f"Для месторождения {target_field!r} не настроено сопоставление листа.")

    source_file = str(path.resolve())
    issues: list[DiagnosticIssue] = []
    with openpyxl_read_warnings() as caught:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    issues.extend(
        DiagnosticIssue(
            code="OPENPYXL_WARNING",
            severity=IssueSeverity.WARNING,
            message=str(warning.message),
            source="well_history",
            locations=(SourceLocation(file=source_file),),
        )
        for warning in caught
    )
    try:
        if sheet_name not in workbook.sheetnames:
            issues.append(
                DiagnosticIssue(
                    code="WELL_HISTORY_SHEET_MISSING",
                    severity=IssueSeverity.WARNING,
                    message=f"В базе истории скважин отсутствует лист {sheet_name!r}.",
                    source="well_history",
                    locations=(sheet_location(source_file, sheet_name),),
                )
            )
            return WellHistoryParseResult((), tuple(issues))

        worksheet = workbook[sheet_name]
        schema = find_well_history_schema(worksheet)
        if schema is None:
            issues.append(
                DiagnosticIssue(
                    code="WELL_HISTORY_SCHEMA_INVALID",
                    severity=IssueSeverity.WARNING,
                    message="На листе истории не распознаны двухстрочные заголовки или блоки ГТМ.",
                    source="well_history",
                    locations=(sheet_location(source_file, sheet_name),),
                )
            )
            return WellHistoryParseResult((), tuple(issues))

        event_counts: Counter[str] = Counter()
        precision_counts: Counter[str] = Counter()
        histories: list[WellHistory] = []
        field_name = FIELD_SHEETS[sheet_name]
        for row_number, cells in enumerate(
            worksheet.iter_rows(min_row=3, max_row=worksheet.max_row, values_only=False),
            3,
        ):
            row = tuple(cell.value for cell in cells)
            raw_well = _value(row, schema.well_column)
            if raw_well is None or not str(raw_well).strip():
                continue
            display_well = display_value(raw_well)
            normalized_well = normalize_well_id(display_well)
            if normalized_well.base_well_id is None:
                issues.append(
                    DiagnosticIssue(
                        code="WELL_HISTORY_INVALID_WELL",
                        severity=IssueSeverity.WARNING,
                        message="Не удалось нормализовать номер скважины в базе истории.",
                        source="well_history",
                        locations=(
                            _cell_location(
                                source_file,
                                sheet_name,
                                row_number,
                                schema.well_column,
                            ),
                        ),
                    )
                )
                continue
            physical_root = normalized_well.base_well_id
            initial_date_cell = cells[schema.launch_date_column]
            launch_date = _parse_cell_date(initial_date_cell)
            launch_raw = initial_date_cell.value
            initial_plast = as_text(_value(row, schema.initial_plast_column))
            initial_plast_norm = normalize_plast(initial_plast)
            if launch_raw is not None:
                precision_counts[
                    launch_date.precision.value.upper() if launch_date else "INVALID"
                ] += 1
                if launch_date is None:
                    issues.append(
                        DiagnosticIssue(
                            code="WELL_HISTORY_INVALID_DATE",
                            severity=IssueSeverity.WARNING,
                            message=f"Не распознана дата ВНС: {launch_raw!r}.",
                            source="well_history",
                            well_id=physical_root,
                            locations=(
                                _cell_location(
                                    source_file,
                                    sheet_name,
                                    row_number,
                                    schema.launch_date_column,
                                ),
                            ),
                        )
                    )
            initial_location = _cell_location(
                source_file, sheet_name, row_number, schema.launch_date_column
            )
            events: list[WellEvent] = []
            stages: list[WellStage] = []
            raw_initial_block = tuple(_value(row, index) for index in schema.initial_block_columns)
            events.append(
                WellEvent(
                    field_name=field_name,
                    physical_root=physical_root,
                    source_display_well=display_well,
                    plast=initial_plast_norm,
                    event_kind=WellEventKind.INITIAL,
                    event_label="ВНС",
                    event_date=launch_date,
                    source_location=initial_location,
                    raw_block=raw_initial_block,
                )
            )
            if launch_date is not None:
                stages.append(
                    WellStage(
                        field_name=field_name,
                        physical_root=physical_root,
                        source_display_well=display_well,
                        target_well_id=normalized_well.target_well_id,
                        plast=initial_plast_norm,
                        event_kind=WellEventKind.INITIAL,
                        started_at=launch_date,
                        stage_sequence=0,
                        source_location=initial_location,
                        raw_block=raw_initial_block,
                    )
                )

            for label_column, date_column, block_number in schema.gtm_blocks:
                raw_block = tuple(_value(row, label_column + offset) for offset in range(6))
                if not any(value is not None and str(value).strip() for value in raw_block):
                    continue
                label = as_text(_value(row, label_column))
                kind = _event_kind(label)
                event_date_cell = cells[date_column]
                event_raw_date = event_date_cell.value
                event_date = _parse_cell_date(event_date_cell)
                event_plast = normalize_plast(as_text(_value(row, label_column + 2)))
                label_key = kind.value if kind is not WellEventKind.OTHER else (label or "OTHER")
                event_counts[label_key] += 1
                if event_raw_date is not None:
                    precision_counts[
                        event_date.precision.value.upper() if event_date else "INVALID"
                    ] += 1
                    if event_date is None:
                        issues.append(
                            DiagnosticIssue(
                                code="WELL_HISTORY_INVALID_DATE",
                                severity=IssueSeverity.WARNING,
                                message=(
                                    f"Не распознана дата события {label!r}: {event_raw_date!r}."
                                ),
                                source="well_history",
                                well_id=physical_root,
                                locations=(
                                    _cell_location(
                                        source_file, sheet_name, row_number, date_column
                                    ),
                                ),
                            )
                        )
                event_location = _cell_location(source_file, sheet_name, row_number, date_column)
                event = WellEvent(
                    field_name=field_name,
                    physical_root=physical_root,
                    source_display_well=display_well,
                    plast=event_plast,
                    event_kind=kind,
                    event_label=label,
                    event_date=event_date,
                    source_location=event_location,
                    raw_block=raw_block,
                    block_number=block_number,
                )
                events.append(event)
                if (
                    kind
                    in {
                        WellEventKind.ZBS,
                        WellEventKind.PVLG,
                        WellEventKind.PNLG,
                    }
                    and event_date is not None
                ):
                    zbs_ordinal = (
                        sum(stage.event_kind is WellEventKind.ZBS for stage in stages) + 1
                        if kind is WellEventKind.ZBS
                        else None
                    )
                    stage_identity = (
                        f"{physical_root}_ЗБС"
                        + (str(zbs_ordinal) if zbs_ordinal and zbs_ordinal > 1 else "")
                        if zbs_ordinal is not None
                        else None
                    )
                    stages.append(
                        WellStage(
                            field_name=field_name,
                            physical_root=physical_root,
                            source_display_well=display_well,
                            target_well_id=stage_identity,
                            plast=event_plast,
                            event_kind=kind,
                            started_at=event_date,
                            stage_sequence=len(stages),
                            source_location=event_location,
                            raw_block=raw_block,
                        )
                    )
            histories.append(
                WellHistory(
                    field_name=field_name,
                    physical_root=physical_root,
                    source_display_well=display_well,
                    events=tuple(events),
                    stages=tuple(stages),
                )
            )
        return WellHistoryParseResult(
            histories=tuple(histories),
            issues=tuple(issues),
            event_counts=dict(sorted(event_counts.items())),
            date_precision_counts=dict(sorted(precision_counts.items())),
        )
    finally:
        workbook.close()
