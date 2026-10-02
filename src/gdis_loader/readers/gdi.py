"""Reader and in-memory index preparation for GDI history."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gdis_loader.domain.models import DiagnosticIssue, GdiKind, GdiRecord, IssueSeverity
from gdis_loader.matching.normalize import normalize_field_name, normalize_plast, normalize_well_id
from gdis_loader.readers.common import (
    as_cell_value,
    as_text,
    cell_location,
    coordinate,
    display_value,
    index_headers,
    normalized_header,
    openpyxl_read_warnings,
    parse_datetime,
    parse_number,
    row_is_empty,
    row_value,
    sheet_location,
)
from gdis_loader.transforms.gdi import classify_gdi_kind

GDI_SHEET = "ГДИС_Сибнефтегаз в целом"
GDI_HEADER_ALIASES: dict[str, set[str]] = {
    "field": {"Месторождение"},
    "well": {"Скважина"},
    "plast": {"Пласт"},
    "research_type": {"Вид исследования"},
    "study_date": {"Дата"},
    "depth": {"Глубина записи, м"},
    "rfp_result": {"Выход на РФП (да/нет)"},
    "lnd_reliability": {"Достоверность согласно ЛНД"},
    "gtm_date": {"Дата проведения ГТМ"},
}


def _find_header(ws: Any) -> tuple[int, tuple[object, ...]] | None:
    rows = ws.iter_rows(min_row=1, max_row=40, max_col=100, values_only=True)  # type: ignore[attr-defined]
    for row_number, row in enumerate(rows, 1):
        values = tuple(row)
        if len(index_headers(values, GDI_HEADER_ALIASES)) == len(GDI_HEADER_ALIASES):
            return row_number, values
    return None


def _find_productivity_a_column(headers: tuple[object, ...]) -> int | None:
    """Find coefficient a by its source heading, including the documented AA form."""

    for index, value in enumerate(headers):
        header = normalized_header(value)
        if header.startswith("a,") or header in {"коэф а", "коэффициент а"}:
            return index
    return None


def _find_productivity_b_column(headers: tuple[object, ...]) -> int | None:
    for index, value in enumerate(headers):
        header = normalized_header(value)
        if header.startswith("b,") or header in {"коэф b", "коэффициент b"}:
            return index
    return None


def read_gdi(path: Path) -> tuple[tuple[GdiRecord, ...], tuple[DiagnosticIssue, ...]]:
    """Read only the required-width portion of the main GDI worksheet."""

    issues: list[DiagnosticIssue] = []
    source_file = str(path.resolve())
    with openpyxl_read_warnings() as caught:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    issues.extend(
        [
            DiagnosticIssue(
                code="OPENPYXL_WARNING",
                severity=IssueSeverity.WARNING,
                message=str(warning.message),
                source="gdi",
                locations=(sheet_location(source_file, GDI_SHEET),),
            )
            for warning in caught
        ]
    )
    try:
        if GDI_SHEET not in workbook.sheetnames:
            issues.append(
                DiagnosticIssue(
                    code="GDI_SHEET_MISSING",
                    severity=IssueSeverity.ERROR,
                    message=f"В книге ГДИ отсутствует лист {GDI_SHEET!r}.",
                    source="gdi",
                    locations=(sheet_location(source_file, GDI_SHEET),),
                )
            )
            return (), tuple(issues)
        worksheet = workbook[GDI_SHEET]
        header = _find_header(worksheet)
        if header is None:
            issues.append(
                DiagnosticIssue(
                    code="GDI_HEADERS_MISSING",
                    severity=IssueSeverity.ERROR,
                    message="На основном листе ГДИ не найдены все обязательные заголовки.",
                    source="gdi",
                    locations=(sheet_location(source_file, GDI_SHEET),),
                )
            )
            return (), tuple(issues)
        header_row, headers = header
        columns = index_headers(headers, GDI_HEADER_ALIASES)
        productivity_a_column = _find_productivity_a_column(headers)
        productivity_b_column = _find_productivity_b_column(headers)
        last_column = (
            max((*columns.values(), productivity_a_column or 0, productivity_b_column or 0)) + 1
        )
        records: list[GdiRecord] = []
        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=header_row + 1,
                max_row=worksheet.max_row,
                max_col=last_column,
                values_only=True,
            ),
            header_row + 1,
        ):
            values = tuple(row)
            identity_values = tuple(
                row_value(values, columns, name)
                for name in ("field", "well", "plast", "study_date")
            )
            if row_is_empty(identity_values):
                continue

            raw_field_value = row_value(values, columns, "field")
            raw_well_value = row_value(values, columns, "well")
            raw_plast_value = row_value(values, columns, "plast")
            raw_date_value = row_value(values, columns, "study_date")
            raw_field = as_text(raw_field_value)
            raw_well = display_value(raw_well_value) if raw_well_value is not None else None
            raw_plast = as_text(raw_plast_value)
            normalized_well = normalize_well_id(raw_well) if raw_well is not None else None
            base_id = normalized_well.base_well_id if normalized_well is not None else None
            target_well_id = normalized_well.target_well_id if normalized_well is not None else None
            study_date = parse_datetime(raw_date_value)
            if raw_date_value is not None and study_date is None:
                issues.append(
                    DiagnosticIssue(
                        code="GDI_INVALID_DATE",
                        severity=IssueSeverity.WARNING,
                        message=(
                            f"Некорректная дата ГДИ: {raw_date_value!r}; запись исключена "
                            "из выбора актуальной."
                        ),
                        source="gdi",
                        coordinate=coordinate(row_number, columns["study_date"]),
                        well_id=base_id,
                        locations=(
                            cell_location(
                                source_file,
                                GDI_SHEET,
                                row_number,
                                columns["study_date"],
                            ),
                        ),
                    )
                )
            elif raw_date_value is None and raw_well is not None:
                issues.append(
                    DiagnosticIssue(
                        code="GDI_MISSING_DATE",
                        severity=IssueSeverity.WARNING,
                        message=(
                            "В строке ГДИ отсутствует дата; запись исключена из выбора актуальной."
                        ),
                        source="gdi",
                        coordinate=coordinate(row_number, columns["study_date"]),
                        well_id=base_id,
                        locations=(
                            cell_location(
                                source_file,
                                GDI_SHEET,
                                row_number,
                                columns["study_date"],
                            ),
                        ),
                    )
                )
            if raw_well is None or base_id is None:
                issues.append(
                    DiagnosticIssue(
                        code="GDI_INVALID_WELL",
                        severity=IssueSeverity.WARNING,
                        message="В строке ГДИ отсутствует или не распознан номер скважины.",
                        source="gdi",
                        coordinate=coordinate(row_number, columns["well"]),
                        locations=(
                            cell_location(source_file, GDI_SHEET, row_number, columns["well"]),
                        ),
                    )
                )
            if raw_plast is None or normalize_plast(raw_plast) is None:
                issues.append(
                    DiagnosticIssue(
                        code="GDI_INVALID_PLAST",
                        severity=IssueSeverity.WARNING,
                        message=(
                            "В строке ГДИ отсутствует обозначение пласта; строка не используется "
                            "для точного match."
                        ),
                        source="gdi",
                        coordinate=coordinate(row_number, columns["plast"]),
                        well_id=base_id,
                        locations=(
                            cell_location(source_file, GDI_SHEET, row_number, columns["plast"]),
                        ),
                    )
                )
            raw_research_type = as_text(row_value(values, columns, "research_type"))
            gdi_kind = classify_gdi_kind(raw_research_type)
            if gdi_kind is GdiKind.UNKNOWN:
                issues.append(
                    DiagnosticIssue(
                        code="GDI_KIND_UNKNOWN",
                        severity=IssueSeverity.INFO,
                        message=(
                            f"Неизвестная подпись вида исследования {raw_research_type!r}; "
                            "запись не учитывается как ГДИ."
                        ),
                        source="gdi",
                        well_id=base_id,
                        locations=(
                            cell_location(
                                source_file,
                                GDI_SHEET,
                                row_number,
                                columns["research_type"],
                            ),
                        ),
                    )
                )
            depth_raw = row_value(values, columns, "depth")
            productivity_a_raw = (
                values[productivity_a_column]
                if productivity_a_column is not None and productivity_a_column < len(values)
                else None
            )
            productivity_b_raw = (
                values[productivity_b_column]
                if productivity_b_column is not None and productivity_b_column < len(values)
                else None
            )
            records.append(
                GdiRecord(
                    raw_field_name=raw_field,
                    normalized_field=normalize_field_name(raw_field)
                    if raw_field is not None
                    else None,
                    raw_well_name=raw_well,
                    base_well_id=base_id,
                    raw_plast=raw_plast,
                    normalized_plast=normalize_plast(raw_plast),
                    research_type=raw_research_type,
                    study_date=study_date,
                    rfp_result=as_text(row_value(values, columns, "rfp_result")),
                    lnd_reliability=as_text(row_value(values, columns, "lnd_reliability")),
                    gtm_date=parse_datetime(row_value(values, columns, "gtm_date")),
                    source_row=row_number,
                    source_file=source_file,
                    source_sheet=GDI_SHEET,
                    field_cell=coordinate(row_number, columns["field"]),
                    well_cell=coordinate(row_number, columns["well"]),
                    plast_cell=coordinate(row_number, columns["plast"]),
                    research_type_cell=coordinate(row_number, columns["research_type"]),
                    date_cell=coordinate(row_number, columns["study_date"]),
                    gtm_date_cell=coordinate(row_number, columns["gtm_date"]),
                    target_well_id=target_well_id,
                    gdi_kind=gdi_kind,
                    depth_raw_value=as_cell_value(depth_raw),
                    depth_value=parse_number(depth_raw),
                    depth_cell=coordinate(row_number, columns["depth"]),
                    rfp_cell=coordinate(row_number, columns["rfp_result"]),
                    lnd_reliability_cell=coordinate(row_number, columns["lnd_reliability"]),
                    productivity_a_raw=as_cell_value(productivity_a_raw),
                    productivity_a_value=parse_number(productivity_a_raw),
                    productivity_a_cell=(
                        coordinate(row_number, productivity_a_column)
                        if productivity_a_column is not None
                        else None
                    ),
                    productivity_b_raw=as_cell_value(productivity_b_raw),
                    productivity_b_value=parse_number(productivity_b_raw),
                    productivity_b_cell=(
                        coordinate(row_number, productivity_b_column)
                        if productivity_b_column is not None
                        else None
                    ),
                    source_rows=(row_number,),
                )
            )
        return tuple(records), tuple(issues)
    finally:
        workbook.close()
