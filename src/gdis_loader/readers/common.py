"""Small reader utilities shared by all Excel source readers."""

from __future__ import annotations

import math
import re
import unicodedata
import warnings
from collections.abc import Iterable, Iterator, Sequence
from contextlib import contextmanager
from datetime import date, datetime, time

from openpyxl.utils.cell import get_column_letter
from openpyxl.utils.datetime import from_excel

from gdis_loader.domain.models import CellValue, SourceLocation

_SPACE_RE = re.compile(r"\s+")
OPENPYXL_CONDITIONAL_FORMATTING_WARNING = (
    "Conditional Formatting extension is not supported and will be removed"
)


@contextmanager
def openpyxl_read_warnings() -> Iterator[list[warnings.WarningMessage]]:
    """Capture useful reader warnings but hide one harmless read-only warning.

    The application never saves these workbooks through openpyxl.  Its
    conditional-formatting extension warning describes a possible future save,
    not a problem with the input currently being read.  Other warnings remain
    visible in diagnostics.
    """

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        warnings.filterwarnings(
            "ignore",
            message=re.escape(OPENPYXL_CONDITIONAL_FORMATTING_WARNING),
        )
        yield caught


def clean_text(value: object) -> str:
    """Normalize whitespace for labels while preserving meaningful content."""

    return _SPACE_RE.sub(" ", str(value).replace("\u00a0", " ").strip())


def normalized_header(value: object) -> str:
    """Compare headers independent of case, whitespace, and Unicode superscripts."""

    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", clean_text(value)).casefold()
    return _SPACE_RE.sub(" ", text)


def find_header_row(
    rows: Iterable[Sequence[object]],
    required: set[str],
    max_rows: int,
) -> tuple[int, tuple[object, ...]] | None:
    """Find the first row containing every normalized required header."""

    for row_number, row in enumerate(rows, 1):
        if row_number > max_rows:
            break
        values = tuple(row)
        present = {normalized_header(value) for value in values if value is not None}
        if required <= present:
            return row_number, values
    return None


def index_headers(
    headers: Sequence[object],
    aliases: dict[str, set[str]],
) -> dict[str, int]:
    """Map semantic names to columns using exact normalized aliases."""

    normalized = {
        normalized_header(value): index for index, value in enumerate(headers) if value is not None
    }
    result: dict[str, int] = {}
    for semantic_name, candidates in aliases.items():
        for candidate in candidates:
            index = normalized.get(normalized_header(candidate))
            if index is not None:
                result[semantic_name] = index
                break
    return result


def missing_header_names(
    headers: Sequence[object],
    aliases: dict[str, set[str]],
) -> list[str]:
    """Return semantic header names not found in a header row."""

    return [name for name in aliases if name not in index_headers(headers, aliases)]


def coordinate(row_number: int, column_index: int) -> str:
    return f"{get_column_letter(column_index + 1)}{row_number}"


def cell_location(
    source_file: str, source_sheet: str | None, row_number: int, column_index: int
) -> SourceLocation:
    """Return a structured source location for a zero-based Excel column."""

    return SourceLocation(
        file=source_file,
        sheet=source_sheet,
        cell=coordinate(row_number, column_index),
    )


def sheet_location(source_file: str, source_sheet: str | None) -> SourceLocation:
    """Return a workbook/sheet location when no exact cell is available."""

    return SourceLocation(file=source_file, sheet=source_sheet, cell=None)


def row_is_empty(values: Sequence[object]) -> bool:
    return not any(value is not None and clean_text(value) for value in values)


def row_value(values: Sequence[object], columns: dict[str, int], name: str) -> object:
    index = columns[name]
    return values[index] if index < len(values) else None


def as_text(value: object) -> str | None:
    if value is None:
        return None
    text = clean_text(value)
    return text or None


def as_cell_value(value: object) -> CellValue:
    """Keep supported raw Excel values typed and stringify uncommon objects."""

    if value is None or isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    return str(value)


def parse_datetime(value: object) -> datetime | None:
    """Parse common Excel dates without treating malformed values as fatal."""

    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            return None
        try:
            parsed = from_excel(value)
        except TypeError, ValueError, OverflowError:
            return None
        if isinstance(parsed, datetime):
            return parsed
        if isinstance(parsed, date):
            return datetime.combine(parsed, time.min)
        return None
    text = clean_text(value)
    if not text:
        return None
    iso_text = text.replace("Z", "+00:00")
    try:
        parsed_iso = datetime.fromisoformat(iso_text)
        return parsed_iso.replace(tzinfo=None)
    except ValueError:
        pass
    for pattern in ("%d.%m.%Y", "%d.%m.%Y %H:%M", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    return None


def parse_number(value: object) -> float | None:
    """Return a finite numeric value; zero is deliberately valid."""

    if isinstance(value, bool) or is_missing_measurement(value):
        return None
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        text = clean_text(value)
        text = text.replace(" ", "").replace(",", ".")
        try:
            result = float(text)
        except ValueError:
            return None
    return result if math.isfinite(result) else None


def is_missing_measurement(value: object) -> bool:
    """Return whether an Excel cell is an expected missing numeric value.

    Excel error tokens and the placeholders used in the source workbooks are
    not data-quality warnings by themselves.  A caller can still distinguish
    them from arbitrary text and report the latter as a useful source note.
    """

    if value is None:
        return True
    if not isinstance(value, str):
        return False
    text = clean_text(value)
    return not text or text in {"-", "—", "–"} or text.startswith("#")


def display_value(value: object) -> str:
    """Keep a readable source representation, including string leading zeroes."""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)
