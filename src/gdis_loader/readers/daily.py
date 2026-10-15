"""Reader/validator for daily telemetry without applying a debit aggregation."""

from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gdis_loader.domain.models import (
    DailyTelemetryRow,
    DiagnosticIssue,
    IssueSeverity,
    SourceLocation,
)
from gdis_loader.matching.normalize import normalize_well_id
from gdis_loader.readers.common import (
    as_text,
    cell_location,
    coordinate,
    display_value,
    is_missing_measurement,
    openpyxl_read_warnings,
    parse_number,
    sheet_location,
)
from gdis_loader.readers.daily_context import day_contexts, reference_evidence

DAILY_SHEET = "Добыча телеметрия"
DAILY_EMPTY_ROW_LIMIT = 200


def _find_daily_header(ws: Any) -> tuple[int, tuple[object, ...]] | None:
    for row_number, row in enumerate(
        ws.iter_rows(min_row=1, max_row=20, max_col=100, values_only=True),
        1,  # type: ignore[attr-defined]
    ):
        values = tuple(row)
        labels = {text.casefold() for value in values[:2] if (text := as_text(value)) is not None}
        if {"гск", "скв."} <= labels:
            return row_number, values
    return None


def read_daily_report(
    path: Path,
) -> tuple[tuple[DailyTelemetryRow, ...], tuple[DiagnosticIssue, ...]]:
    """Read daily values and retain errors as missing, never as debit."""

    issues: list[DiagnosticIssue] = []
    source_file = str(path.resolve())
    with openpyxl_read_warnings() as caught:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    issues.extend(
        [
            DiagnosticIssue(
                code="OPENPYXL_WARNING",
                severity=IssueSeverity.WARNING,
                message=str(warning.message),
                source="daily_report",
                locations=(sheet_location(source_file, DAILY_SHEET),),
            )
            for warning in caught
        ]
    )
    try:
        names = [
            name
            for name in workbook.sheetnames
            if name.strip().casefold() == DAILY_SHEET.casefold()
        ]
        if len(names) != 1:
            missing = not names
            code = "DAILY_SHEET_MISSING" if missing else "DAILY_SHEET_AMBIGUOUS"
            message = (
                f"В суточном рапорте отсутствует лист {DAILY_SHEET!r}."
                if missing
                else f"В суточном рапорте несколько листов совпадают с {DAILY_SHEET!r}."
            )
            issues.append(
                DiagnosticIssue(
                    code=code,
                    severity=IssueSeverity.ERROR,
                    message=message,
                    source="daily_report",
                    locations=tuple(sheet_location(source_file, name) for name in names)
                    or (sheet_location(source_file, DAILY_SHEET),),
                )
            )
            return (), tuple(issues)
        sheet_name = names[0]
        worksheet = workbook[sheet_name]
        formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            contexts = day_contexts(formulas)
            formula_rows = list(formulas[sheet_name].iter_rows(max_col=100, values_only=True))
        finally:
            formulas.close()
        months = {context.report_date.strftime("%Y-%m") for context in contexts.values()}
        if len(months) > 1:
            return (), (
                DiagnosticIssue(
                    code="DAILY_PERIOD_CONFLICT",
                    severity=IssueSeverity.ERROR,
                    message="Датированные вкладки рапорта относятся к разным месяцам.",
                    source="daily_report",
                    locations=(sheet_location(source_file, sheet_name),),
                ),
            )
        report_month = next(iter(months), None)
        header = _find_daily_header(worksheet)
        if header is None:
            issues.append(
                DiagnosticIssue(
                    code="DAILY_HEADERS_MISSING",
                    severity=IssueSeverity.ERROR,
                    message="На листе суточного рапорта не найдены заголовки ГСК и СКВ.",
                    source="daily_report",
                    locations=(sheet_location(source_file, DAILY_SHEET),),
                )
            )
            return (), tuple(issues)
        header_row, headers = header
        day_columns: list[tuple[int, int]] = []
        for column_index, value in enumerate(headers[2:], 2):
            if isinstance(value, bool):
                continue
            try:
                day = int(value) if isinstance(value, (int, float, str)) else 0
            except TypeError, ValueError:
                continue
            if 1 <= day <= 31:
                day_columns.append((day, column_index))
        if report_month is not None:
            import calendar

            year, month = map(int, report_month.split("-"))
            day_columns = [
                (day, column)
                for day, column in day_columns
                if day <= calendar.monthrange(year, month)[1]
            ]
        if len(day_columns) < 28:
            issues.append(
                DiagnosticIssue(
                    code="DAILY_DAY_HEADERS_INCOMPLETE",
                    severity=IssueSeverity.INFO,
                    message=(
                        f"На листе суточного рапорта найдено только {len(day_columns)} "
                        "дневных колонок."
                    ),
                    source="daily_report",
                    locations=(sheet_location(source_file, DAILY_SHEET),),
                )
            )
        last_column = max((column for _, column in day_columns), default=1) + 1
        rows: list[DailyTelemetryRow] = []
        blank_streak = 0
        invalid_counts: Counter[str] = Counter()
        invalid_samples: list[str] = []
        missing_counts: Counter[str] = Counter()
        invalid_locations: list[SourceLocation] = []
        cache_missing_locations: list[SourceLocation] = []
        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=header_row + 1,
                max_row=worksheet.max_row,
                max_col=last_column,
                values_only=True,
            ),
            header_row + 1,
        ):
            values = tuple(row)
            field_value = values[0] if values else None
            collector_value = display_value(field_value) if field_value is not None else None
            well_value = values[1] if len(values) > 1 else None
            if field_value is None and well_value is None:
                blank_streak += 1
                if blank_streak >= DAILY_EMPTY_ROW_LIMIT:
                    break
                continue
            blank_streak = 0
            raw_field = as_text(field_value)
            raw_well = display_value(well_value) if well_value is not None else None
            normalized = normalize_well_id(raw_well) if raw_well is not None else None
            evidence: list[SourceLocation] = []
            fields: set[str] = set()
            linked_days = 0
            for day, column in day_columns:
                formula = formula_rows[row_number - 1][column]
                proof = reference_evidence(
                    formula,
                    day,
                    normalized.target_well_id if normalized else None,
                    contexts,
                    source_file,
                )
                if proof is not None:
                    fields.add(proof[0])
                    evidence.extend(proof[1])
                    linked_days += 1
            # ГСК is a collector, never a field. Only complete, consistent links
            # across the reporting period can establish this row's field scope.
            raw_field = (
                next(iter(fields)) if len(fields) == 1 and linked_days == len(day_columns) else None
            )
            if raw_field is None:
                evidence.clear()
            daily_values: list[float | None] = []
            for _day, column_index in day_columns:
                raw_value = values[column_index] if column_index < len(values) else None
                parsed = parse_number(raw_value)
                formula = formula_rows[row_number - 1][column_index]
                if raw_value is None and isinstance(formula, str) and formula.startswith("="):
                    cache_missing_locations.append(
                        cell_location(source_file, sheet_name, row_number, column_index)
                    )
                if raw_value is not None and parsed is None:
                    invalid_locations.append(
                        cell_location(source_file, sheet_name, row_number, column_index)
                    )
                    if is_missing_measurement(raw_value):
                        missing_counts["daily"] += 1
                    else:
                        invalid_counts["daily"] += 1
                        if len(invalid_samples) < 3:
                            invalid_samples.append(repr(raw_value))
                daily_values.append(parsed)
            rows.append(
                DailyTelemetryRow(
                    raw_field_name=raw_field,
                    raw_well_name=raw_well,
                    base_well_id=normalized.base_well_id if normalized else None,
                    target_well_id=normalized.target_well_id if normalized else None,
                    source_row=row_number,
                    daily_values=tuple(daily_values),
                    source_file=source_file,
                    source_sheet=sheet_name,
                    field_evidence=tuple(dict.fromkeys(evidence)),
                    report_month=report_month,
                    field_scope_invalid=report_month is not None and raw_field is None,
                    raw_collector_id=collector_value,
                    field_cell=coordinate(row_number, 0),
                    well_cell=coordinate(row_number, 1),
                    daily_cells=tuple(
                        coordinate(row_number, column_index) for _day, column_index in day_columns
                    ),
                )
            )
        if missing_counts["daily"] or invalid_counts["daily"]:
            details: list[str] = []
            if missing_counts["daily"]:
                details.append(
                    f"{missing_counts['daily']} пустых, placeholder- или Excel error-значений"
                )
            if invalid_counts["daily"]:
                details.append(
                    f"{invalid_counts['daily']} содержательных нечисловых значений "
                    f"(примеры: {', '.join(invalid_samples)})"
                )
            issues.append(
                DiagnosticIssue(
                    code="DAILY_INVALID_VALUE",
                    severity=IssueSeverity.INFO,
                    message=(
                        "В суточном рапорте прочитано как пропуск: "
                        + "; ".join(details)
                        + ". Суточные значения участвуют в расчёте среднего Дебита по правилу "
                        "работающего дня: учитываются числовые ставки строго выше порога."
                    ),
                    source="daily_report",
                    locations=tuple(invalid_locations),
                )
            )
        if cache_missing_locations:
            issues.append(
                DiagnosticIssue(
                    code="DAILY_FORMULA_CACHE_MISSING",
                    severity=IssueSeverity.INFO,
                    message="Формулы без сохранённого результата прочитаны как пропуски.",
                    source="daily_report",
                    locations=tuple(cache_missing_locations),
                )
            )
        return tuple(rows), tuple(issues)
    finally:
        workbook.close()
