"""Reader for per-research GDM adaptation deltas."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gdis_loader.domain.dates import parse_partial_date
from gdis_loader.domain.models import (
    DiagnosticIssue,
    IssueSeverity,
    ResearchStatusRecord,
    SourceLocation,
)
from gdis_loader.matching.normalize import normalize_plast, normalize_well_id
from gdis_loader.readers.common import (
    as_cell_value,
    as_text,
    cell_location,
    coordinate,
    display_value,
    normalized_header,
    openpyxl_read_warnings,
    parse_number,
    row_is_empty,
    sheet_location,
)
from gdis_loader.readers.well_history import FIELD_SHEETS

STATUS_SHEETS: dict[str, str] = {
    "БЛУ": "АР_таблица_БЛУ",
    "ПЛУ": "АР_таблица_ПЛУ",
    "ХЛУ": "АР_таблица_ХЛУ",
}
_HEADER_ALIASES: dict[str, set[str]] = {
    "well": {"Скважина"},
    "plast": {"Пласт"},
    "research_type": {"Исследования", "Исследвоания"},
    "study_date": {"Дата исследования"},
    "delta": {"Дельта"},
}


@dataclass(frozen=True, slots=True)
class ResearchStatusSchema:
    """Detected columns and the first source row after the multirow header."""

    columns: dict[str, int]
    data_start_row: int


def find_research_status_schema(worksheet: Any) -> ResearchStatusSchema | None:
    """Resolve semantic columns across the source's two-level headings."""

    max_column = min(max(worksheet.max_column, 10), 20)  # type: ignore[attr-defined]
    header_rows = min(8, worksheet.max_row)  # type: ignore[attr-defined]
    rows = tuple(
        tuple(row)
        for row in worksheet.iter_rows(
            min_row=1,
            max_row=header_rows,
            max_col=max_column,
            values_only=True,
        )
    )
    columns: dict[str, int] = {}
    used_header_rows: list[int] = []
    for semantic, aliases in _HEADER_ALIASES.items():
        normalized_aliases = {normalized_header(alias) for alias in aliases}
        for column_index in range(max_column):
            labels = [
                normalized_header(row[column_index])
                for row in rows
                if column_index < len(row) and row[column_index] is not None
            ]
            if any(label in normalized_aliases for label in labels):
                columns[semantic] = column_index
                used_header_rows.extend(
                    index + 1
                    for index, row in enumerate(rows)
                    if column_index < len(row)
                    and normalized_header(row[column_index]) in normalized_aliases
                )
                break
    if len(columns) != len(_HEADER_ALIASES):
        return None
    data_start_row = max(used_header_rows) + 1
    # Some workbooks put labels such as "Объект" and units immediately below
    # the semantic headers. Do not parse that level as a research record.
    while data_start_row <= worksheet.max_row:  # type: ignore[attr-defined]
        row = next(
            worksheet.iter_rows(
                min_row=data_start_row,
                max_row=data_start_row,
                min_col=columns["well"] + 1,
                max_col=columns["well"] + 1,
                values_only=True,
            ),
            (),
        )
        well_label = normalized_header(row[0] if row else None)
        if well_label not in {"объект", "параметры", "скважина"}:
            break
        data_start_row += 1
    return ResearchStatusSchema(columns, data_start_row)


def _delta_is_missing(value: object) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return not text or text in {"-", "—", "–"}


def read_research_status(
    path: Path,
    lu_code: str,
) -> tuple[tuple[ResearchStatusRecord, ...], tuple[DiagnosticIssue, ...]]:
    """Read one LU sheet while retaining raw values and exact cell provenance."""

    if lu_code not in STATUS_SHEETS or lu_code not in FIELD_SHEETS:
        raise ValueError(f"Неизвестный код лицензионного участка: {lu_code!r}.")
    sheet_name = STATUS_SHEETS[lu_code]
    source_file = str(path.resolve())
    issues: list[DiagnosticIssue] = []
    with openpyxl_read_warnings() as caught:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        formula_workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    issues.extend(
        DiagnosticIssue(
            code="OPENPYXL_WARNING",
            severity=IssueSeverity.WARNING,
            message=str(warning.message),
            source="research_status",
            locations=(sheet_location(source_file, sheet_name),),
        )
        for warning in caught
    )
    try:
        if sheet_name not in workbook.sheetnames:
            issues.append(
                DiagnosticIssue(
                    code="GDM_STATUS_SHEET_MISSING",
                    severity=IssueSeverity.WARNING,
                    message=f"В книге статуса исследований отсутствует лист {sheet_name!r}.",
                    source="research_status",
                    locations=(sheet_location(source_file, sheet_name),),
                )
            )
            return (), tuple(issues)
        worksheet = workbook[sheet_name]
        formula_worksheet = formula_workbook[sheet_name]
        schema = find_research_status_schema(worksheet)
        if schema is None:
            issues.append(
                DiagnosticIssue(
                    code="GDM_STATUS_HEADERS_MISSING",
                    severity=IssueSeverity.WARNING,
                    message="Не распознаны многострочные заголовки таблицы статуса исследований.",
                    source="research_status",
                    locations=(sheet_location(source_file, sheet_name),),
                )
            )
            return (), tuple(issues)

        max_column = max(schema.columns.values()) + 1
        formula_rows = formula_worksheet.iter_rows(
            min_row=schema.data_start_row,
            max_row=worksheet.max_row,
            max_col=max_column,
            values_only=False,
        )
        formula_iter = iter(formula_rows)
        records: list[ResearchStatusRecord] = []
        invalid_well_locations: list[SourceLocation] = []
        invalid_date_locations: list[SourceLocation] = []
        invalid_delta_locations: list[SourceLocation] = []
        missing_delta_locations: list[SourceLocation] = []
        formula_cache_locations: list[SourceLocation] = []
        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=schema.data_start_row,
                max_row=worksheet.max_row,
                max_col=max_column,
                values_only=False,
            ),
            schema.data_start_row,
        ):
            formula_row = next(formula_iter, ())
            values = tuple(cell.value for cell in row)
            column = schema.columns
            relevant = tuple(
                values[index] if index < len(values) else None for index in column.values()
            )
            if row_is_empty(relevant):
                continue

            raw_well_value = values[column["well"]]
            raw_well = display_value(raw_well_value) if raw_well_value is not None else None
            normalized_well = normalize_well_id(raw_well) if raw_well is not None else None
            raw_plast = as_text(values[column["plast"]])
            date_cell = row[column["study_date"]]
            study_date = parse_partial_date(date_cell.value, date_cell.number_format)
            delta_index = column["delta"]
            delta_cell = row[delta_index]
            raw_delta = delta_cell.value
            delta_value = parse_number(raw_delta)
            formula_value = (
                formula_row[delta_index].value if delta_index < len(formula_row) else None
            )
            formula_without_cache = (
                isinstance(formula_value, str)
                and formula_value.startswith("=")
                and raw_delta is None
            )
            location = cell_location(source_file, sheet_name, row_number, delta_index)
            if formula_without_cache:
                formula_cache_locations.append(location)
            elif delta_value is None:
                if _delta_is_missing(raw_delta):
                    missing_delta_locations.append(location)
                else:
                    invalid_delta_locations.append(location)
            if raw_well is not None and normalized_well is None:
                invalid_well_locations.append(
                    cell_location(source_file, sheet_name, row_number, column["well"])
                )
            if date_cell.value is not None and study_date is None:
                invalid_date_locations.append(
                    cell_location(source_file, sheet_name, row_number, column["study_date"])
                )
            records.append(
                ResearchStatusRecord(
                    raw_well_name=raw_well,
                    base_well_id=normalized_well.base_well_id if normalized_well else None,
                    target_well_id=normalized_well.target_well_id if normalized_well else None,
                    raw_plast=raw_plast,
                    normalized_plast=normalize_plast(raw_plast),
                    research_type=as_text(values[column["research_type"]]),
                    study_date=study_date,
                    delta_raw=as_cell_value(raw_delta),
                    delta_value=delta_value,
                    delta_formula_without_cache=formula_without_cache,
                    source_row=row_number,
                    source_file=source_file,
                    source_sheet=sheet_name,
                    well_cell=coordinate(row_number, column["well"]),
                    plast_cell=coordinate(row_number, column["plast"]),
                    research_type_cell=coordinate(row_number, column["research_type"]),
                    date_cell=coordinate(row_number, column["study_date"]),
                    delta_cell=coordinate(row_number, delta_index),
                )
            )

        for code, severity, message, locations in (
            (
                "GDM_STATUS_WELL_INVALID",
                IssueSeverity.WARNING,
                "Не удалось нормализовать номера скважин в строках статуса ГДМ.",
                invalid_well_locations,
            ),
            (
                "GDM_STATUS_DATE_INVALID",
                IssueSeverity.WARNING,
                "Не распознаны даты исследования в строках статуса ГДМ.",
                invalid_date_locations,
            ),
            (
                "GDM_DELTA_INVALID",
                IssueSeverity.WARNING,
                "Некорректные значения Delta оставлены неизвестными.",
                invalid_delta_locations,
            ),
            (
                "GDM_DELTA_MISSING",
                IssueSeverity.INFO,
                "Пустые/дефисные значения Delta оставлены неизвестными.",
                missing_delta_locations,
            ),
            (
                "GDM_DELTA_FORMULA_CACHE_MISSING",
                IssueSeverity.WARNING,
                "Формулы Delta без cached numeric value оставлены неизвестными.",
                formula_cache_locations,
            ),
        ):
            if locations:
                issues.append(
                    DiagnosticIssue(
                        code=code,
                        severity=severity,
                        message=f"{len(locations)} строк. {message}",
                        source="research_status",
                        locations=tuple(locations),
                    )
                )
        return tuple(records), tuple(issues)
    finally:
        workbook.close()
        formula_workbook.close()
