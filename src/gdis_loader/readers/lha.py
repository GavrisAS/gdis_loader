"""Streaming reader and index for the large LHA workbook."""

from __future__ import annotations

import re
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gdis_loader.domain.models import (
    DiagnosticIssue,
    IssueSeverity,
    LhaRecord,
    LhaSelection,
    LhaValues,
    SourceLocation,
    WellHistory,
    WellStage,
)
from gdis_loader.matching.lifecycle import assign_stage_by_date
from gdis_loader.matching.normalize import normalize_field_name, normalize_plast, normalize_well_id
from gdis_loader.readers.common import (
    as_cell_value,
    as_text,
    cell_location,
    coordinate,
    display_value,
    index_headers,
    is_missing_measurement,
    openpyxl_read_warnings,
    parse_datetime,
    parse_number,
    row_is_empty,
    row_value,
    sheet_location,
)

LHA_SHEET = "БД Сибнефтегаз"
LHA_EMPTY_IDENTITY_LIMIT = 2_000
LHA_HEADER_ALIASES: dict[str, set[str]] = {
    "field": {"Месторождение"},
    "well": {"Номер скважины"},
    "plast": {"Пласт"},
    "measured_at": {"Дата и время исследования"},
    "cl": {
        "Данные ЛХА, CL-, мг/дм3",
        "Данные ЛХА, CL-, мг/дм³",
    },
    "vgf": {"ВГФ, г/м3", "ВГФ, г/м³"},
}


def _find_header(ws: Any) -> tuple[int, tuple[object, ...]] | None:
    for row_number, row in enumerate(
        ws.iter_rows(min_row=1, max_row=40, max_col=100, values_only=True),
        1,  # type: ignore[attr-defined]
    ):
        values = tuple(row)
        if len(index_headers(values, LHA_HEADER_ALIASES)) == len(LHA_HEADER_ALIASES):
            return row_number, values
    return None


class LhaIndex:
    """One-pass index grouped by lifecycle identity, field, and exact plast."""

    def __init__(self, records: Iterable[LhaRecord]) -> None:
        grouped: defaultdict[tuple[str, str | None, str], list[LhaRecord]] = defaultdict(list)
        retained: list[LhaRecord] = []
        for record in records:
            target_id = record.target_well_id or (
                normalize_well_id(record.raw_well_name).target_well_id
                if record.raw_well_name is not None
                else None
            )
            if target_id is None or record.normalized_plast is None or record.measured_at is None:
                continue
            grouped[(target_id, record.normalized_field, record.normalized_plast)].append(record)
            retained.append(record)
        self._by_key = {key: tuple(items) for key, items in grouped.items()}
        self._records = tuple(retained)

    def select(
        self,
        *,
        base_well_id: str,
        normalized_field: str,
        normalized_plast: str,
        cutoff: datetime | None = None,
        target_well_id: str | None = None,
        physical_root: str | None = None,
        histories: tuple[WellHistory, ...] = (),
        lifecycle_stage: WellStage | None = None,
    ) -> LhaSelection:
        """Select the last two distinct timestamps independently for Cl and VGF.

        Rows with the same timestamp are collapsed by metric using the maximum
        valid numeric value. This exact identity lookup never joins a suffixed
        LHA row to its physical root.
        """

        identity = target_well_id or base_well_id
        exact = self._by_key.get((identity, normalized_field, normalized_plast), ())
        if not exact:
            # A few legacy LHA rows have no field value.  They can be used only
            # when no field-qualified row exists; rows from another field never
            # cross the boundary.
            exact = self._by_key.get((identity, None, normalized_plast), ())
        eligible = [
            record
            for record in exact
            if cutoff is None or (record.measured_at is not None and record.measured_at <= cutoff)
        ]
        lifecycle_issues: list[DiagnosticIssue] = []
        root = physical_root or base_well_id
        if identity != root and lifecycle_stage is not None:
            # Legacy LHA rows can omit a confirmed suffix. Attach those rows
            # only when field, plast, timestamp, and the history timeline all
            # resolve them to the same lifecycle stage as the selected GDI.
            root_records = tuple(
                record
                for record in self._records
                if (
                    record.target_well_id
                    or (
                        normalize_well_id(record.raw_well_name).target_well_id
                        if record.raw_well_name is not None
                        else None
                    )
                )
                == root
                and record.normalized_field == normalized_field
                and record.normalized_plast == normalized_plast
                and (cutoff is None or (record.measured_at and record.measured_at <= cutoff))
            )
            for record in root_records:
                if record.measured_at is None:
                    continue
                location = SourceLocation(
                    file=record.source_file,
                    sheet=record.source_sheet,
                    cell=record.date_cell,
                )
                assignment = assign_stage_by_date(
                    physical_root=root,
                    normalized_field=normalized_field,
                    normalized_plast=normalized_plast,
                    observed_at=record.measured_at,
                    histories=histories,
                    observation_location=location,
                )
                if assignment.issue is not None:
                    lifecycle_issues.append(assignment.issue)
                if assignment.stage == lifecycle_stage:
                    eligible.append(record)

        cl_current, cl_previous, cl_current_locations, cl_previous_locations, cl_issues = (
            _select_metric(
                eligible,
                metric="cl_value",
                metric_label="Cl",
                base_well_id=base_well_id,
                normalized_field=normalized_field,
                normalized_plast=normalized_plast,
            )
        )
        vgf_current, vgf_previous, vgf_current_locations, vgf_previous_locations, vgf_issues = (
            _select_metric(
                eligible,
                metric="vgf_value",
                metric_label="ВГФ",
                base_well_id=base_well_id,
                normalized_field=normalized_field,
                normalized_plast=normalized_plast,
            )
        )
        return LhaSelection(
            values=LhaValues(
                cl_current,
                cl_previous,
                vgf_current,
                vgf_previous,
                cl_current_locations,
                cl_previous_locations,
                vgf_current_locations,
                vgf_previous_locations,
            ),
            issues=(*lifecycle_issues, *cl_issues, *vgf_issues),
        )


@dataclass(frozen=True, slots=True)
class _MetricPoint:
    measured_at: datetime
    value: float
    values: tuple[float, ...]
    records: tuple[LhaRecord, ...]
    selected_record: LhaRecord


def _select_metric(
    records: Iterable[LhaRecord],
    *,
    metric: str,
    metric_label: str,
    base_well_id: str,
    normalized_field: str,
    normalized_plast: str,
) -> tuple[
    float | None,
    float | None,
    tuple[SourceLocation, ...],
    tuple[SourceLocation, ...],
    tuple[DiagnosticIssue, ...],
]:
    records_by_timestamp: defaultdict[datetime, list[LhaRecord]] = defaultdict(list)
    values_by_timestamp: defaultdict[datetime, list[tuple[float, LhaRecord]]] = defaultdict(list)
    for record in records:
        measured_at = record.measured_at
        value = getattr(record, metric)
        if measured_at is not None:
            records_by_timestamp[measured_at].append(record)
            if value is not None:
                values_by_timestamp[measured_at].append((value, record))

    points_list: list[_MetricPoint] = []
    duplicate_issues: list[DiagnosticIssue] = []
    for measured_at in sorted(records_by_timestamp, reverse=True):
        all_records = tuple(sorted(records_by_timestamp[measured_at], key=_record_location_key))
        valid_values = values_by_timestamp.get(measured_at, [])
        if not valid_values:
            continue
        selected_value, selected_record = min(
            valid_values,
            key=lambda item: (-item[0], _record_location_key(item[1])),
        )
        point = _MetricPoint(
            measured_at=measured_at,
            value=selected_value,
            values=tuple(sorted(value for value, _ in valid_values)),
            records=all_records,
            selected_record=selected_record,
        )
        points_list.append(point)
        if len(all_records) > 1:
            duplicate_issues.append(
                _duplicate_timestamp_resolved_issue(
                    point,
                    metric=metric,
                    metric_label=metric_label,
                    base_well_id=base_well_id,
                    normalized_field=normalized_field,
                    normalized_plast=normalized_plast,
                )
            )
    points = tuple(points_list)
    if not points:
        return None, None, (), (), ()
    current = points[0].value
    current_locations = points[0].selected_record.metric_locations(metric)
    previous = None
    if len(points) <= 1:
        return current, previous, current_locations, (), tuple(duplicate_issues)
    previous = points[1].value
    previous_locations = points[1].selected_record.metric_locations(metric)
    return current, previous, current_locations, previous_locations, tuple(duplicate_issues)


def _record_location_key(record: LhaRecord) -> tuple[str, str, int, str]:
    cell = record.cl_cell or record.vgf_cell or ""
    match = re.search(r"(\d+)$", cell)
    row_number = int(match.group(1)) if match else 0
    return (record.source_file, record.source_sheet or "", row_number, cell)


def _duplicate_timestamp_resolved_issue(
    point: _MetricPoint,
    *,
    metric: str,
    metric_label: str,
    base_well_id: str,
    normalized_field: str,
    normalized_plast: str,
) -> DiagnosticIssue:
    selected_locations = point.selected_record.metric_locations(metric)
    selected_text = selected_locations[0].cell if selected_locations else "неизвестная ячейка"
    return DiagnosticIssue(
        code="LHA_DUPLICATE_TIMESTAMP_RESOLVED_MAX",
        severity=IssueSeverity.INFO,
        message=(
            f"ЛХА: скважина {base_well_id!r}, месторождение {normalized_field!r}, "
            f"пласт {normalized_plast!r}, метрика {metric_label}: для timestamp "
            f"{point.measured_at.isoformat(sep=' ')} выбрано максимальное валидное "
            f"значение из {len(point.records)} строк; ячейка {selected_text}."
        ),
        source="lha",
        well_id=base_well_id,
        locations=tuple(
            dict.fromkeys(
                location for record in point.records for location in record.metric_locations(metric)
            )
        ),
    )


def read_lha(path: Path) -> tuple[tuple[LhaRecord, ...], tuple[DiagnosticIssue, ...]]:
    """Read the relevant columns and stop after a long empty identity tail."""

    issues: list[DiagnosticIssue] = []
    source_file = str(path.resolve())
    with openpyxl_read_warnings() as caught:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    issues.extend(
        [
            DiagnosticIssue(
                code="OPENPYXL_WARNING",
                severity=IssueSeverity.WARNING,
                message=str(warning.message),
                source="lha",
                locations=(sheet_location(source_file, LHA_SHEET),),
            )
            for warning in caught
        ]
    )
    try:
        if LHA_SHEET not in workbook.sheetnames:
            issues.append(
                DiagnosticIssue(
                    code="LHA_SHEET_MISSING",
                    severity=IssueSeverity.ERROR,
                    message=f"В книге ЛХА отсутствует лист {LHA_SHEET!r}.",
                    source="lha",
                    locations=(sheet_location(source_file, LHA_SHEET),),
                )
            )
            return (), tuple(issues)
        worksheet = workbook[LHA_SHEET]
        header = _find_header(worksheet)
        if header is None:
            issues.append(
                DiagnosticIssue(
                    code="LHA_HEADERS_MISSING",
                    severity=IssueSeverity.ERROR,
                    message="На листе ЛХА не найдены все обязательные заголовки.",
                    source="lha",
                    locations=(sheet_location(source_file, LHA_SHEET),),
                )
            )
            return (), tuple(issues)
        header_row, headers = header
        columns = index_headers(headers, LHA_HEADER_ALIASES)
        last_column = max(columns.values()) + 1
        records: list[LhaRecord] = []
        blank_identity_streak = 0
        invalid_value_counts: defaultdict[str, int] = defaultdict(int)
        invalid_value_samples: defaultdict[str, list[str]] = defaultdict(list)
        invalid_value_locations: defaultdict[str, list[SourceLocation]] = defaultdict(list)
        missing_value_counts: defaultdict[str, int] = defaultdict(int)
        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=header_row + 1,
                max_row=worksheet.max_row,
                max_col=last_column,
                values_only=True,
            ),
            header_row + 1,
        ):
            values = tuple(row)
            identity_values = tuple(
                row_value(values, columns, name)
                for name in ("field", "well", "plast", "measured_at")
            )
            if row_is_empty(identity_values):
                blank_identity_streak += 1
                if blank_identity_streak >= LHA_EMPTY_IDENTITY_LIMIT:
                    break
                continue
            blank_identity_streak = 0

            raw_field_value = row_value(values, columns, "field")
            raw_well_value = row_value(values, columns, "well")
            raw_plast_value = row_value(values, columns, "plast")
            raw_date_value = row_value(values, columns, "measured_at")
            raw_field = as_text(raw_field_value)
            raw_well = display_value(raw_well_value) if raw_well_value is not None else None
            raw_plast = as_text(raw_plast_value)
            normalized_well = normalize_well_id(raw_well) if raw_well is not None else None
            measured_at = parse_datetime(raw_date_value)
            if measured_at is None:
                issues.append(
                    DiagnosticIssue(
                        code="LHA_INVALID_DATE",
                        severity=IssueSeverity.WARNING,
                        message=(
                            f"Некорректная дата/время ЛХА: {raw_date_value!r}; "
                            "точка не используется."
                        ),
                        source="lha",
                        coordinate=coordinate(row_number, columns["measured_at"]),
                        well_id=normalized_well.base_well_id if normalized_well else None,
                        locations=(
                            cell_location(
                                source_file,
                                LHA_SHEET,
                                row_number,
                                columns["measured_at"],
                            ),
                        ),
                    )
                )
            if normalized_well is None or normalized_well.base_well_id is None:
                issues.append(
                    DiagnosticIssue(
                        code="LHA_INVALID_WELL",
                        severity=IssueSeverity.WARNING,
                        message="В строке ЛХА отсутствует или не распознан номер скважины.",
                        source="lha",
                        coordinate=coordinate(row_number, columns["well"]),
                        locations=(
                            cell_location(source_file, LHA_SHEET, row_number, columns["well"]),
                        ),
                    )
                )
            normalized_plast = normalize_plast(raw_plast)
            if normalized_plast is None:
                issues.append(
                    DiagnosticIssue(
                        code="LHA_INVALID_PLAST",
                        severity=IssueSeverity.WARNING,
                        message=(
                            "В строке ЛХА отсутствует обозначение пласта; точная связь невозможна."
                        ),
                        source="lha",
                        coordinate=coordinate(row_number, columns["plast"]),
                        well_id=normalized_well.base_well_id if normalized_well else None,
                        locations=(
                            cell_location(source_file, LHA_SHEET, row_number, columns["plast"]),
                        ),
                    )
                )
            cl_raw = row_value(values, columns, "cl")
            vgf_raw = row_value(values, columns, "vgf")
            cl_value = parse_number(cl_raw)
            vgf_value = parse_number(vgf_raw)
            for semantic_name, raw_value, parsed_value in (
                ("cl", cl_raw, cl_value),
                ("vgf", vgf_raw, vgf_value),
            ):
                if raw_value is not None and parsed_value is None:
                    if is_missing_measurement(raw_value):
                        missing_value_counts[semantic_name] += 1
                    else:
                        invalid_value_counts[semantic_name] += 1
                        invalid_value_locations[semantic_name].append(
                            cell_location(
                                source_file,
                                LHA_SHEET,
                                row_number,
                                columns[semantic_name],
                            )
                        )
                        if len(invalid_value_samples[semantic_name]) < 3:
                            invalid_value_samples[semantic_name].append(repr(raw_value))
            records.append(
                LhaRecord(
                    raw_field_name=raw_field,
                    normalized_field=normalize_field_name(raw_field)
                    if raw_field is not None
                    else None,
                    raw_well_name=raw_well,
                    base_well_id=normalized_well.base_well_id if normalized_well else None,
                    raw_plast=raw_plast,
                    normalized_plast=normalized_plast,
                    measured_at=measured_at,
                    cl_value=cl_value,
                    vgf_value=vgf_value,
                    source_row=row_number,
                    source_file=source_file,
                    source_sheet=LHA_SHEET,
                    field_cell=coordinate(row_number, columns["field"]),
                    well_cell=coordinate(row_number, columns["well"]),
                    plast_cell=coordinate(row_number, columns["plast"]),
                    date_cell=coordinate(row_number, columns["measured_at"]),
                    cl_cell=coordinate(row_number, columns["cl"]),
                    vgf_cell=coordinate(row_number, columns["vgf"]),
                    target_well_id=normalized_well.target_well_id if normalized_well else None,
                    raw_cl_value=as_cell_value(cl_raw),
                    raw_vgf_value=as_cell_value(vgf_raw),
                )
            )
        for semantic_name, count in missing_value_counts.items():
            issues.append(
                DiagnosticIssue(
                    code="LHA_MISSING_VALUE_SUMMARY",
                    severity=IssueSeverity.INFO,
                    message=(
                        f"В ЛХА для {semantic_name.upper()} пропущено {count} пустых, "
                        "placeholder- или Excel error-значений; они считаются отсутствующими."
                    ),
                    source="lha",
                )
            )
        for semantic_name, count in invalid_value_counts.items():
            samples = ", ".join(invalid_value_samples[semantic_name])
            issues.append(
                DiagnosticIssue(
                    code="LHA_INVALID_VALUE",
                    severity=IssueSeverity.INFO,
                    message=(
                        f"В ЛХА для {semantic_name.upper()} найдено {count} содержательных "
                        f"нечисловых значений; они прочитаны как пропуски. Примеры: {samples}."
                    ),
                    source="lha",
                )
            )
        return tuple(records), tuple(issues)
    finally:
        workbook.close()
