"""Input-path and workbook-schema validation performed before ETL."""

from __future__ import annotations

import os
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from gdis_loader.domain.models import (
    TARGET_FIELDS,
    DiagnosticIssue,
    IssueSeverity,
    SourceLocation,
    SourcePaths,
)
from gdis_loader.matching.normalize import normalize_field_name
from gdis_loader.progress import ProgressCallback, ProgressUpdate
from gdis_loader.readers.common import normalized_header, openpyxl_read_warnings
from gdis_loader.readers.daily import DAILY_SHEET, _find_daily_header
from gdis_loader.readers.fund import DATE_SHEET_RE, read_fund
from gdis_loader.readers.gdi import GDI_HEADER_ALIASES, GDI_SHEET, read_gdi
from gdis_loader.readers.gdi import _find_header as find_gdi_header
from gdis_loader.readers.lha import LHA_HEADER_ALIASES, LHA_SHEET, read_lha
from gdis_loader.readers.lha import _find_header as find_lha_header
from gdis_loader.readers.research_status import (
    STATUS_SHEETS,
    find_research_status_schema,
    read_research_status,
)
from gdis_loader.readers.well_history import (
    FIELD_SHEETS,
    find_well_history_schema,
    read_well_history,
)


@dataclass(frozen=True, slots=True)
class ValidationResult:
    """All validation findings, not just the first failure."""

    issues: tuple[DiagnosticIssue, ...]
    available_fields: tuple[str, ...] = ()
    field_coverage: dict[str, dict[str, int | bool | str | None]] = field(default_factory=dict)

    @property
    def errors(self) -> tuple[DiagnosticIssue, ...]:
        return tuple(issue for issue in self.issues if issue.severity is IssueSeverity.ERROR)

    @property
    def warnings(self) -> tuple[DiagnosticIssue, ...]:
        return tuple(issue for issue in self.issues if issue.severity is IssueSeverity.WARNING)

    @property
    def valid(self) -> bool:
        return not self.errors


def _path_issue(path: Path, code: str, message: str) -> DiagnosticIssue:
    return DiagnosticIssue(
        code=code,
        severity=IssueSeverity.ERROR,
        message=message,
        source=str(path),
        locations=(SourceLocation(file=str(path.resolve())),),
    )


def _check_path(path: Path, expected_suffixes: set[str], label: str) -> list[DiagnosticIssue]:
    issues: list[DiagnosticIssue] = []
    if not str(path):
        issues.append(_path_issue(path, "PATH_EMPTY", f"Не задан путь: {label}."))
        return issues
    if not path.exists():
        issues.append(_path_issue(path, "PATH_NOT_FOUND", f"Файл {label} не существует: {path}."))
        return issues
    if not path.is_file():
        issues.append(
            _path_issue(path, "PATH_NOT_FILE", f"Путь {label} не является файлом: {path}.")
        )
    if path.suffix.casefold() not in expected_suffixes:
        suffixes = ", ".join(sorted(expected_suffixes))
        issues.append(
            _path_issue(
                path, "PATH_EXTENSION", f"Файл {label} должен иметь расширение {suffixes}: {path}."
            )
        )
    try:
        with path.open("rb"):
            pass
    except OSError as error:
        issues.append(
            _path_issue(path, "PATH_UNREADABLE", f"Файл {label} недоступен для чтения: {error}.")
        )
    return issues


def _load_schema(
    path: Path,
    label: str,
    checker: Callable[[Any, Path, list[DiagnosticIssue]], None],
    *,
    read_only: bool = True,
) -> list[DiagnosticIssue]:
    issues: list[DiagnosticIssue] = []
    try:
        with openpyxl_read_warnings() as caught:
            workbook = load_workbook(path, read_only=read_only, data_only=False, keep_links=False)
        issues.extend(
            [
                DiagnosticIssue(
                    code="OPENPYXL_WARNING",
                    severity=IssueSeverity.WARNING,
                    message=str(warning.message),
                    source=label,
                    locations=(SourceLocation(file=str(path.resolve())),),
                )
                for warning in caught
            ]
        )
        try:
            checker(workbook, path, issues)
        finally:
            workbook.close()
    except (OSError, InvalidFileException, KeyError, TypeError, ValueError) as error:
        issues.append(
            DiagnosticIssue(
                code="WORKBOOK_OPEN_FAILED",
                severity=IssueSeverity.ERROR,
                message=f"Не удалось открыть {label}: {error}.",
                source=str(path),
                locations=(SourceLocation(file=str(path.resolve())),),
            )
        )
    return issues


def _check_calculator(workbook: Any, path: Path, issues: list[DiagnosticIssue]) -> None:
    sheetnames = workbook.sheetnames  # type: ignore[attr-defined]
    if "Данные" not in sheetnames:
        issues.append(
            DiagnosticIssue(
                code="CALCULATOR_DATA_SHEET_MISSING",
                severity=IssueSeverity.ERROR,
                message="В калькуляторе отсутствует лист 'Данные'.",
                source="calculator",
                locations=(SourceLocation(file=str(path.resolve()), sheet="Данные"),),
            )
        )
        return
    worksheet = workbook["Данные"]  # type: ignore[index]
    headers = tuple(cell.value for cell in worksheet[6])
    header_map = {
        normalized_header(value): index for index, value in enumerate(headers) if value is not None
    }
    missing = [name for name in TARGET_FIELDS if normalized_header(name) not in header_map]
    if missing:
        issues.append(
            DiagnosticIssue(
                code="CALCULATOR_HEADERS_MISSING",
                severity=IssueSeverity.ERROR,
                message=f"В строке 6 калькулятора отсутствуют заголовки: {', '.join(missing)}.",
                source="calculator",
                locations=(SourceLocation(file=str(path.resolve()), sheet="Данные"),),
            )
        )
    if "tblData" not in worksheet.tables:
        issues.append(
            DiagnosticIssue(
                code="CALCULATOR_TABLE_MISSING",
                severity=IssueSeverity.ERROR,
                message="На листе 'Данные' отсутствует таблица tblData.",
                source="calculator",
                locations=(SourceLocation(file=str(path.resolve()), sheet="Данные"),),
            )
        )


def _check_fund(workbook: Any, path: Path, issues: list[DiagnosticIssue]) -> None:
    dated = [
        str(worksheet.title).strip()
        for worksheet in workbook.worksheets  # type: ignore[attr-defined]
        if DATE_SHEET_RE.fullmatch(str(worksheet.title).strip())
        and not str(worksheet.title).strip().upper().startswith("RSPN")
    ]
    if not dated:
        issues.append(
            DiagnosticIssue(
                code="FUND_DATE_SHEET_NOT_FOUND",
                severity=IssueSeverity.ERROR,
                message="В фонде отсутствует подходящий датированный лист.",
                source="fund",
                locations=(SourceLocation(file=str(path.resolve())),),
            )
        )


def _check_gdi(workbook: Any, path: Path, issues: list[DiagnosticIssue]) -> None:
    if GDI_SHEET not in workbook.sheetnames:  # type: ignore[attr-defined]
        issues.append(
            DiagnosticIssue(
                code="GDI_SHEET_MISSING",
                severity=IssueSeverity.ERROR,
                message=f"В книге ГДИ отсутствует лист {GDI_SHEET!r}.",
                source="gdi",
                locations=(SourceLocation(file=str(path.resolve()), sheet=GDI_SHEET),),
            )
        )
        return
    if find_gdi_header(workbook[GDI_SHEET]) is None:  # type: ignore[index]
        aliases = ", ".join(GDI_HEADER_ALIASES)
        issues.append(
            DiagnosticIssue(
                code="GDI_HEADERS_MISSING",
                severity=IssueSeverity.ERROR,
                message=f"На листе ГДИ не найдены обязательные заголовки: {aliases}.",
                source="gdi",
                locations=(SourceLocation(file=str(path.resolve()), sheet=GDI_SHEET),),
            )
        )


def _check_lha(workbook: Any, path: Path, issues: list[DiagnosticIssue]) -> None:
    if LHA_SHEET not in workbook.sheetnames:  # type: ignore[attr-defined]
        issues.append(
            DiagnosticIssue(
                code="LHA_SHEET_MISSING",
                severity=IssueSeverity.ERROR,
                message=f"В книге ЛХА отсутствует лист {LHA_SHEET!r}.",
                source="lha",
                locations=(SourceLocation(file=str(path.resolve()), sheet=LHA_SHEET),),
            )
        )
        return
    if find_lha_header(workbook[LHA_SHEET]) is None:  # type: ignore[index]
        aliases = ", ".join(LHA_HEADER_ALIASES)
        issues.append(
            DiagnosticIssue(
                code="LHA_HEADERS_MISSING",
                severity=IssueSeverity.ERROR,
                message=f"На листе ЛХА не найдены обязательные заголовки: {aliases}.",
                source="lha",
                locations=(SourceLocation(file=str(path.resolve()), sheet=LHA_SHEET),),
            )
        )


def _check_daily(workbook: Any, path: Path, issues: list[DiagnosticIssue]) -> None:
    if DAILY_SHEET not in workbook.sheetnames:  # type: ignore[attr-defined]
        issues.append(
            DiagnosticIssue(
                code="DAILY_SHEET_MISSING",
                severity=IssueSeverity.ERROR,
                message=f"В суточном рапорте отсутствует лист {DAILY_SHEET!r}.",
                source="daily_report",
                locations=(SourceLocation(file=str(path.resolve()), sheet=DAILY_SHEET),),
            )
        )
        return
    if _find_daily_header(workbook[DAILY_SHEET]) is None:  # type: ignore[index]
        issues.append(
            DiagnosticIssue(
                code="DAILY_HEADERS_MISSING",
                severity=IssueSeverity.ERROR,
                message="На листе суточного рапорта не найдены заголовки ГСК и СКВ.",
                source="daily_report",
                locations=(SourceLocation(file=str(path.resolve()), sheet=DAILY_SHEET),),
            )
        )


def _check_well_history(workbook: Any, path: Path, issues: list[DiagnosticIssue]) -> None:
    present = [sheet for sheet in FIELD_SHEETS if sheet in workbook.sheetnames]  # type: ignore[attr-defined]
    if not present:
        issues.append(
            DiagnosticIssue(
                code="WELL_HISTORY_SHEET_MISSING",
                severity=IssueSeverity.ERROR,
                message="В базе истории скважин отсутствуют поддерживаемые листы ЛУ.",
                source="well_history",
                locations=(SourceLocation(file=str(path.resolve())),),
            )
        )
        return
    valid_sheets = [
        sheet
        for sheet in present
        if find_well_history_schema(workbook[sheet]) is not None  # type: ignore[index]
    ]
    if not valid_sheets:
        issues.append(
            DiagnosticIssue(
                code="WELL_HISTORY_SCHEMA_INVALID",
                severity=IssueSeverity.ERROR,
                message="Не распознаны двухстрочные заголовки и блоки ГТМ на листах истории.",
                source="well_history",
                locations=tuple(
                    SourceLocation(file=str(path.resolve()), sheet=sheet) for sheet in present
                ),
            )
        )


def _check_research_status(workbook: Any, path: Path, issues: list[DiagnosticIssue]) -> None:
    present = [
        sheet
        for sheet in STATUS_SHEETS.values()
        if sheet in workbook.sheetnames  # type: ignore[attr-defined]
    ]
    if not present:
        issues.append(
            DiagnosticIssue(
                code="GDM_STATUS_SHEET_MISSING",
                severity=IssueSeverity.ERROR,
                message="В книге статуса исследований отсутствуют поддерживаемые листы ЛУ.",
                source="research_status",
                locations=(SourceLocation(file=str(path.resolve())),),
            )
        )
        return
    valid_sheets = [
        sheet
        for sheet in present
        if find_research_status_schema(workbook[sheet]) is not None  # type: ignore[index]
    ]
    if not valid_sheets:
        issues.append(
            DiagnosticIssue(
                code="GDM_STATUS_HEADERS_MISSING",
                severity=IssueSeverity.ERROR,
                message="Не распознаны многострочные заголовки таблиц статуса исследований.",
                source="research_status",
                locations=tuple(
                    SourceLocation(file=str(path.resolve()), sheet=sheet) for sheet in present
                ),
            )
        )


def _canonical_field(raw_field: str) -> str:
    known_names = {normalize_field_name(name): name for name in FIELD_SHEETS.values()}
    return known_names.get(normalize_field_name(raw_field), " ".join(raw_field.split()))


def _field_mapping(raw_field: str) -> tuple[str | None, str | None]:
    key = normalize_field_name(raw_field)
    for code, name in FIELD_SHEETS.items():
        if normalize_field_name(name) == key:
            return code, STATUS_SHEETS.get(code)
    return None, None


def _discover_fields(
    paths: SourcePaths,
    issues: list[DiagnosticIssue],
    progress: ProgressCallback | None = None,
    progress_range: tuple[int, int] = (35, 80),
) -> tuple[tuple[str, ...], dict[str, dict[str, int | bool | str | None]]]:
    """Discover target fields from the fund and measure source coverage."""

    start, end = progress_range
    span = end - start
    if progress is not None:
        progress(ProgressUpdate(start, "Чтение фонда для обнаружения месторождений"))
    fund = read_fund(paths.fund)
    issues.extend(fund.issues)
    available_fields: list[str] = []
    fund_counts: dict[str, int] = {}
    for well in fund.wells:
        canonical = _canonical_field(well.field_name)
        if canonical not in available_fields:
            available_fields.append(canonical)
        fund_counts[canonical] = fund_counts.get(canonical, 0) + 1

    if progress is not None:
        progress(ProgressUpdate(start + span // 4, "Чтение данных ГДИ и ЛХА"))
    gdi_records, gdi_issues = read_gdi(paths.gdi)
    lha_records, lha_issues = read_lha(paths.lha)
    issues.extend((*gdi_issues, *lha_issues))
    history_workbook = load_workbook(
        paths.well_history, read_only=True, data_only=True, keep_links=False
    )
    status_workbook = load_workbook(
        paths.research_status, read_only=True, data_only=True, keep_links=False
    )
    if progress is not None:
        progress(ProgressUpdate(start + span // 2, "Проверка листов истории и статуса"))
    coverage: dict[str, dict[str, int | bool | str | None]] = {}
    try:
        for field_index, field_name in enumerate(available_fields):
            normalized = normalize_field_name(field_name)
            code, status_sheet = _field_mapping(field_name)
            gdi_count = sum(record.normalized_field == normalized for record in gdi_records)
            lha_count = sum(record.normalized_field == normalized for record in lha_records)
            history_sheet = code
            history_present = (
                history_sheet is not None and history_sheet in history_workbook.sheetnames
            )
            status_present = status_sheet is not None and status_sheet in status_workbook.sheetnames
            history_rows: int | None = None
            status_rows: int | None = None
            if code is None or history_sheet is None:
                issues.append(
                    DiagnosticIssue(
                        code="FIELD_SOURCE_SHEET_MAPPING_MISSING",
                        severity=IssueSeverity.WARNING,
                        message=(
                            f"Для месторождения {field_name!r} не настроено "
                            "сопоставление листа базы истории."
                        ),
                        source="well_history",
                        locations=(SourceLocation(file=str(paths.well_history.resolve())),),
                    )
                )
            elif not history_present:
                issues.append(
                    DiagnosticIssue(
                        code="FIELD_SOURCE_SHEET_MISSING",
                        severity=IssueSeverity.WARNING,
                        message=f"В базе истории нет листа {history_sheet!r} для {field_name!r}.",
                        source="well_history",
                        locations=(
                            SourceLocation(
                                file=str(paths.well_history.resolve()), sheet=history_sheet
                            ),
                        ),
                    )
                )
            else:
                history_schema = find_well_history_schema(history_workbook[history_sheet])
                if history_schema is None:
                    issues.append(
                        DiagnosticIssue(
                            code="FIELD_SOURCE_SCHEMA_INVALID",
                            severity=IssueSeverity.WARNING,
                            message=(
                                f"На листе {history_sheet!r} не распознана схема "
                                f"для {field_name!r}."
                            ),
                            source="well_history",
                            locations=(
                                SourceLocation(
                                    file=str(paths.well_history.resolve()), sheet=history_sheet
                                ),
                            ),
                        )
                    )
                else:
                    history_result = read_well_history(paths.well_history, field_name)
                    history_rows = len(history_result.histories)
                    issues.extend(history_result.issues)

            if status_sheet is None:
                issues.append(
                    DiagnosticIssue(
                        code="FIELD_SOURCE_SHEET_MAPPING_MISSING",
                        severity=IssueSeverity.WARNING,
                        message=(
                            f"Для месторождения {field_name!r} не настроено "
                            "сопоставление листа статуса исследований."
                        ),
                        source="research_status",
                        locations=(SourceLocation(file=str(paths.research_status.resolve())),),
                    )
                )
            elif not status_present:
                issues.append(
                    DiagnosticIssue(
                        code="FIELD_SOURCE_SHEET_MISSING",
                        severity=IssueSeverity.WARNING,
                        message=(
                            f"В статусе исследований нет листа {status_sheet!r} для {field_name!r}."
                        ),
                        source="research_status",
                        locations=(
                            SourceLocation(
                                file=str(paths.research_status.resolve()), sheet=status_sheet
                            ),
                        ),
                    )
                )
            else:
                status_schema = find_research_status_schema(status_workbook[status_sheet])
                if status_schema is None:
                    issues.append(
                        DiagnosticIssue(
                            code="FIELD_SOURCE_SCHEMA_INVALID",
                            severity=IssueSeverity.WARNING,
                            message=(
                                f"На листе {status_sheet!r} не распознана схема для {field_name!r}."
                            ),
                            source="research_status",
                            locations=(
                                SourceLocation(
                                    file=str(paths.research_status.resolve()), sheet=status_sheet
                                ),
                            ),
                        )
                    )
                elif code is not None:
                    status_records, status_issues = read_research_status(
                        paths.research_status, code
                    )
                    status_rows = len(status_records)
                    issues.extend(status_issues)

            coverage[field_name] = {
                "fund_rows": fund_counts[field_name],
                "gdi_rows": gdi_count,
                "lha_rows": lha_count,
                "well_history_sheet_present": history_present,
                "well_history_rows": history_rows,
                "research_status_sheet_present": status_present,
                "research_status_rows": status_rows,
            }
            if progress is not None:
                progress_percent = (
                    start
                    + span // 2
                    + (span // 2) * (field_index + 1) // max(len(available_fields), 1)
                )
                progress(ProgressUpdate(progress_percent, f"Проверка покрытия для {field_name}"))
    finally:
        history_workbook.close()
        status_workbook.close()
    if progress is not None:
        progress(ProgressUpdate(end, "Обнаружение месторождений завершено"))
    return tuple(available_fields), coverage


def validate_inputs(
    paths: SourcePaths,
    progress: ProgressCallback | None = None,
    progress_range: tuple[int, int] = (10, 80),
) -> ValidationResult:
    """Validate all paths and workbook schemas in one pass of findings."""

    issues: list[DiagnosticIssue] = []
    progress_start, progress_end = progress_range
    progress_span = progress_end - progress_start
    schema_end = progress_start + progress_span * 2 // 5
    if progress is not None:
        progress(ProgressUpdate(progress_start, "Проверка путей источников"))
    source_specs = (
        (paths.calculator, {".xlsx"}, "калькулятор", _check_calculator),
        (paths.fund, {".xlsx"}, "фонд скважин", _check_fund),
        (paths.gdi, {".xlsx"}, "сводная таблица ГДИ", _check_gdi),
        (paths.lha, {".xlsm"}, "база ЛХА", _check_lha),
        (paths.daily_report, {".xlsm"}, "суточный рапорт", _check_daily),
        (paths.well_history, {".xlsx"}, "база истории скважин", _check_well_history),
        (paths.research_status, {".xlsx"}, "статус исследований", _check_research_status),
    )
    for index, (path, suffixes, label, checker) in enumerate(source_specs):
        path_issues = _check_path(path, suffixes, label)
        issues.extend(path_issues)
        if not any(issue.severity is IssueSeverity.ERROR for issue in path_issues):
            issues.extend(_load_schema(path, label, checker, read_only=label != "калькулятор"))
        if progress is not None:
            percent = progress_start + (schema_end - progress_start) * (index + 1) // len(
                source_specs
            )
            progress(ProgressUpdate(percent, f"Проверка структуры: {label}"))

    available_fields: tuple[str, ...] = ()
    field_coverage: dict[str, dict[str, int | bool | str | None]] = {}
    if not any(issue.severity is IssueSeverity.ERROR for issue in issues):
        available_fields, field_coverage = _discover_fields(
            paths,
            issues,
            progress,
            (schema_end, progress_end),
        )
    elif progress is not None:
        progress(ProgressUpdate(schema_end, "Проверка завершена с ошибками структуры"))
    return ValidationResult(tuple(issues), available_fields, field_coverage)


def validate_output_path(output: Path, paths: SourcePaths) -> tuple[DiagnosticIssue, ...]:
    """Validate a result destination immediately before update only."""

    issues: list[DiagnosticIssue] = []
    if output.suffix.casefold() != ".xlsx":
        issues.append(
            _path_issue(
                output, "OUTPUT_EXTENSION", "Путь результата должен иметь расширение .xlsx."
            )
        )
    input_paths = (
        (paths.calculator, "калькулятором"),
        (paths.fund, "фондом"),
        (paths.gdi, "сводной ГДИ"),
        (paths.lha, "базой ЛХА"),
        (paths.daily_report, "суточным рапортом"),
        (paths.well_history, "базой истории скважин"),
        (paths.research_status, "статусом исследований"),
    )
    collision = next(
        (
            label
            for source_path, label in input_paths
            if output.resolve(strict=False) == source_path.resolve(strict=False)
        ),
        None,
    )
    if collision is not None:
        issues.append(
            _path_issue(
                output,
                "OUTPUT_EQUALS_INPUT",
                f"Путь результата совпадает с исходным {collision}; "
                "исходный файл нельзя перезаписать.",
            )
        )
    if not output.parent.is_dir():
        issues.append(
            _path_issue(
                output.parent,
                "OUTPUT_DIRECTORY_MISSING",
                "Каталог результата не существует.",
            )
        )
    elif not os.access(output.parent, os.W_OK):
        issues.append(
            _path_issue(
                output.parent,
                "OUTPUT_DIRECTORY_NOT_WRITABLE",
                "Нет прав на запись в каталог результата.",
            )
        )
    if output.exists() and output.is_dir():
        issues.append(
            _path_issue(output, "OUTPUT_IS_DIRECTORY", "Путь результата указывает на каталог.")
        )
    elif output.exists() and not os.access(output, os.W_OK):
        issues.append(
            _path_issue(output, "OUTPUT_NOT_WRITABLE", "Нет прав на перезапись файла результата.")
        )
    return tuple(issues)
