"""Minimal OOXML patch writer for the calculator's `Данные` sheet."""

from __future__ import annotations

import math
import posixpath
import re
import tempfile
from collections.abc import Iterable, Mapping
from copy import deepcopy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import (
    column_index_from_string,
    get_column_letter,
    range_boundaries,
)
from openpyxl.utils.datetime import to_excel

from gdis_loader.domain.models import TARGET_FIELDS, TargetRow
from gdis_loader.readers.common import normalized_header, openpyxl_read_warnings

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
X14AC_NS = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"


def Q(namespace: str, name: str) -> str:
    return f"{{{namespace}}}{name}"


ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", REL_NS)
ET.register_namespace("mc", MC_NS)
ET.register_namespace("x14ac", X14AC_NS)

_ROOT_START_RE = re.compile(rb"<(?P<prefix>[A-Za-z_][\w.-]*:)?(?P<name>\w+)\b[^>]*>")
_XMLNS_RE = re.compile(rb"\s+xmlns(?::([A-Za-z_][\w.-]*))?\s*=\s*([\"'])(.*?)\2")
_CELL_REF_RE = re.compile(r"^([A-Za-z]+)(\d+)$")


@dataclass(frozen=True, slots=True)
class TargetWriteResult:
    """Writer outcome and exact package entries intentionally changed."""

    output_path: Path
    rows_written: int
    warnings: tuple[str, ...] = ()
    changed_entries: tuple[str, ...] = ()


class TargetWorkbookError(ValueError):
    """Raised when the calculator package is not safe to patch."""


@dataclass(frozen=True, slots=True)
class _TableBounds:
    min_col: int
    min_row: int
    max_col: int
    max_row: int

    @property
    def data_start_row(self) -> int:
        return self.min_row + 1


@dataclass(frozen=True, slots=True)
class _TargetSchema:
    header_positions: dict[str, int]
    table_bounds: _TableBounds


def _relationship_target(base_path: str, target: str) -> str:
    if target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_path), target))


def _resolve_package_targets(archive: ZipFile) -> tuple[str, str]:
    workbook_path = "xl/workbook.xml"
    workbook_root = ET.fromstring(archive.read(workbook_path))
    workbook_rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationships = {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in workbook_rels.findall(Q(PACKAGE_REL_NS, "Relationship"))
    }
    data_sheet_path: str | None = None
    for sheet in workbook_root.findall(Q(MAIN_NS, "sheets") + "/" + Q(MAIN_NS, "sheet")):
        if sheet.attrib.get("name") == "Данные":
            relation_id = sheet.attrib.get(Q(REL_NS, "id"))
            if relation_id in relationships:
                data_sheet_path = _relationship_target(
                    "xl/workbook.xml", relationships[relation_id]
                )
            break
    if data_sheet_path is None:
        raise TargetWorkbookError("Не удалось определить XML листа 'Данные'.")
    rels_path = posixpath.join(
        posixpath.dirname(data_sheet_path), "_rels", f"{posixpath.basename(data_sheet_path)}.rels"
    )
    if rels_path not in archive.namelist():
        raise TargetWorkbookError("У листа 'Данные' отсутствует relationships-файл.")
    sheet_rels = ET.fromstring(archive.read(rels_path))
    table_paths = [
        _relationship_target(data_sheet_path, relationship.attrib["Target"])
        for relationship in sheet_rels.findall(Q(PACKAGE_REL_NS, "Relationship"))
        if relationship.attrib.get("Type", "").endswith("/table")
    ]
    for table_path in table_paths:
        root = ET.fromstring(archive.read(table_path))
        if root.attrib.get("name") == "tblData":
            return data_sheet_path, table_path
    raise TargetWorkbookError("Для листа 'Данные' не найдено связанное определение tblData.")


def _bounds_from_ref(reference: str) -> _TableBounds:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(reference)
    except ValueError as error:
        raise TargetWorkbookError(
            f"Некорректный диапазон таблицы tblData: {reference!r}."
        ) from error
    return _TableBounds(min_col, min_row, max_col, max_row)


def _table_bounds_from_xml(data: bytes) -> _TableBounds:
    root = ET.fromstring(data)
    if root.attrib.get("name") != "tblData":
        raise TargetWorkbookError("Связанная таблица не называется tblData.")
    reference = root.attrib.get("ref")
    if not reference:
        raise TargetWorkbookError("У таблицы tblData отсутствует диапазон ref.")
    return _bounds_from_ref(reference)


def _cell_value_element(cell: ET.Element) -> None:
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)


def _set_cell_value(cell: ET.Element, value: object) -> None:
    _cell_value_element(cell)
    if value is None:
        return
    if isinstance(value, str):
        cell.set("t", "inlineStr")
        inline = ET.SubElement(cell, Q(MAIN_NS, "is"))
        text = ET.SubElement(inline, Q(MAIN_NS, "t"))
        if value[:1].isspace() or value[-1:].isspace():
            text.set(Q(XML_NS, "space"), "preserve")
        text.text = value
        return
    if isinstance(value, bool):
        cell.set("t", "b")
        ET.SubElement(cell, Q(MAIN_NS, "v")).text = "1" if value else "0"
        return
    if isinstance(value, (datetime, date)):
        value = to_excel(value)
    if isinstance(value, (int, float)):
        numeric = float(value)
        if not math.isfinite(numeric):
            raise TargetWorkbookError(f"Нельзя записать нечисловое значение {value!r}.")
        ET.SubElement(cell, Q(MAIN_NS, "v")).text = str(value)
        return
    raise TargetWorkbookError(
        f"Неподдерживаемый тип значения целевой строки: {type(value).__name__}."
    )


def _cell_column(cell: ET.Element) -> int | None:
    match = _CELL_REF_RE.fullmatch(cell.attrib.get("r", ""))
    if match is None:
        return None
    return column_index_from_string(match.group(1))


def _clone_data_row(
    template: ET.Element,
    row_number: int,
    values: Mapping[int, object],
    bounds: _TableBounds,
    preserve_row: ET.Element | None = None,
) -> ET.Element:
    row = deepcopy(preserve_row if preserve_row is not None else template)
    row.set("r", str(row_number))
    row.set("spans", f"{bounds.min_col}:{bounds.max_col}")
    for child in list(row):
        if child.tag == Q(MAIN_NS, "c") and (
            (column := _cell_column(child)) is not None
            and bounds.min_col <= column <= bounds.max_col
        ):
            row.remove(child)

    template_cells = {
        column: cell
        for cell in template.findall(Q(MAIN_NS, "c"))
        if (column := _cell_column(cell)) is not None
    }
    for column in range(bounds.min_col, bounds.max_col + 1):
        cell = deepcopy(template_cells.get(column, ET.Element(Q(MAIN_NS, "c"))))
        cell.set("r", f"{get_column_letter(column)}{row_number}")
        _set_cell_value(cell, values.get(column))
        row.append(cell)
    return row


def _clear_table_cells(row: ET.Element, bounds: _TableBounds) -> ET.Element | None:
    for child in list(row):
        if child.tag == Q(MAIN_NS, "c"):
            column = _cell_column(child)
            if column is not None and bounds.min_col <= column <= bounds.max_col:
                row.remove(child)
    return row if list(row) else None


def _expand_dimension(root: ET.Element, maximum_row: int, maximum_column: int) -> None:
    dimension = root.find(Q(MAIN_NS, "dimension"))
    if dimension is None:
        dimension = ET.Element(Q(MAIN_NS, "dimension"), {"ref": "A1:A1"})
        root.insert(0, dimension)
    current_ref = dimension.attrib.get("ref", "A1:A1")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(current_ref)
    except ValueError as error:
        raise TargetWorkbookError(f"Некорректный dimension листа: {current_ref!r}.") from error
    dimension.set(
        "ref",
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max(max_col, maximum_column))}{max(max_row, maximum_row)}",
    )


def _expand_matching_conditional_formatting(
    root: ET.Element, bounds: _TableBounds, new_end_row: int
) -> None:
    if bounds.max_row < bounds.data_start_row or new_end_row < bounds.data_start_row:
        return
    old_reference = f"A{bounds.data_start_row}:A{bounds.max_row}"
    new_reference = f"A{bounds.data_start_row}:A{new_end_row}"
    if old_reference == new_reference:
        return
    for conditional_formatting in root.findall(Q(MAIN_NS, "conditionalFormatting")):
        if conditional_formatting.attrib.get("sqref", "").strip() == old_reference:
            conditional_formatting.set("sqref", new_reference)


def _patch_sheet_xml(
    data: bytes,
    rows: tuple[TargetRow, ...],
    header_positions: Mapping[str, int],
    bounds: _TableBounds,
) -> bytes:
    root = ET.fromstring(data)
    sheet_data = root.find(Q(MAIN_NS, "sheetData"))
    if sheet_data is None:
        raise TargetWorkbookError("В XML листа 'Данные' отсутствует sheetData.")

    existing_rows = {
        int(row.attrib["r"]): row
        for row in sheet_data.findall(Q(MAIN_NS, "row"))
        if row.attrib.get("r", "").isdigit()
    }
    new_end_row = bounds.min_row + len(rows)
    expansion_conflicts = sorted(
        row_number for row_number in existing_rows if bounds.max_row < row_number <= new_end_row
    )
    if expansion_conflicts:
        joined = ", ".join(str(row_number) for row_number in expansion_conflicts[:10])
        suffix = "…" if len(expansion_conflicts) > 10 else ""
        raise TargetWorkbookError(
            "TARGET_TABLE_EXPANSION_CONFLICT: расширение tblData пересекает "
            f"нетабличные строки {joined}{suffix}; исходный результат не заменен."
        )

    table_rows = [
        row_number
        for row_number in existing_rows
        if bounds.data_start_row <= row_number <= bounds.max_row
    ]
    template = existing_rows.get(bounds.max_row)
    if template is None and table_rows:
        template = existing_rows[table_rows[0]]
    if template is None:
        template = ET.Element(
            Q(MAIN_NS, "row"),
            {"spans": f"{bounds.min_col}:{bounds.max_col}"},
        )

    row_map = dict(existing_rows)
    for row_number in table_rows:
        row_map.pop(row_number, None)

    for row_number, target_row in enumerate(rows, bounds.data_start_row):
        mapping = target_row.as_mapping()
        values = {header_positions[field]: mapping[field] for field in TARGET_FIELDS}
        row_map[row_number] = _clone_data_row(
            template,
            row_number,
            values,
            bounds,
            preserve_row=existing_rows.get(row_number),
        )

    for row_number in table_rows[len(rows) :]:
        cleared = _clear_table_cells(existing_rows[row_number], bounds)
        if cleared is not None:
            row_map[row_number] = cleared

    for child in list(sheet_data):
        if child.tag == Q(MAIN_NS, "row"):
            sheet_data.remove(child)
    for row_number in sorted(row_map):
        sheet_data.append(row_map[row_number])

    maximum_existing_row = max(row_map, default=bounds.min_row)
    _expand_dimension(root, max(maximum_existing_row, new_end_row), bounds.max_col)
    _expand_matching_conditional_formatting(root, bounds, new_end_row)
    serialized = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    return _restore_root_namespace_declarations(data, serialized, "worksheet")


def _patch_table_xml(data: bytes, bounds: _TableBounds, rows_written: int) -> bytes:
    root = ET.fromstring(data)
    new_end_row = bounds.min_row + rows_written
    reference = (
        f"{get_column_letter(bounds.min_col)}{bounds.min_row}:"
        f"{get_column_letter(bounds.max_col)}{max(new_end_row, bounds.min_row)}"
    )
    root.set("ref", reference)
    auto_filter = root.find(Q(MAIN_NS, "autoFilter"))
    if auto_filter is not None:
        auto_filter.set("ref", reference)
    serialized = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    return _restore_root_namespace_declarations(data, serialized, "table")


def _restore_root_namespace_declarations(data: bytes, serialized: bytes, name: str) -> bytes:
    """Restore declarations that ET drops when a prefix is only ignorable."""

    source_match = _ROOT_START_RE.search(data)
    if source_match is None or source_match.group("name").decode("ascii") != name:
        return serialized
    declarations = [
        (match.group(1).decode("utf-8") if match.group(1) else None, match.group(3).decode("utf-8"))
        for match in _XMLNS_RE.finditer(source_match.group(0))
    ]
    text = serialized.decode("utf-8")
    for prefix, uri in declarations:
        attribute = "xmlns" if prefix is None else f"xmlns:{prefix}"
        root_match = re.search(rf"<{re.escape(name)}\b[^>]*>", text)
        if root_match is None:
            return serialized
        if re.search(rf"\s{re.escape(attribute)}\s*=", root_match.group(0)):
            continue
        position = root_match.end() - 1
        text = f'{text[:position]} {attribute}="{uri}"{text[position:]}'
    return text.encode("utf-8")


def _validate_target_headers(source: Path) -> _TargetSchema:
    with openpyxl_read_warnings() as caught:
        workbook = load_workbook(source, read_only=False, data_only=False, keep_links=False)
    _ = caught
    try:
        if "Данные" not in workbook.sheetnames:
            raise TargetWorkbookError("В калькуляторе отсутствует лист 'Данные'.")
        worksheet = workbook["Данные"]
        if "tblData" not in worksheet.tables:
            raise TargetWorkbookError("На листе 'Данные' отсутствует таблица tblData.")
        bounds = _bounds_from_ref(worksheet.tables["tblData"].ref)
        headers = tuple(cell.value for cell in worksheet[bounds.min_row])
        positions = {
            normalized_header(value): index
            for index, value in enumerate(headers, 1)
            if value is not None and bounds.min_col <= index <= bounds.max_col
        }
        missing = [field for field in TARGET_FIELDS if normalized_header(field) not in positions]
        if missing:
            raise TargetWorkbookError(
                f"В калькуляторе отсутствуют заголовки: {', '.join(missing)}."
            )
        return _TargetSchema(
            header_positions={
                field: positions[normalized_header(field)] for field in TARGET_FIELDS
            },
            table_bounds=bounds,
        )
    finally:
        workbook.close()


def _verify_output(output: Path, rows: tuple[TargetRow, ...]) -> None:
    with openpyxl_read_warnings():
        workbook = load_workbook(output, read_only=True, data_only=True, keep_links=False)
    try:
        worksheet = workbook["Данные"]
        header_values = next(worksheet.iter_rows(min_row=6, max_row=6, values_only=True))
        header_map = {
            normalized_header(value): index
            for index, value in enumerate(header_values)
            if value is not None
        }
        actual_rows = worksheet.iter_rows(
            min_row=7, max_row=6 + len(rows), max_col=len(header_values), values_only=True
        )
        for row_number, (expected, actual) in enumerate(zip(rows, actual_rows, strict=True), 7):
            for name, expected_value in expected.as_mapping().items():
                index = header_map.get(normalized_header(name))
                if index is None:
                    raise TargetWorkbookError(f"После записи не найден заголовок {name!r}.")
                actual_value = actual[index]
                if expected_value is None:
                    if actual_value not in (None, ""):
                        raise TargetWorkbookError(f"Строка {row_number}, поле {name!r} не очищено.")
                elif isinstance(expected_value, (int, float)) and isinstance(
                    actual_value, (int, float)
                ):
                    if float(actual_value) != float(expected_value):
                        raise TargetWorkbookError(
                            f"Строка {row_number}, поле {name!r} записано неверно."
                        )
                elif isinstance(expected_value, date) and isinstance(actual_value, datetime):
                    if actual_value.date() != expected_value:
                        raise TargetWorkbookError(
                            f"Строка {row_number}, поле {name!r} записано неверно."
                        )
                elif actual_value != expected_value:
                    raise TargetWorkbookError(
                        f"Строка {row_number}, поле {name!r} записано неверно."
                    )
    finally:
        workbook.close()


def write_target_rows(source: Path, output: Path, rows: Iterable[TargetRow]) -> TargetWriteResult:
    """Copy the source package and patch only sheet XML plus tblData XML."""

    if source.resolve() == output.resolve():
        raise TargetWorkbookError("Исходный калькулятор нельзя перезаписывать.")
    target_rows = tuple(rows)
    schema = _validate_target_headers(source)
    warnings_out: list[str] = []
    from gdis_loader.excel.template_capabilities import detect_template_capabilities

    capabilities = detect_template_capabilities(source)
    if capabilities.exceeds_capacity(len(target_rows)):
        warnings_out.append("DATA_ROWS_EXCEED_CALCULATOR_FORMULA_CAPACITY")
    if capabilities.calculation_capacity is None:
        warnings_out.append("CALCULATOR_FORMULA_CAPACITY_UNKNOWN")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_name: str | None = None
    with ZipFile(source, "r") as archive:
        sheet_path, table_path = _resolve_package_targets(archive)
        table_data = archive.read(table_path)
        raw_bounds = _table_bounds_from_xml(table_data)
        if raw_bounds != schema.table_bounds:
            raise TargetWorkbookError(
                "Диапазон tblData в OOXML не совпадает со схемой, прочитанной из книги."
            )
        patched_sheet = _patch_sheet_xml(
            archive.read(sheet_path),
            target_rows,
            schema.header_positions,
            raw_bounds,
        )
        patched_table = _patch_table_xml(table_data, raw_bounds, len(target_rows))
        original_names = archive.namelist()
        with tempfile.NamedTemporaryFile(
            mode="wb",
            suffix=".xlsx",
            prefix=f".{output.name}.",
            dir=output.parent,
            delete=False,
        ) as temporary:
            temporary_name = temporary.name
        try:
            with ZipFile(temporary_name, "w") as destination:
                for info in archive.infolist():
                    data = (
                        patched_sheet
                        if info.filename == sheet_path
                        else patched_table
                        if info.filename == table_path
                        else archive.read(info.filename)
                    )
                    destination.writestr(info, data)
        except BaseException:
            Path(temporary_name).unlink(missing_ok=True)
            raise
    try:
        _verify_output(Path(temporary_name), target_rows)
        with ZipFile(source, "r") as before, ZipFile(temporary_name, "r") as after:
            for name in original_names:
                if name in {sheet_path, table_path}:
                    continue
                if before.read(name) != after.read(name):
                    raise TargetWorkbookError(
                        f"Нетаргетная OOXML entry неожиданно изменилась: {name}."
                    )
        Path(temporary_name).replace(output)
    except BaseException:
        Path(temporary_name).unlink(missing_ok=True)
        raise
    return TargetWriteResult(
        output_path=output,
        rows_written=len(target_rows),
        warnings=tuple(warnings_out),
        changed_entries=(sheet_path, table_path),
    )
