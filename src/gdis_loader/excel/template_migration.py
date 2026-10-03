"""Offline, SHA-pinned structural migration. Never called by loader runtime."""

from __future__ import annotations

import hashlib
import html
import re
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange

SOURCE_SHA256 = "bb046cc365749a546d1bbe141165b08a664ba427ab9493b678ddce008fc85c6a"
TEMPLATE_VERSION = 2
CAPACITY = 1000
REMOVED_SLOTS = (8, 14)
EXPECTED_CHANGED = frozenset(
    {
        "xl/styles.xml",
        "xl/worksheets/sheet2.xml",
        "xl/worksheets/sheet4.xml",
        "xl/worksheets/sheet5.xml",
        "xl/worksheets/sheet6.xml",
        "xl/tables/table3.xml",
        "xl/tables/table4.xml",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "[Content_Types].xml",
        "xl/calcChain.xml",
    }
)
CELL = re.compile(r'<c\b[^>]*\br="([A-Z]+)(\d+)"[^>]*?(?:/>|>.*?</c>)', re.DOTALL)
FORMULA = re.compile(r"<f(?:\s[^>]*)?>(.*?)</f>", re.DOTALL)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _arguments(expression: str) -> list[str]:
    """Split a formula function's arguments, respecting strings and nesting."""
    depth = 0
    quoted = False
    start = 0
    args = []
    for index, char in enumerate(expression):
        if char == '"':
            quoted = not quoted
        elif not quoted:
            if char == "(":
                depth += 1
            elif char == ")":
                if depth == 0:
                    args.append(expression[start:index])
                    return args
                depth -= 1
            elif char == "," and depth == 0:
                args.append(expression[start:index])
                start = index + 1
    raise ValueError("Unbalanced pinned formula")


def _mu(formula: str, raw: str) -> tuple[str, str]:
    prefix = f'IF(OR({raw}="#FIELD"'
    position = formula.index(prefix)
    condition, missing, _known = _arguments(formula[position + 3 :])
    assert missing.startswith("IF($") and '"Недостаточно данных"' in missing
    return formula.replace(missing, "0", 1), condition


def _formula_cell(column: str, row: int, formula: str, style: str = "") -> str:
    return f'<c r="{column}{row}"{style}><f>{html.escape(formula, quote=False)}</f></c>'


def _capacity_formula(formula: str) -> str:
    return formula.replace("$106", f"${CAPACITY + 6}").replace("$103", f"${CAPACITY + 3}")


def _patch_sheet(text: str, number: int) -> str:
    masks: dict[int, list[str]] = {}

    def patch(match: re.Match[str]) -> str:
        col, row_text = match.group(1, 2)
        row = int(row_text)
        cell = match.group(0)
        f = FORMULA.search(cell)
        if number == 2:
            opening = re.match(r"<c\b([^>]*?)(?:/>|>)", cell)
            assert opening is not None
            attributes = re.sub(r'\s(?:r|t)="[^"]*"', "", opening.group(1))
            if row in (28, 34):
                if col in ("A", "K"):
                    return f'<c r="{col}{row}"{attributes} t="inlineStr"><is><t>Нет</t></is></c>'
                if col in ("D",):
                    return f'<c r="{col}{row}"{attributes}><v>0</v></c>'
                if col in ("P", "T"):
                    return _formula_cell(col, row, '""', attributes)
                return re.sub(r">.*</c>$", "/>", re.sub(r' t="[^"]*"', "", cell), flags=re.DOTALL)
            if (col == "L" and 21 <= row <= 50) or (col == "B" and row == 12):
                return (
                    f'<c r="{col}{row}"{attributes} t="inlineStr"><is><t>Не учитывать</t></is></c>'
                )
            if col == "T" and 21 <= row <= 50:
                return _formula_cell(col, row, '"IGNORE"', attributes)
            if col == "B" and row == 18:
                return _formula_cell(col, row, 'IF(B17>0,"OK","Проверить веса")', attributes)
            return cell
        if f is None:
            return cell
        formula = _capacity_formula(html.unescape(f.group(1)))
        if number == 4 and 20 <= row <= 119:
            idx = sum((ord(c) - 64) * 26**i for i, c in enumerate(reversed(col)))
            if 2 <= idx <= 31:
                slot = idx - 1
                raw = f"{get_column_letter(32 + slot)}{row}"
                mask_col = get_column_letter(94 + slot)
                if slot in REMOVED_SLOTS:
                    formula, mask = '""', "0"
                else:
                    formula, missing = _mu(formula, raw)
                    mask = (
                        f'IF($A{row}="",0,IF(${col}$4<>"Да",0,'
                        f"IF({missing},0,IF(ISNUMBER({col}{row}),1,0))))"
                    )
                masks.setdefault(row, []).append(_formula_cell(mask_col, row, mask))
            elif 33 <= idx <= 62 or 64 <= idx <= 93:
                slot = idx - 32 if idx <= 62 else idx - 63
                if slot in REMOVED_SLOTS:
                    formula = '""'
                else:
                    # Excel INDEX(blank) returns 0: guard before returning raw data.
                    lookup = re.search(r"INDEX\(Данные!.*", formula)
                    assert lookup is not None
                    expression = "INDEX(" + ",".join(_arguments(lookup.group(0)[6:])) + ")"
                    assert expression.count("(") == expression.count(")")
                    formula = formula.replace(expression, f'IF({expression}="","",{expression})')
        if number == 5 and 4 <= row <= 103:
            criterion_row = row + 16
            terms = []
            weights = []
            counts = []
            grounds = []
            reasons = []
            for slot in range(1, 31):
                if slot in REMOVED_SLOTS:
                    continue
                c = get_column_letter(slot + 1)
                mask = f"Критерии!{get_column_letter(94 + slot)}{criterion_row}"
                mu = f"Критерии!{c}{criterion_row}"
                weight = f"Критерии!{c}$6"
                counts.append(mask)
                grounds.append(f'IF(AND({mask}=1,{mu}>=Настройки!$B$16,Критерии!{c}$7="Да"),1,0)')
                settings_row = slot + 20
                reasons.append(
                    f"IF(AND({mask}=1,{mu}>=Настройки!$B$16,"
                    f'Настройки!$K${settings_row}="Да"),'
                    f'Настройки!$B${settings_row}&"; ","")'
                )
                weights.append(f"IF({mask}=1,{weight},0)")
                terms.append(f"IF({mask}=1,{weight}*LN(MAX(Настройки!$B$14,{mu})),0)")
            if col == "E":
                formula = f'IF($A{row}="","",' + "+".join(counts) + ")"
            elif col == "F":
                formula = f'IF($A{row}="","",D{row}-E{row})'
            elif col == "G":
                formula = f'IF($A{row}="","",' + "+".join(grounds) + ")"
            elif col == "N":
                formula = (
                    f'IF($A{row}="","",IF(G{row}=0,"Нет сработавших оснований",'
                    + "&".join(reasons)
                    + "))"
                )
            elif col == "H":
                formula = f'IF($A{row}="","",' + "+".join(weights) + ")"
            elif col == "J":
                formula = (
                    f'IF($A{row}="","",IF(OR(H{row}=0,COUNTIF('
                    f'Критерии!$B{criterion_row}:$AE{criterion_row},"Некорректные настройки")>0),'
                    '"",EXP((' + "+".join(terms) + f")/H{row})))"
                )
        updated = FORMULA.sub(lambda _: "<f>" + html.escape(formula, quote=False) + "</f>", cell)
        return re.sub(r"<v>.*?</v>", "", updated, flags=re.DOTALL)

    text = CELL.sub(patch, text)
    if number in (5, 6):

        def ranking_row(match: re.Match[str]) -> str:
            row = int(match.group(1))
            body = match.group(0)
            if not 4 <= row <= 103:
                return body
            if number == 5:
                formula = (
                    f'IF($A{row}="","",IF(ISNUMBER(K{row}),K{row},'
                    f"COUNT($J$4:$J${CAPACITY + 3})+"
                    f'COUNTIFS($A$4:A{row},"<>",$J$4:J{row},"")))'
                )
                helper = _formula_cell("O", row, formula)
            else:
                formula = (
                    f'IF(ROW()-3>COUNTA(Данные!$A$7:$A${CAPACITY + 6}),"",'
                    f"MATCH(ROW()-3,Расчет!$O$4:$O${CAPACITY + 3},0))"
                )
                helper = _formula_cell("M", row, formula)

                def rating_cell(cell_match: re.Match[str]) -> str:
                    col = cell_match.group(1)
                    cell = cell_match.group(0)
                    if col not in "ABCDEFGHI":
                        return cell
                    f = FORMULA.search(cell)
                    if f is None:
                        return cell
                    value = html.unescape(f.group(1))
                    if col == "A":
                        lookup = f"INDEX(Расчет!$K$4:$K${CAPACITY + 3},$M{row})"
                        value = f'IF($M{row}="","",IF({lookup}="","",{lookup}))'
                    else:
                        value = value.replace(f'$A{row}=""', f'$M{row}=""')
                        value = value.replace(
                            f"MATCH($A{row},Расчет!$K$4:$K${CAPACITY + 3},0)", f"$M{row}"
                        )
                    return FORMULA.sub(
                        lambda _: "<f>" + html.escape(value, quote=False) + "</f>", cell
                    )

                body = CELL.sub(rating_cell, body)
            return body.replace("</row>", helper + "</row>")

        text = re.sub(r'<row\b[^>]* r="(\d+)"[^>]*>.*?</row>', ranking_row, text, flags=re.DOTALL)
        letter, index = ("O", 15) if number == 5 else ("M", 13)
        text = re.sub(r'(<dimension ref="[^"]*?:)[A-Z]+', lambda m: m.group(1) + letter, text)
        text = text.replace(
            "</cols>",
            f'<col min="{index}" max="{index}" width="10" hidden="1" customWidth="1"/></cols>',
        )
    if number == 4:
        text = re.sub(r'(<dimension ref="[^"]*?:)[A-Z]+', r"\g<1>DT", text)
        for row, cells in masks.items():
            text = re.sub(
                rf'(<row\b[^>]* r="{row}"[^>]*>)(.*?)(</row>)',
                lambda m, cells=cells: m.group(1) + m.group(2) + "".join(cells) + m.group(3),
                text,
                count=1,
                flags=re.DOTALL,
            )
        text = text.replace(
            "</cols>", '<col min="95" max="124" width="10" hidden="1" customWidth="1"/></cols>'
        )
    if number in (4, 5, 6):
        max_column = {4: 124, 5: 15, 6: 13}[number]
        text = re.sub(
            r'(<row\b[^>]*\bspans=")\d+:\d+("[^>]*>)',
            lambda m: m.group(1) + f"1:{max_column}" + m.group(2),
            text,
        )
        end = 119 if number == 4 else 103
        template = re.search(rf'<row\b[^>]* r="{end}"[^>]*>.*?</row>', text, re.DOTALL)
        assert template is not None
        base = template.group(0)
        rows = []
        for target in range(end + 1, end + CAPACITY - 99):

            def translate(m: re.Match[str], target: int = target) -> str:
                col = m.group(1)
                cell = m.group(0)
                f = FORMULA.search(cell)
                if f:
                    formula = Translator(
                        "=" + html.unescape(f.group(1)), origin=f"{col}{end}"
                    ).translate_formula(f"{col}{target}")[1:]
                    cell = FORMULA.sub(
                        lambda _: "<f>" + html.escape(formula, quote=False) + "</f>", cell
                    )
                return cell.replace(f'r="{col}{end}"', f'r="{col}{target}"')

            rows.append(CELL.sub(translate, base).replace(f'r="{end}"', f'r="{target}"'))
        text = text.replace("</sheetData>", "".join(rows) + "</sheetData>")
        new_end = end + CAPACITY - 100
        text = re.sub(r'(<dimension ref="[^"]*?:[A-Z]+)\d+("/>)', rf"\g<1>{new_end}\2", text)
        text = re.sub(r'(sqref="[^"]*)\b([A-Z]+)' + str(end) + r"\b", rf"\g<1>\g<2>{new_end}", text)
    if number == 2:
        text = text.replace(
            "</sheetData>",
            '<row r="52"><c r="B52" t="inlineStr"><is><t>'
            "Не используется: Расходомер; Последний АДД; Не проведен АДД; "
            "Техническая возможность. Manual μ: резерв, не используется."
            "</t></is></c></row></sheetData>",
        )
        # Fixed policies and retired slots are locked, neutral cells. Protect
        # settings while keeping the existing yellow input styles unlocked.
        text = re.sub(
            r'<dataValidation\b[^>]*sqref="(?:B12|L21:L50)"[^>]*>.*?</dataValidation>',
            "",
            text,
            flags=re.DOTALL,
        )
        text = re.sub(r'(<dataValidations count=")\d+', r"\g<1>5", text)

        def neutral_policy(m: re.Match[str]) -> str:
            col, row = m.group(1), int(m.group(2))
            if (col == "B" and row == 12) or (col == "L" and 21 <= row <= 50):
                return re.sub(r' s="\d+"', ' s="0"', m.group(0))
            return m.group(0)

        text = CELL.sub(neutral_policy, text)
        text = text.replace('ref="A1:U50"', 'ref="A1:U52"')
        for row in (28, 34):
            text = re.sub(rf'(<row\b[^>]* r="{row}"[^>]*)(>)', r'\1 hidden="1"\2', text)
        text = text.replace(
            "</sheetData>", '</sheetData><sheetProtection sheet="1" objects="1" scenarios="1"/>'
        )
    return text


def migrate(source: Path, output: Path) -> tuple[str, ...]:
    """Produce a separate candidate from the exact accepted source; no overwrite."""
    if sha256(source) != SOURCE_SHA256:
        raise ValueError("Unexpected source SHA-256 (or already migrated template)")
    if source.resolve() == output.resolve() or output.exists():
        raise ValueError("Output must be a new separate file")
    patches = {}
    with ZipFile(source) as archive:
        for number in (2, 4, 5, 6):
            name = f"xl/worksheets/sheet{number}.xml"
            patches[name] = _patch_sheet(archive.read(name).decode(), number).encode()
        for number in (3, 4):
            name = f"xl/tables/table{number}.xml"
            patches[name] = archive.read(name).replace(b"103", str(CAPACITY + 3).encode())
        name = "xl/styles.xml"
        styles = archive.read(name).decode()

        def unlock_inputs(match: re.Match[str]) -> str:
            body = match.group(0)
            if 'fillId="4"' not in body:
                return body
            body = re.sub(r"<protection\b[^>]*/>", "", body)
            if body.endswith("/>"):
                body = body[:-2] + '><protection locked="0"/></xf>'
            else:
                body = body.replace("</xf>", '<protection locked="0"/></xf>')
            return body.replace("<xf ", '<xf applyProtection="1" ', 1)

        patches[name] = re.sub(r"<xf\b[^>]*(?:/>|>.*?</xf>)", unlock_inputs, styles).encode()
        name = "xl/workbook.xml"
        patches[name] = re.sub(
            rb"<calcPr[^>]*/>",
            b'<calcPr calcId="144525" calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>',
            archive.read(name),
        ).replace(
            b"<calcPr",
            b'<definedNames><definedName name="GDIS_TEMPLATE_VERSION">2</definedName>'
            b"</definedNames><calcPr",
        )
        name = "xl/_rels/workbook.xml.rels"
        patches[name] = re.sub(
            rb'<Relationship\b[^>]*Type="[^"]*/calcChain"[^>]*/>', b"", archive.read(name)
        )
        name = "[Content_Types].xml"
        patches[name] = re.sub(
            rb'<Override\b[^>]*PartName="/xl/calcChain.xml"[^>]*/>', b"", archive.read(name)
        )
        changed = {name for name, data in patches.items() if archive.read(name) != data} | {
            "xl/calcChain.xml"
        }
        if changed != EXPECTED_CHANGED:
            raise ValueError(f"Unexpected migration parts: {changed ^ EXPECTED_CHANGED}")
        for data in patches.values():
            root = ET.fromstring(data)
            for element in root.iter():
                if "sqref" in element.attrib:
                    MultiCellRange(element.attrib["sqref"])
        output.parent.mkdir(parents=True, exist_ok=True)
        with ZipFile(output, "w") as destination:
            destination.comment = archive.comment
            for info in archive.infolist():
                if info.filename != "xl/calcChain.xml":
                    destination.writestr(
                        info, patches.get(info.filename, archive.read(info.filename))
                    )
    return tuple(sorted(changed))
