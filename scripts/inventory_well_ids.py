"""Privacy-safe inventory of no-separator identities in runtime Excel inputs."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gdis_loader.matching.normalize import normalize_field_name, normalize_well_id
from gdis_loader.readers.fund import read_fund

_NO_SEPARATOR = re.compile(r"^\d+(?:ЗБС|ГС|МЗС|ПВЛГ|ПНЛГ|ПВР|КЛК)", re.IGNORECASE)


def inventory(runtime_manifest: dict[str, Any], source_dir: Path) -> list[dict[str, Any]]:
    """Scan every source sheet, counting rows once and exposing no row payload."""
    fund_entry = next(item for item in runtime_manifest["sources"] if item["role"] == "fund")
    fund = read_fund(source_dir / fund_entry["name"])
    field = normalize_field_name(runtime_manifest["selected_field"])
    targets = tuple(
        well for well in fund.wells if well.is_rating_eligible and well.normalized_field == field
    )
    result: list[dict[str, Any]] = []
    for entry in runtime_manifest["sources"]:
        if entry["role"] == "generated_calculator":
            continue
        workbook = load_workbook(source_dir / entry["name"], read_only=True, data_only=True)
        row_keys: set[tuple[str, int]] = set()
        canonical: set[str] = set()
        roots: set[str] = set()
        try:
            for sheet in workbook:
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    for value in row:
                        if not isinstance(value, str) or not _NO_SEPARATOR.match(value.strip()):
                            continue
                        parsed = normalize_well_id(value)
                        if parsed.recognized_variant is None:
                            continue
                        row_keys.add((sheet.title, number))
                        if parsed.target_well_id is not None:
                            canonical.add(parsed.target_well_id)
                        if parsed.base_well_id is not None:
                            roots.add(parsed.base_well_id)
        finally:
            workbook.close()
        result.append(
            {
                "source_role": entry["role"],
                "count_rows": len(row_keys),
                "distinct_canonical_forms": sorted(canonical),
                "affected_target_counts": sum(well.base_well_id in roots for well in targets),
            }
        )
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runtime-report", type=Path, required=True)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    manifest = json.loads(args.runtime_report.read_text(encoding="utf-8"))["source_manifest"]
    args.output.write_text(
        json.dumps(inventory(manifest, args.source_dir), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
