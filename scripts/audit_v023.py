"""Private production reconciliation and minimal publishable v0.23 evidence."""

from __future__ import annotations

import argparse
import json
from dataclasses import replace
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from gdis_loader.domain.models import SourcePaths
from gdis_loader.service import run_update

FIELDS = ("Береговое", "Пырейное", "Хадырьяхинское")


def audit(data: Path, template: Path, directory: Path, operation_date: date) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    paths = SourcePaths(
        calculator=template,
        fund=data / "Фонд 2026 на 01.08.2026.xlsx",
        gdi=data / "Сводная таблица ГДИ.xlsx",
        lha=data / "База ЛХА_Итог!.xlsm",
        daily_report=data / "07.2026 Cуточный рапорт БГП_полный.xlsm",
        well_history=data / "База по скважинам.xlsx",
        research_status=data / "Статус исследований_2026.xlsx",
    )
    before_path = directory / "before/summary.json"
    before = json.loads(before_path.read_text()) if before_path.exists() else {}
    summaries: dict[str, Any] = {}
    after: dict[str, Any] = {}
    evidence: dict[str, Any] = {}
    for field in FIELDS:
        output = directory / f"UPDATED_{field}.xlsx"
        result = run_update(
            replace(paths, output=output),
            directory / f"{field}.log",
            target_field=field,
            calculation_date=operation_date,
        )
        if result.has_errors:
            raise ValueError(f"Production update failed: {field}")
        report = result.report.to_dict()
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            sheet = workbook["Данные"]
            headers = [c.value for c in sheet[6]]
            rows = {
                str(row[0]): dict(zip(headers, row, strict=True))
                for row in sheet.iter_rows(min_row=7, values_only=True)
                if row[0] is not None
            }
        finally:
            workbook.close()
        audit_rows = report["gdi_reconciliation"]
        counted = [r for r in audit_rows if r["counted"]]
        assert sum(row["Количество ГДИ"] or 0 for row in rows.values()) == len(counted)
        previous = before.get(field, {}).get("rows", {})
        previous_report = before.get(field, {}).get("report", {})
        base_previous = {k.split("_")[0]: row for k, row in previous.items()}
        changed = sum(
            row["Количество ГДИ"] != base_previous.get(k.split("_")[0], {}).get("Количество ГДИ")
            for k, row in rows.items()
        )
        summaries[field] = {
            "before": {
                key: previous_report.get(key)
                for key in (
                    "rating_eligible",
                    "matched",
                    "ambiguous",
                    "no_gdi",
                    "invalid_source",
                    "downhole_yes",
                    "gdi_pre_launch_excluded",
                )
            },
            "after": {
                key: report[key]
                for key in (
                    "rating_eligible",
                    "matched",
                    "ambiguous",
                    "no_gdi",
                    "invalid_source",
                    "downhole_yes",
                    "gdi_pre_launch_excluded",
                    "gdi_pre_launch_included",
                    "lnd_previous_fallbacks",
                )
            },
            "total_counted_gdi_before": sum(
                row["Количество ГДИ"] or 0 for row in previous.values()
            ),
            "total_counted_gdi_after": len(counted),
            "wells_count_changed": changed,
            "counted_ksd": sum(r["canonical_research_type"] == "KSD" for r in counted),
            "counted_kvd_id": sum(
                (r["raw_research_type"] or "").upper().startswith("КВД+ИД") for r in counted
            ),
            "restored_by_lifecycle": sum(
                r["counted"]
                for r in audit_rows
                if any(
                    i["code"] == "SUMMARY_GDI_CURRENT_HISTORY_INFERRED"
                    and (i["well_id"] or "").split("_")[0] == (r["target_well"] or "").split("_")[0]
                    for i in report["issues"]
                )
            ),
            "restored_by_plast_normalization": 0,
            "plast_normalization_note": (
                "Canonicalizer unchanged from v0.22; no new normalization recovery."
            ),
        }
        for root in ("10306", "10405"):
            control = [r for r in audit_rows if (r["canonical_well"] or "").split("_")[0] == root]
            if not control:
                continue
            # Minimum necessary customer-approved evidence: no absolute user paths.
            for row in control:
                row["source_file"] = Path(row["source_file"]).name
            evidence[root] = {
                "field": field,
                "old_count": base_previous.get(root, {}).get("Количество ГДИ"),
                "new_count": sum(r["counted"] for r in control),
                "rows": control,
            }
        after[field] = {"report": report, "rows": rows}
        print(field, summaries[field], flush=True)
    (directory / "after-summary.json").write_text(
        json.dumps(after, ensure_ascii=False, indent=2, default=str)
    )
    (directory / "production-summary.json").write_text(
        json.dumps(summaries, ensure_ascii=False, indent=2)
    )
    (directory / "controls-reconciliation.json").write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2)
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, default=Path("data"))
    parser.add_argument(
        "--template",
        type=Path,
        default=Path("templates/calculator/GDIS_Rating_Calculator_STRUCTURAL_v3.xlsx"),
    )
    parser.add_argument("--directory", type=Path, default=Path("artifacts/v0.23"))
    parser.add_argument("--calculation-date", type=date.fromisoformat, default=date.today())
    args = parser.parse_args()
    audit(args.data, args.template, args.directory, args.calculation_date)
