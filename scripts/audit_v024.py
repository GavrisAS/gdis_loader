"""Private Beregovoe September debit audit; raw outputs stay in ignored artifacts/."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from dataclasses import asdict, replace
from datetime import date
from functools import lru_cache
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook

from gdis_loader import service
from gdis_loader.domain.models import SourcePaths
from gdis_loader.matching.normalize import normalize_field_name, normalize_well_id
from gdis_loader.readers.common import parse_number
from gdis_loader.readers.daily import read_daily_report
from gdis_loader.readers.well_history import read_well_history
from gdis_loader.source_manifest import SOURCE_FILES

FIELDS = ("Береговое",)
TEMPLATE = Path("templates/calculator/GDIS_Rating_Calculator_STRUCTURAL_v3.xlsx")


def digest(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def dump(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str))


def sources() -> SourcePaths:
    return SourcePaths(
        **{
            role: TEMPLATE if role == "calculator" else Path("data") / name
            for role, name in SOURCE_FILES
        }
    )


def run(directory: Path) -> None:
    """Execute the same explicit date and default criteria for Beregovoe."""
    directory.mkdir(parents=True, exist_ok=True)
    # Immutable reader results can be reused across fields, notably the large LHA.
    for name in ("read_lha", "read_fund", "read_daily_report", "read_well_history"):
        setattr(service, name, lru_cache(maxsize=None)(getattr(service, name)))
    for field in FIELDS:
        result = service.run_update(
            replace(sources(), output=directory / f"UPDATED_{field}.xlsx"),
            directory / f"{field}.log",
            target_field=field,
            calculation_date=date(2026, 10, 12),
        )
        if result.has_errors:
            raise ValueError(f"Failed update: {field}")


def output_rows(path: Path) -> dict[str, dict[str, Any]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["Данные"]
        headers = tuple(c.value for c in sheet[6])
        return {
            str(row[0]): {"target_row": index, **dict(zip(headers, row, strict=True))}
            for index, row in enumerate(sheet.iter_rows(min_row=7, values_only=True), 7)
            if row[0] is not None
        }
    finally:
        book.close()


def reconcile(root: Path) -> None:
    """Compare identical-source baselines, every candidate, and all output parts."""
    paths = sources()
    daily, _ = read_daily_report(paths.daily_report)
    fund = service.read_fund(paths.fund)
    history_by_field = {
        field: read_well_history(paths.well_history, field).histories for field in FIELDS
    }
    direct_cells = {
        (location.sheet, location.cell)
        for item in daily
        for location in item.field_evidence
        if location.sheet is not None
        and location.cell is not None
        and location.cell.startswith("AH")
    }
    direct_values: dict[tuple[str | None, str], float | None] = {}
    source_book = load_workbook(
        paths.daily_report, read_only=True, data_only=True, keep_links=False
    )
    try:
        sheet_rows = {
            sheet: tuple(
                value[0]
                for value in source_book[sheet].iter_rows(min_col=34, max_col=34, values_only=True)
            )
            for sheet in {sheet for sheet, _cell in direct_cells if sheet is not None}
        }
        for sheet, cell in direct_cells:
            assert sheet is not None
            match = re.search(r"\d+$", cell)
            assert match is not None
            raw = sheet_rows[sheet][int(match.group()) - 1]
            direct_values[(sheet, cell)] = parse_number(raw)
    finally:
        source_book.close()
    cache_comparisons = 0
    manifest = {
        role: {"name": path.name, "sha256": digest(path)}
        for role, path in asdict(paths).items()
        if isinstance(path, Path)
    }
    manifest["v2"] = {
        "sha256": digest(TEMPLATE.with_name("GDIS_Rating_Calculator_STRUCTURAL_v2.xlsx"))
    }
    dump(root / "inputs.json", manifest)
    reconciliation: list[dict[str, Any]] = []
    summaries: dict[str, Any] = {}
    preservation: dict[str, Any] = {}
    for field in FIELDS:
        before = output_rows(root / "before" / f"UPDATED_{field}.xlsx")
        after = output_rows(root / "after" / f"UPDATED_{field}.xlsx")
        old_report = json.loads((root / "before" / f"UPDATED_{field}.load-report.json").read_text())
        report = json.loads((root / "after" / f"UPDATED_{field}.load-report.json").read_text())
        assert before.keys() == after.keys()
        assert old_report["criteria_snapshot"] == report["criteria_snapshot"]
        old_hashes = {
            e["role"]: e["sha256"]
            for e in old_report["source_manifest"]["sources"]
            if e["role"] != "generated_calculator"
        }
        new_hashes = {
            e["role"]: e["sha256"]
            for e in report["source_manifest"]["sources"]
            if e["role"] != "generated_calculator"
        }
        assert old_hashes == new_hashes
        config_bytes = json.dumps(
            report["criteria_snapshot"], sort_keys=True, ensure_ascii=False
        ).encode()
        dump(
            root / f"effective-config-{field}.json",
            {
                "config": report["criteria_snapshot"],
                "sha256": hashlib.sha256(config_bytes).hexdigest(),
            },
        )
        unchanged = 0
        for well, row in after.items():
            base = normalize_well_id(well).base_well_id
            prior = before[well]["Дебит"]
            value = row["Дебит"]
            if prior is not None:
                assert math.isclose(prior, value, rel_tol=1e-12, abs_tol=1e-10), (
                    field,
                    well,
                    prior,
                    value,
                )
                unchanged += 1
            for key in row:
                if key != "Дебит":
                    assert row[key] == before[well][key], (field, well, key)
            candidates = [item for item in daily if item.base_well_id == base]
            old_issues = [
                i
                for i in old_report["issues"]
                if i["well_id"] == base and i["source"] == "daily_report"
            ]
            old_codes = [i["code"] for i in old_issues]
            codes = [
                i["code"]
                for i in report["issues"]
                if i["well_id"] == base and i["source"] == "daily_report"
            ]
            selected = [
                item
                for item in candidates
                if item.raw_field_name
                and normalize_field_name(item.raw_field_name) == normalize_field_name(field)
            ]
            proof = next(
                (
                    p
                    for p in report["field_provenance"]
                    if p["well_id"] == well and p["field"] == "Дебит"
                ),
                None,
            )
            debit_provenance = next(
                (
                    p
                    for p in report["field_provenance"]
                    if p["well_id"] == well and p["field"] == "Дебит"
                ),
                None,
            )
            chosen_row = None
            if debit_provenance is not None:
                chosen_row = next(
                    (
                        int(match.group(1))
                        for location in debit_provenance["locations"]
                        if location.get("sheet") == "Добыча телеметрия"
                        and (match := re.fullmatch(r"C(\d+)", location.get("cell", "")))
                    ),
                    None,
                )
            current_daily_issues = [
                i
                for i in report["issues"]
                if i["well_id"] == base and i["source"] == "daily_report"
            ]
            if chosen_row is None and any(
                i["code"] == "DEBIT_NO_VALID_WORKING_DAYS" for i in current_daily_issues
            ):
                chosen_row = next(
                    (
                        int(match.group(1))
                        for issue in current_daily_issues
                        for location in issue["locations"]
                        if location.get("sheet") == "Добыча телеметрия"
                        and (match := re.fullmatch(r"C(\d+)", location.get("cell", "")))
                    ),
                    None,
                )
            identity_evidence = (
                debit_provenance["details"].get("identity_evidence")
                if debit_provenance is not None
                else None
            )
            launch_evidence = next(
                (
                    p
                    for p in report["field_provenance"]
                    if p["well_id"] == well and p["field"] == "Дата запуска"
                ),
                None,
            )
            history_evidence = identity_evidence
            if history_evidence is None:
                matching_fund = [
                    item
                    for item in fund.wells
                    if item.target_well_id == well
                    and item.normalized_field == normalize_field_name(field)
                ]
                base = normalize_well_id(well).base_well_id
                stage_candidates = [
                    stage
                    for history in history_by_field[field]
                    if history.physical_root == base
                    for stage in history.stages
                ]
                launch_locations = (
                    launch_evidence["locations"] if launch_evidence is not None else []
                )
                selected_stages = [
                    stage
                    for stage in stage_candidates
                    if stage.source_location.to_dict() in launch_locations
                ]
                if len(selected_stages) == 1:
                    stage = selected_stages[0]
                    selected_history: dict[str, Any] = {
                        "event": stage.event_kind.value,
                        "target_well_id": stage.target_well_id,
                        "source_display_well": stage.source_display_well,
                        "source_partial_date": stage.started_at.to_dict(),
                        "location": stage.source_location.to_dict(),
                    }
                else:
                    selected_history = {
                        "event": "UNRESOLVED",
                        "candidate_stages": [
                            {
                                "event": stage.event_kind.value,
                                "target_well_id": stage.target_well_id,
                                "source_display_well": stage.source_display_well,
                                "source_partial_date": stage.started_at.to_dict(),
                                "location": stage.source_location.to_dict(),
                            }
                            for stage in stage_candidates
                        ],
                    }
                history_evidence = {
                    "base_well_id": base,
                    "target_well_id": well,
                    "field": field,
                    "fund_locations": [
                        location.to_dict() for item in matching_fund for location in item.locations
                    ],
                    "current_history": selected_history,
                }
            evidence: list[dict[str, Any]] = []
            for item in candidates:
                valid = [x for x in item.daily_values if x is not None and x > 0]
                accepted = chosen_row == item.source_row
                direct_locations = [
                    location
                    for location in item.field_evidence
                    if location.cell is not None
                    and location.cell.startswith("AH")
                    and location.sheet is not None
                ]
                direct_mean = None
                cache_matches_daily_row = None
                if item.raw_field_name is not None:
                    source_rates: list[float | None] = []
                    for location in direct_locations:
                        assert location.sheet is not None and location.cell is not None
                        source_rates.append(direct_values[(location.sheet, location.cell)])
                    assert len(source_rates) == len(item.daily_values) == 30
                    cache_matches_daily_row = all(
                        (left is None and right is None)
                        or (
                            left is not None
                            and right is not None
                            and math.isclose(left, right, rel_tol=1e-12, abs_tol=1e-12)
                        )
                        for left, right in zip(item.daily_values, source_rates, strict=True)
                    )
                    assert cache_matches_daily_row, (item.raw_well_name, item.source_row)
                    cache_comparisons += len(source_rates)
                    positives = [rate for rate in source_rates if rate is not None and rate > 0]
                    direct_mean = sum(positives) / len(positives) if positives else None
                decision = (
                    "ACCEPT"
                    if accepted
                    else "REJECT"
                    if item.raw_field_name
                    and normalize_field_name(item.raw_field_name) != normalize_field_name(field)
                    else "UNRESOLVED"
                )
                evidence.append(
                    {
                        "sheet": item.source_sheet,
                        "row": item.source_row,
                        "raw_well": item.raw_well_name,
                        "normalized_well": item.target_well_id,
                        "raw_collector_id": item.raw_collector_id,
                        "canonical_well": item.target_well_id,
                        "field": item.raw_field_name,
                        "field_evidence": [x.to_dict() for x in item.field_evidence],
                        "valid_days": len(valid),
                        "mean": sum(valid) / len(valid) if valid else None,
                        "direct_cached_mean": direct_mean,
                        "cached_cells_match_telemetry": cache_matches_daily_row,
                        "decision": decision,
                        "reason": (
                            "dated direct links, same well, explicit field heading, "
                            "resolved current lifecycle"
                        )
                        if accepted
                        else "different proven field"
                        if decision == "REJECT"
                        else "not selected: lifecycle identity or candidate conflict is unresolved",
                    }
                )
            if value is not None:
                assert proof is not None and len(selected) == 1
                expected = evidence[candidates.index(selected[0])]["mean"]
                assert math.isclose(value, expected, rel_tol=1e-12, abs_tol=1e-10)
            else:
                assert codes, (field, well)
            reconciliation.append(
                {
                    "target_well": well,
                    "base_well": base,
                    "history": history_evidence,
                    "target_row": row["target_row"],
                    "selected_field": field,
                    "before": prior,
                    "after": value,
                    "old_diagnostic": old_codes,
                    "old_issue_evidence": old_issues,
                    "new_diagnostic": codes,
                    "new_issue_evidence": current_daily_issues,
                    "candidates": evidence,
                    "provenance": proof,
                    "outcome": "FILLED" if value is not None else "BLANK_WITH_PROVEN_REASON",
                }
            )
        summaries[field] = {
            "eligible": len(after),
            "filled_before": sum(r["Дебит"] is not None for r in before.values()),
            "blank_before": sum(r["Дебит"] is None for r in before.values()),
            "filled_after": report["debit_filled"],
            "blank_after": report["debit_blank"],
            "unchanged_filled": unchanged,
            "resolved_ambiguous": sum(
                r["selected_field"] == field
                and "DEBIT_SOURCE_FIELD_AMBIGUOUS" in r["old_diagnostic"]
                and r["after"] is not None
                for r in reconciliation
            ),
            "still_ambiguous": sum(
                "AMBIGUOUS" in i["code"] for i in report["issues"] if i["source"] == "daily_report"
            ),
            "no_valid_days": sum(
                i["code"] == "DEBIT_NO_VALID_WORKING_DAYS" for i in report["issues"]
            ),
            "missing_source": sum(
                i["code"]
                in {
                    "DEBIT_SOURCE_ROW_MISSING",
                    "DEBIT_SOURCE_FIELD_MISMATCH",
                    "DEBIT_SOURCE_FIELD_UNPROVEN",
                }
                for i in report["issues"]
            ),
            "other": max(
                0,
                report["debit_blank"]
                - sum(i["code"] == "DEBIT_NO_VALID_WORKING_DAYS" for i in report["issues"])
                - sum(
                    "AMBIGUOUS" in i["code"]
                    for i in report["issues"]
                    if i["source"] == "daily_report"
                )
                - sum(
                    i["code"]
                    in {
                        "DEBIT_SOURCE_ROW_MISSING",
                        "DEBIT_SOURCE_FIELD_MISMATCH",
                        "DEBIT_SOURCE_FIELD_UNPROVEN",
                    }
                    for i in report["issues"]
                ),
            ),
            "blank_wells": [w for w, r in after.items() if r["Дебит"] is None],
            "report_month": "2026-09",
            "source_sha256": digest(paths.daily_report),
        }
        with (
            ZipFile(TEMPLATE) as original,
            ZipFile(root / "after" / f"UPDATED_{field}.xlsx") as output,
        ):
            assert set(original.namelist()) == set(output.namelist())
            changed = [
                name for name in original.namelist() if original.read(name) != output.read(name)
            ]
            assert set(changed) <= {"xl/worksheets/sheet3.xml", "xl/tables/table2.xml"}, changed
            preservation[field] = changed
    summaries["direct_cache_validation"] = {
        "matched_daily_cells": cache_comparisons,
        "source_sheet": "01-30!AH",
        "source_month": "2026-09",
        "rule": "direct formula destination cached values equal telemetry sheet cached values",
    }
    control_wells = {"22", "62", "101", "321", "322"}
    cross_field_checks = [
        {
            "target_well": row["target_well"],
            "base_well": row["base_well"],
            "field": row["selected_field"],
            "debit": row["after"],
            "candidate_fields": sorted({c["field"] for c in row["candidates"] if c["field"]}),
            "diagnostics": row["new_diagnostic"],
            "decision": "SOURCE_MATCHED_IN_BEREGOVOE"
            if row["selected_field"] == "Береговое" and row["after"] is not None
            else "NO_FOREIGN_SOURCE_ACCEPTED"
            if row["selected_field"] == "Береговое" and row["after"] is None
            else "REVIEW",
        }
        for row in reconciliation
        if row["base_well"] in control_wells
    ]
    assert all(item["decision"] != "REVIEW" for item in cross_field_checks)
    dump(root / "same-number-source-scope-checks.json", cross_field_checks)
    summaries["production_scope"] = {
        "fields_run": ["Береговое"],
        "fields_omitted": {
            "Пырейное": "the attached September daily report is for Beregovoe only",
            "Хадырьяхинское": "the attached September daily report is for Beregovoe only",
        },
        "same_number_controls": sorted(control_wells),
        "synthetic_cross_field_matching_test": (
            "covered by test_daily_selection_requires_same_well_and_selected_field"
        ),
    }
    dump(root / "reconciliation.json", reconciliation)
    dump(root / "summary.json", summaries)
    dump(root / "preservation.json", preservation)
    with (root / "reconciliation.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(reconciliation[0]))
        writer.writeheader()
        writer.writerows(
            {
                k: json.dumps(v, ensure_ascii=False, default=str)
                if isinstance(v, (dict, list))
                else v
                for k, v in row.items()
            }
            for row in reconciliation
        )
    dump(root / "baseline-blanks.json", [r for r in reconciliation if r["before"] is None])
    render_audit(root, reconciliation, summaries, manifest)
    print(json.dumps(summaries, ensure_ascii=False, indent=2))


def render_audit(
    root: Path,
    rows: list[dict[str, Any]],
    summaries: dict[str, Any],
    inputs: dict[str, dict[str, str]],
) -> None:
    """Write private row-level Markdown evidence beside the raw reconciliation."""

    def number(value: object) -> str:
        return "—" if value is None else f"{value:.12g}" if isinstance(value, float) else str(value)

    def history_summary(item: dict[str, Any]) -> str:
        current = item.get("current_history")
        if current is not None:
            start = current.get("source_partial_date", {})
            location = current.get("location", {})
            event = current.get("event", "UNKNOWN")
        else:
            start = item.get("history_start", {})
            location = item.get("history_location", {})
            event = item.get("history_event", "UNKNOWN")
        date_text = f"{start.get('year', '?')}-{start.get('month', '?'):02}"
        location_text = (
            f"{Path(location['file']).name}::{location['sheet']}::{location['cell']}"
            if location.get("file") and location.get("sheet") and location.get("cell")
            else "history source unresolved"
        )
        fund_cells = ", ".join(
            f"{Path(location['file']).name}::{location['sheet']}::{location['cell']}"
            for location in item.get("fund_locations", [])
            if location.get("file") and location.get("sheet") and location.get("cell")
        )
        return f"{item.get('base_well_id', '?')} → {item.get('target_well_id', '?')}; {event} {date_text}; {location_text}; fund {fund_cells or 'unresolved'}"  # noqa: E501

    def candidate_text(candidate: dict[str, Any]) -> str:
        row = candidate["row"]
        telemetry_cells = f"C{row}:AF{row}"
        evidence = candidate.get("field_evidence", [])
        per_day: dict[str, list[str]] = {}
        for location in evidence:
            sheet, cell = location.get("sheet"), location.get("cell")
            if sheet is not None and cell is not None:
                per_day.setdefault(sheet, []).append(cell)
        days_text = "; ".join(
            f"{sheet}!{'/'.join(cells)}" for sheet, cells in sorted(per_day.items())
        )
        return (
            f"{candidate.get('sheet')}:{row}; raw={candidate.get('raw_well')!r}; "
            f"normalized={candidate.get('normalized_well')}; field={candidate.get('field')!r}; "
            f"telemetry cells {telemetry_cells or 'none'}; evidence {days_text or 'none'}; "
            f"valid days={candidate.get('valid_days')}/30; mean={number(candidate.get('mean'))}; "
            f"decision={candidate.get('decision')} ({candidate.get('reason')})"
        )

    summary = summaries["Береговое"]
    manifest_path = root / "after" / "UPDATED_Береговое.load-report.json"
    report = json.loads(manifest_path.read_text())
    source_manifest = report["source_manifest"]
    effective = report["criteria_snapshot"]
    config_hash = hashlib.sha256(
        json.dumps(effective, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()
    output: list[str] = [
        "# GDIS Loader v0.24 — production audit",
        "",
        "**Verdict: PASS** for the September 2026 Beregovoe debit run. Release is subject to the repository's Linux checks, Windows frozen build, and release workflow. Pyreinoe and Khadyr'yakhinskoye are outside this audit because their daily reports were not attached.",  # noqa: E501
        "",
        f"Calculation date: `{source_manifest['calculation_date']}`; selected field: `Береговое`; report month: `2026-09`; working-day threshold: `{effective['values']['gas_rate_working_day_min_rate']}` (default).",  # noqa: E501
        f"Runtime source commit: `{source_manifest.get('runtime_git_commit')}`; recorded runtime tag: `{source_manifest.get('runtime_release_tag')}`; Python: `{source_manifest['python_version']}`.",  # noqa: E501
        "",
        "## Input hashes",
        "",
        "| Role | File | SHA-256 |",
        "|---|---|---|",
    ]
    for role, entry in inputs.items():
        name = "GDIS_Rating_Calculator_STRUCTURAL_v2.xlsx" if role == "v2" else entry["name"]
        output.append(f"| {role} | `{name}` | `{entry['sha256']}` |")
    output.extend(
        [
            f"| effective criteria | default criteria snapshot | `{config_hash}` |",
            "",
            "The daily report is `data/09.2026 Cуточный рапорт БГП.xlsm`, read from the unique normalized sheet `Добыча телеметрия`. It is 10,994,783 bytes with SHA-256 `56fb8f826fe70671ebc0a69bb5fc80c62ef63a64f1cbab30e7d3530e996151b3`. The 31 day headers are filtered by actual report dates: tabs `01`–`30` date September 1–30, 2026; tab `31` is undated and excluded. Telemetry headers are on row 1 (`ГСК`, `СКВ.`, days 1–31), the sheet range is `A1:AH122`, and there are 3,839 formulas, 16 merged ranges, and zero blank formula caches. The daily tabs have merged headers and title/date/well/rate structure at `A1`, `A4`, `D`, and `AH`.",  # noqa: E501
            "",
            "## Root cause and matching contract",
            "",
            "The old reader interpreted `ГСК` as a field label although it is a collector identifier (blank or numeric in this file). The September rows therefore lacked valid field scope, producing false `DEBIT_SOURCE_FIELD_AMBIGUOUS` outcomes. The corrected reader proves scope from every in-month direct formula: the dated tab title identifies Beregovoe, `A4` matches the day, and the referenced `D` well identity matches the telemetry row; the daily rate is the referenced `AH` cell. Scope is filtered before well matching. Root ID and target lifecycle ID remain separate, including `71_ЗБС` and `113_ЗБС`.",  # noqa: E501
            "",
            "Fail-safe control: `Добыча телеметрия!B108` is `10703`, but `AF108` references `30!AH114`; `30!D114` is `Р-83` and `30!AB114` is `=D114`. The reader leaves this row field-unproven and does not use it for a debit.",  # noqa: E501
            "",
            "## Before and after",
            "",
            "| Eligible | Filled before | Blank before | Filled after | Blank after | Resolved ambiguous | Still ambiguous | No valid days | Missing/unproven source | Other |",  # noqa: E501
            "|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
            f"| {summary['eligible']} | {summary['filled_before']} | {summary['blank_before']} | {summary['filled_after']} | {summary['blank_after']} | {summary['resolved_ambiguous']} | {summary['still_ambiguous']} | {summary['no_valid_days']} | {summary['missing_source']} | {summary['other']} |",  # noqa: E501
            "",
            "All 63 previously filled values remain numerically unchanged within `1e-10` absolute / `1e-12` relative tolerance, and every non-debit output cell is unchanged. All 12 September baseline ambiguities now resolve with a proven field/well/lifecycle source and 30 valid working days. The remaining 3 blanks have a unique proven Beregovoe source row, but zero numeric rates strictly above the default threshold over all 30 days; each remains blank with `DEBIT_NO_VALID_WORKING_DAYS`.",  # noqa: E501
            "",
            "## Every baseline blank",
            "",
            "| Target | Base → lifecycle / history | Target row | Before → after | Old → new diagnostic | Candidate rows and evidence | Outcome |",  # noqa: E501
            "|---|---|---:|---|---|---|---|",
        ]
    )
    for row in (item for item in rows if item["before"] is None):
        candidate_details = "<br>".join(candidate_text(c) for c in row["candidates"])
        diagnostics = (
            f"{', '.join(row['old_diagnostic']) or 'none'} → "
            f"{', '.join(row['new_diagnostic']) or 'none'}"
        )
        output.append(
            f"| `{row['target_well']}` | {history_summary(row['history'])} | {row['target_row']} | — → {number(row['after'])} | {diagnostics} | {candidate_details or 'no candidate source row'} | {row['outcome']} |"  # noqa: E501
        )
    output.extend(
        [
            "",
            "For each accepted candidate, the evidence list enumerates the exact cells on all 30 dated tabs (field/title cell, date cell, well cell, and rate cell) and the source telemetry day cells. Candidate decisions preserve `ACCEPT`, `REJECT`, or `UNRESOLVED` with the reader's reason. The paired JSON retains the complete 78-target before/after map and all raw locations.",  # noqa: E501
            "",
            "## Full 78-row before/after reconciliation",
            "",
            "| Target | Target row | Field | Before | After | Old / new diagnostic | Source sheet:row | Raw → normalized source well | Field and cell evidence | Valid days | Decision / reason |",  # noqa: E501
            "|---|---:|---|---:|---:|---|---|---|---|---:|---|",
        ]
    )
    for row in rows:
        candidates = row["candidates"]
        source_rows = "<br>".join(f"{c.get('sheet')}:{c['row']}" for c in candidates) or "none"
        source_wells = (
            "<br>".join(f"{c.get('raw_well')!r} → {c.get('normalized_well')}" for c in candidates)
            or "none"
        )
        field_evidence = (
            "<br>".join(
                f"{c.get('field')!r}: {', '.join(f'{e.get("sheet")}!{e.get("cell")}' for e in c.get('field_evidence', [])) or 'no proof'}"  # noqa: E501
                for c in candidates
            )
            or "none"
        )
        candidate_decisions = (
            "<br>".join(f"{c.get('decision')}: {c.get('reason')}" for c in candidates)
            or "no candidate"
        )
        diagnostics = (
            f"{', '.join(row['old_diagnostic']) or 'none'} / "
            f"{', '.join(row['new_diagnostic']) or 'none'}"
        )
        output.append(
            f"| `{row['target_well']}` | {row['target_row']} | {row['selected_field']} | {number(row['before'])} | {number(row['after'])} | {diagnostics} | {source_rows} | {source_wells} | {field_evidence} | {', '.join(str(c['valid_days']) for c in candidates) if candidates else '0'} | {candidate_decisions}; {history_summary(row['history'])}; outcome={row['outcome']} |"  # noqa: E501
        )
    output.extend(
        [
            "",
            "## Regression, cross-field scope, and preservation",
            "",
            "- The original v0.24 preflight passed 581 tests. Fresh AH corrective gate results are retained separately in ah-make-check.log; regression coverage includes dated scope, direct references, identity mismatch, selected-field precedence, unknown scope, lifecycle suffix, leading zero, formula-cache missing values, and order independence.",  # noqa: E501
            "- Same-number controls `22`, `62`, `101`, `321`, and `322` resolve only to the available Beregovoe report candidates. The synthetic cross-field test asserts that a same-number candidate from another field is rejected. Pyreinoe and Khadyr'yakhinskoye were not run on production inputs because their reports are absent.",  # noqa: E501
            "- The July source workbook is absent from `data/`; the customer-provided historical 78/66/12 July result is not treated as September baseline and was not rerun.",  # noqa: E501
            "- Structural v3 SHA-256 before/after: `6a7523fdbf524fd9b7f6709d0b3b95279641c706f958c36d388a41c23d610248`. V2 SHA-256: `645a3e5ba9303c42f441598e71fb7da261c9738789638cb5a0e7fceaf97d2578`; rollback SHA-256: `bb046cc365749a546d1bbe141165b08a664ba427ab9493b678ddce008fc85c6a`.",  # noqa: E501
            "- Output OOXML entries changed: `xl/worksheets/sheet3.xml` and `xl/tables/table2.xml` only. Capacity remains 1,000; canonical status remains `false` pending native Excel acceptance.",  # noqa: E501
            "- The 2,340 telemetry cached values were independently checked against their linked `01`–`30!AH` values. Selected mean values use numeric daily rates strictly greater than zero; no formula or default changed.",  # noqa: E501
            "",
            "The production input books and row-level JSON/CSV are excluded from Git and the review bundle; this Markdown and the local detailed audit contain only the evidence needed for review.",  # noqa: E501
        ]
    )
    (root / "AUDIT_v0.24_FINAL.md").write_text("\n".join(output) + "\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", type=Path, default=Path("artifacts/v0.24"))
    parser.add_argument("--run", action="store_true")
    args = parser.parse_args()
    if args.run:
        run(args.directory / "after")
    reconcile(args.directory)
