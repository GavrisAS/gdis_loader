"""Local Beregovoe acceptance evidence without production row dumps."""

from __future__ import annotations

import argparse
import json
from dataclasses import replace
from datetime import date
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from gdis_loader.domain.models import SourcePaths
from gdis_loader.excel.template_capabilities import detect_template_capabilities
from gdis_loader.excel.template_migration import SOURCE_SHA256, sha256
from gdis_loader.service import run_update, run_validate


def verify(candidate: Path, data: Path, directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    paths = SourcePaths(
        calculator=candidate,
        fund=data / "Фонд 2026 на 01.08.2026.xlsx",
        gdi=data / "Сводная таблица ГДИ.xlsx",
        lha=data / "База ЛХА_Итог!.xlsm",
        daily_report=data / "09.2026 Cуточный рапорт БГП.xlsm",
        well_history=data / "База по скважинам.xlsx",
        research_status=data / "Статус исследований_2026.xlsx",
        output=None,
    )
    today = date(2026, 10, 12)
    validation = run_validate(paths, directory / "validate.log", calculation_date=today)
    assert not validation.has_errors
    evidence: dict[str, object] = {
        "template_sha256": sha256(candidate),
        "source_sha256": SOURCE_SHA256,
        "capacity": detect_template_capabilities(candidate).calculation_capacity,
        "calculation_date": today.isoformat(),
        "validation_errors": validation.error_count,
    }
    fields = {}
    examples = {"12": 21, "13": 21, "21": 19, "22": 19}
    for field, expected in (("Береговое", 78),):
        output = directory / f"UPDATED_{field}.xlsx"
        result = run_update(
            replace(paths, output=output),
            directory / f"{field}.log",
            target_field=field,
            calculation_date=today,
        )
        assert not result.has_errors
        assert result.report.rows_written == expected
        with ZipFile(candidate) as before, ZipFile(output) as after:
            assert before.namelist() == after.namelist()
            changed = [n for n in before.namelist() if before.read(n) != after.read(n)]
            assert set(changed) == {"xl/worksheets/sheet3.xml", "xl/tables/table2.xml"}
            formulas = 0
            for name in after.namelist():
                if name.endswith(".xml"):
                    xml = after.read(name)
                    ET.fromstring(xml)
                    assert b"#REF!" not in xml
                    formulas += xml.count(b"<f>")
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            sheet = workbook["Данные"]
            headers = [c.value for c in sheet[6]]
            rows = {
                str(r[headers.index("Скважина")]): r
                for r in sheet.iter_rows(min_row=7, values_only=True)
                if r[headers.index("Скважина")] is not None
            }
            assert len(rows) == expected
            for well, count in examples.items():
                assert rows[well][headers.index("Количество ГДИ")] == count
        finally:
            workbook.close()
        fields[field] = {
            "rows": expected,
            "errors": result.error_count,
            "changed_entries": changed,
            "formula_cells": formulas,
            "customer_examples": "PASS",
            "non_target_byte_identity": "PASS",
            "capacity_warning": result.has_capacity_warning,
            "stale_migration_message": any(
                "требуется отдельная structural migration" in issue.message
                for issue in result.report.issues
            ),
        }
    evidence["fields"] = fields
    (directory / "acceptance.json").write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(evidence, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("candidate", type=Path)
    parser.add_argument("data", type=Path)
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    verify(args.candidate, args.data, args.directory)
