from hashlib import sha256
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from gdis_loader.domain.models import TARGET_FIELDS, TargetRow
from gdis_loader.excel.template_capabilities import detect_template_capabilities
from gdis_loader.excel.template_v3 import EXPECTED_PARTS, V2_SHA256, migrate_v3, shift_reference
from gdis_loader.excel.writer import write_target_rows

ROOT = Path(__file__).parents[1] / "templates/calculator"
V2 = ROOT / "GDIS_Rating_Calculator_STRUCTURAL_v2.xlsx"
V3 = ROOT / "GDIS_Rating_Calculator_STRUCTURAL_v3.xlsx"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def test_v3_exact_part_gate_and_immutable_v2() -> None:
    assert sha256(V2.read_bytes()).hexdigest() == V2_SHA256
    with ZipFile(V2) as before, ZipFile(V3) as after:
        assert before.namelist() == after.namelist()
        changed = {n for n in before.namelist() if before.read(n) != after.read(n)}
        assert changed == EXPECTED_PARTS
        assert "xl/calcChain.xml" not in after.namelist()
        for name in after.namelist():
            if name.endswith(".xml"):
                data = after.read(name)
                ET.fromstring(data)
                assert b"#REF!" not in data
                assert "Расходомер" not in data.decode()
        # All formula sheets, lambda/rank, charts, styles, relationships are exact.
        for name in before.namelist():
            if name not in EXPECTED_PARTS:
                assert before.read(name) == after.read(name)
        workbook_before = ET.fromstring(before.read("xl/workbook.xml"))
        workbook_after = ET.fromstring(after.read("xl/workbook.xml"))
        for root in (workbook_before, workbook_after):
            marker = root.find(f"{NS}definedNames/{NS}definedName")
            assert marker is not None
            marker.text = "VERSION"
        assert ET.tostring(workbook_before) == ET.tostring(workbook_after)
        settings_before = before.read("xl/worksheets/sheet2.xml").decode()
        assert (
            settings_before.replace("Не используется: Расходомер; ", "Не используется: ")
            == after.read("xl/worksheets/sheet2.xml").decode()
        )


def test_v3_every_surviving_cell_moves_with_style_value_and_header() -> None:
    with ZipFile(V2) as before, ZipFile(V3) as after:
        original = ET.fromstring(before.read("xl/worksheets/sheet3.xml"))
        migrated = ET.fromstring(after.read("xl/worksheets/sheet3.xml"))
        cells = {c.attrib["r"]: c for c in migrated.iter(NS + "c")}
        expected_count = 0
        for cell in original.iter(NS + "c"):
            address = cell.attrib["r"]
            column = "".join(c for c in address if c.isalpha())
            if column_index_from_string(column) == 7:
                continue
            new_address = shift_reference(address)
            cell.set("r", new_address)
            assert ET.tostring(cell) == ET.tostring(cells[new_address])
            expected_count += 1
        assert len(cells) == expected_count
        for tag in ("conditionalFormatting", "dataValidations", "sheetProtection", "tableParts"):
            assert [ET.tostring(e) for e in original.findall(NS + tag)] == [
                ET.tostring(e) for e in migrated.findall(NS + tag)
            ]
        table = ET.fromstring(after.read("xl/tables/table2.xml"))
        assert table.attrib["ref"] == "A6:Y106"
        auto_filter = table.find(NS + "autoFilter")
        assert auto_filter is not None
        assert auto_filter.attrib["ref"] == "A6:Y106"
        columns = table.find(NS + "tableColumns")
        assert columns is not None and columns.attrib["count"] == "25"
        assert tuple(c.attrib["name"] for c in columns) == TARGET_FIELDS
        assert [int(c.attrib["id"]) for c in columns] == list(range(1, 26))
    capabilities = detect_template_capabilities(V3)
    assert capabilities.schema_version == 3
    assert capabilities.calculation_capacity == 1000
    assert capabilities.supports_missing_exclusion


def test_v3_reproducible_payload_and_reject_mutating_v2(tmp_path: Path) -> None:
    output = tmp_path / "v3.xlsx"
    proof = migrate_v3(V2, output)
    assert set(proof) == EXPECTED_PARTS
    with ZipFile(V3) as checked, ZipFile(output) as rebuilt:
        assert checked.namelist() == rebuilt.namelist()
        assert all(checked.read(n) == rebuilt.read(n) for n in checked.namelist())
    with pytest.raises(ValueError, match="new output"):
        migrate_v3(V2, V2)


def test_v3_writer_mapping_expansion_and_auto_recalc(tmp_path: Path) -> None:
    output = tmp_path / "Калькулятор результат.xlsx"
    rows = tuple(TargetRow(well=str(i), gtm="Нет", lnd="Да", debit=0) for i in range(1000))
    result = write_target_rows(V3, output, rows)
    assert set(result.changed_entries) == {"xl/worksheets/sheet3.xml", "xl/tables/table2.xml"}
    with ZipFile(V3) as before, ZipFile(output) as after:
        for name in before.namelist():
            if name not in result.changed_entries:
                assert before.read(name) == after.read(name)
        table = ET.fromstring(after.read("xl/tables/table2.xml"))
        assert table.attrib["ref"] == "A6:Y1006"
        workbook = ET.fromstring(after.read("xl/workbook.xml"))
        calc = workbook.find(NS + "calcPr")
        assert calc is not None
        assert calc.attrib["calcMode"] == "auto"
        assert calc.attrib["fullCalcOnLoad"] == calc.attrib["forceFullCalc"] == "1"
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        sheet = workbook["Данные"]
        assert tuple(c.value for c in sheet[6]) == TARGET_FIELDS
        first = dict(
            zip(
                TARGET_FIELDS,
                next(sheet.iter_rows(min_row=7, max_row=7, values_only=True)),
                strict=True,
            )
        )
        assert first["ГТМ"] == "Нет"
        assert first["ЛНД"] == "Да"
        assert first["Дебит"] == 0
        assert sheet.max_column == 25
    finally:
        workbook.close()
