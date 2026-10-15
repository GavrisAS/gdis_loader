from pathlib import Path

from openpyxl import Workbook

from gdis_loader.readers.daily import read_daily_report
from gdis_loader.readers.daily_context import day_contexts, reference_evidence


def _book() -> Workbook:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "01"
    sheet["A1"] = "БЕРЕГОВОЙ (ОПОРНЫЙ) ГАЗОВЫЙ ПРОМЫСЕЛ"
    sheet["A4"] = "01 сентября 2026 г."
    sheet["D6"] = "№"
    sheet["D7"] = "скв."
    sheet["AB6"] = "№"
    sheet["AB7"] = "скв."
    sheet["D8"] = "0022"
    sheet["AB8"] = "=D8"
    return book


def test_direct_link_proves_field_date_and_leading_zero_identity() -> None:
    contexts = day_contexts(_book())
    proof = reference_evidence("='01'!$AH8", 1, "22", contexts, "daily.xlsm")
    assert proof is not None
    assert proof[0] == "Береговое"
    assert [p.cell for p in proof[1]] == ["A1", "A4", "D8", "AH8"]
    assert reference_evidence("='01'!$AH8", 2, "22", contexts, "daily.xlsm") is None
    assert reference_evidence("='01'!$AH8", 1, "22_ЗБС", contexts, "daily.xlsm") is None


def test_section_boundaries_and_unknown_marker_clear_scope() -> None:
    book = _book()
    sheet = book["01"]
    sheet["A9"] = "ПЛУ"
    sheet["D10"] = "22"
    sheet["A11"] = "неизвестный ЛУ"
    sheet["D12"] = "22"
    contexts = day_contexts(book)
    proof = reference_evidence("='01'!AH10", 1, "22", contexts, "x")
    assert proof is not None and proof[0] == "пырейное"
    assert proof[1][0].cell == "A9"
    assert reference_evidence("='01'!AH12", 1, "22", contexts, "x") is None


def test_conflicting_secondary_identity_rejects_link() -> None:
    book = _book()
    book["01"]["AB8"] = "23"
    assert reference_evidence("='01'!AH8", 1, "22", day_contexts(book), "x") is None


def test_normalized_sheet_name_and_real_month_exclude_day_31(tmp_path: Path) -> None:
    book = _book()
    sheet = book.create_sheet("добыча телеметрия")
    sheet.append(["ГСК", "СКВ.", *range(1, 32)])
    sheet.append([99, "22", *["='01'!AH8" for _ in range(31)]])
    path = tmp_path / "сентябрь.xlsm"
    book.save(path)
    rows, issues = read_daily_report(path)
    assert len(rows[0].daily_values) == 30
    assert rows[0].report_month == "2026-09"
    assert rows[0].source_sheet == "добыча телеметрия"
    assert rows[0].raw_field_name is None  # Wrong daily references cannot prove scope.
    assert all(i.code != "DAILY_SHEET_MISSING" for i in issues)


def test_formula_cache_missing_is_not_zero(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Добыча телеметрия"
    sheet.append(["ГСК", "СКВ.", 1, 2])
    sheet.append([1, "22", "=1+2", 0])
    path = tmp_path / "cache.xlsm"
    book.save(path)
    rows, issues = read_daily_report(path)
    assert rows[0].daily_values == (None, 0.0)
    issue = next(i for i in issues if i.code == "DAILY_FORMULA_CACHE_MISSING")
    assert issue.locations[0].cell == "C2"


def test_casefold_duplicate_authoritative_sheet_is_rejected(monkeypatch) -> None:
    from gdis_loader.readers import daily

    class DuplicateSheetWorkbook:
        sheetnames = ("Добыча телеметрия", "добыча телеметрия")

        def close(self) -> None:
            pass

    monkeypatch.setattr(daily, "load_workbook", lambda *args, **kwargs: DuplicateSheetWorkbook())

    rows, issues = daily.read_daily_report(Path("duplicate.xlsm"))

    assert rows == ()
    assert any(issue.code == "DAILY_SHEET_AMBIGUOUS" for issue in issues)


def test_real_september_period_and_independent_cached_means() -> None:
    import hashlib
    import math

    import pytest
    from openpyxl import load_workbook

    path = Path(__file__).parents[1] / "data/09.2026 Cуточный рапорт БГП.xlsm"
    if not path.exists():
        pytest.skip("September production report unavailable")
    with path.open("rb") as source:
        source_hash = hashlib.file_digest(source, "sha256").hexdigest()
    if source_hash != "56fb8f826fe70671ebc0a69bb5fc80c62ef63a64f1cbab30e7d3530e996151b3":
        pytest.skip("Different source revision requires fresh production qualification")
    rows, _ = read_daily_report(path)
    assert {r.report_month for r in rows} == {"2026-09"}
    assert {len(r.daily_values) for r in rows} == {30}
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        raw = list(book["Добыча телеметрия"].iter_rows(max_col=33, values_only=True))
        for row in rows:
            values = [
                x
                for x in raw[row.source_row - 1][2:32]
                if isinstance(x, (float, int))
                and not isinstance(x, bool)
                and math.isfinite(x)
                and x > 0
            ]
            parsed = [x for x in row.daily_values if x is not None and x > 0]
            assert len(values) == len(parsed)
            if values:
                assert math.isclose(
                    sum(values) / len(values), sum(parsed) / len(parsed), rel_tol=1e-12
                )
        for well in ("12", "21", "22", "42", "62", "71", "101", "102", "111", "113", "321", "322"):
            candidate = next(r for r in rows if r.target_well_id == well)
            assert candidate.raw_field_name == "Береговое"
            assert len(candidate.field_evidence) == 120
        wrong_link = next(r for r in rows if r.target_well_id == "10703")
        assert wrong_link.field_scope_invalid
    finally:
        book.close()
