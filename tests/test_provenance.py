from __future__ import annotations

import json
import logging
import tomllib
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from gdis_loader.domain.models import (
    DiagnosticIssue,
    IssueSeverity,
    LoadReport,
    SourceLocation,
)
from gdis_loader.readers.gdi import GDI_SHEET, read_gdi
from gdis_loader.readers.lha import LHA_SHEET, LhaIndex, read_lha
from gdis_loader.service import _append_issues, configure_logging, format_diagnostic_issue
from gdis_loader.validation.sources import _check_gdi


def _write_gdi(path: Path, rows: list[list[object]], headers: list[str] | None = None) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = GDI_SHEET
    worksheet.append(
        headers
        or [
            "Месторождение",
            "Скважина",
            "Пласт",
            "Вид исследования",
            "Дата",
            "Выход на РФП (да/нет)",
            "Достоверность согласно ЛНД",
            "Дата проведения ГТМ",
        ]
    )
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def _write_lha(path: Path, rows: list[list[object]], headers: list[str] | None = None) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = LHA_SHEET
    worksheet.append(
        headers
        or [
            "Месторождение",
            "Номер скважины",
            "Пласт",
            "Дата и время исследования",
            "Данные ЛХА, CL-, мг/дм3",
            "ВГФ, г/м3",
        ]
    )
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def test_gdi_invalid_well_and_plast_have_exact_cells(tmp_path: Path) -> None:
    path = tmp_path / "Сводная таблица ГДИ.xlsx"
    _write_gdi(
        path,
        [
            ["Береговое", None, "ПК-1", "ИД", datetime(2026, 1, 1), None, None, None],
            ["Береговое", 11, None, "ИД", datetime(2026, 1, 2), None, None, None],
        ],
    )

    _records, issues = read_gdi(path)

    by_code = {issue.code: issue for issue in issues}
    assert by_code["GDI_INVALID_WELL"].locations == (
        SourceLocation(str(path.resolve()), GDI_SHEET, "B2"),
    )
    assert by_code["GDI_INVALID_PLAST"].locations == (
        SourceLocation(str(path.resolve()), GDI_SHEET, "C3"),
    )


def test_lha_invalid_date_has_exact_date_cell(tmp_path: Path) -> None:
    path = tmp_path / "База ЛХА_Итог!.xlsm"
    _write_lha(path, [["Береговое", 11, "ПК-1", "not-a-date", 10, 1]])

    _records, issues = read_lha(path)

    issue = next(issue for issue in issues if issue.code == "LHA_INVALID_DATE")
    assert issue.locations == (SourceLocation(str(path.resolve()), LHA_SHEET, "D2"),)


def test_lha_conflict_contains_all_conflicting_measurement_cells(tmp_path: Path) -> None:
    path = tmp_path / "lha.xlsm"
    latest = datetime(2026, 1, 2)
    _write_lha(
        path,
        [
            ["Береговое", 11, "ПК-1", latest, 10, None],
            ["Береговое", 11, "ПК-1", latest, 11, None],
            ["Береговое", 11, "ПК-1", datetime(2026, 1, 1), 5, None],
        ],
    )
    records, _ = read_lha(path)

    selection = LhaIndex(records).select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )
    issue = next(
        issue for issue in selection.issues if issue.code == "LHA_DUPLICATE_TIMESTAMP_CONFLICT"
    )

    assert issue.locations == (
        SourceLocation(str(path.resolve()), LHA_SHEET, "E2"),
        SourceLocation(str(path.resolve()), LHA_SHEET, "E3"),
    )


def test_missing_header_has_file_and_sheet_provenance(tmp_path: Path) -> None:
    path = tmp_path / "gdi.xlsx"
    _write_gdi(path, [], headers=["wrong"])

    _records, issues = read_gdi(path)

    issue = next(issue for issue in issues if issue.code == "GDI_HEADERS_MISSING")
    assert issue.locations == (SourceLocation(str(path.resolve()), GDI_SHEET, None),)


def test_input_validation_missing_sheet_has_file_and_expected_sheet(tmp_path: Path) -> None:
    path = tmp_path / "gdi.xlsx"
    workbook = Workbook()
    issues: list[DiagnosticIssue] = []

    _check_gdi(workbook, path, issues)

    assert issues[0].code == "GDI_SHEET_MISSING"
    assert issues[0].locations == (SourceLocation(str(path.resolve()), GDI_SHEET, None),)


def test_json_serialization_preserves_structured_locations(tmp_path: Path) -> None:
    location = SourceLocation(str(tmp_path / "source.xlsx"), "Sheet1", "B7")
    issue = DiagnosticIssue(
        code="TEST",
        severity=IssueSeverity.WARNING,
        message="problem",
        locations=(location,),
    )
    report = LoadReport("now", "0.1", "3.14", {}, issues=[issue])

    serialized = json.loads(json.dumps(report.to_dict(), ensure_ascii=False))

    assert serialized["issues"][0]["locations"] == [
        {"file": str(tmp_path / "source.xlsx"), "sheet": "Sheet1", "cell": "B7"}
    ]


def test_text_log_format_contains_file_sheet_and_cell(tmp_path: Path) -> None:
    log_path = configure_logging(tmp_path / "loader.log")
    issue = DiagnosticIssue(
        code="LHA_INVALID_DATE",
        severity=IssueSeverity.WARNING,
        message="Некорректная дата",
        locations=(SourceLocation("source.xlsm", "БД Сибнефтегаз", "G1524"),),
    )
    report = LoadReport("now", "0.1", "3.14", {})
    _append_issues(report, (issue,))
    for handler in logging.getLogger("gdis_loader").handlers:
        handler.flush()

    text = log_path.read_text(encoding="utf-8")
    assert "WARNING [LHA_INVALID_DATE]" in text
    assert "source.xlsm::БД Сибнефтегаз::G1524" in text


def test_issue_without_cell_is_serializable_and_formatting_is_safe() -> None:
    issue = DiagnosticIssue(
        code="MISSING_SHEET",
        severity=IssueSeverity.ERROR,
        message="Лист не найден",
        locations=(SourceLocation("source.xlsx", "Sheet1", None),),
    )

    assert issue.to_dict()["locations"] == [
        {"file": "source.xlsx", "sheet": "Sheet1", "cell": None}
    ]
    assert "source.xlsx::Sheet1" in format_diagnostic_issue(issue)


def test_project_metadata_uses_two_component_version_and_declared_author() -> None:
    with (Path(__file__).parents[1] / "pyproject.toml").open("rb") as stream:
        project = tomllib.load(stream)["project"]

    assert project["version"] == "0.1"
    assert project["authors"] == [{"name": "Александр Гаврись"}]
