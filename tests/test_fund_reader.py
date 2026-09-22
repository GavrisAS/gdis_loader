from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from gdis_loader.domain.models import FundCategory, IssueSeverity, SourceLocation
from gdis_loader.readers.fund import parse_well_list_cell, read_fund


def _make_fund(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    old = workbook.create_sheet("01.01.2025")
    old["C1"] = datetime(2025, 1, 1)
    old["A11"] = "Старое поле"
    old["B11"] = "Старое (ПК 1)"
    old["C11"] = 1
    old["C12"] = "999"
    ignored = workbook.create_sheet("RSPN 03.03.2025")
    ignored["C1"] = datetime(2025, 3, 3)
    ignored["A11"] = "RSPN"
    ignored["B11"] = "RSPN object"
    ignored["C11"] = 1
    ignored["C12"] = "888"
    latest = workbook.create_sheet("02.02.2025")
    latest["C1"] = datetime(2025, 2, 2)
    latest["A11"] = "Береговой"
    latest["B11"] = "Береговой (ПК 1)"
    latest["C11"] = 2
    latest["D11"] = 1
    latest["G11"] = 2
    latest["H11"] = 2
    latest["J11"] = 1
    latest["C12"] = "011, 012"
    latest["D12"] = 13
    latest["G12"] = "14\n15"
    latest["H12"] = "016, 017"
    latest["J12"] = "018"
    workbook.save(path)


def test_fund_chooses_latest_non_rspn_sheet_and_all_categories(tmp_path: Path) -> None:
    path = tmp_path / "fund.xlsx"
    _make_fund(path)

    result = read_fund(path)

    assert result.sheet_name == "02.02.2025"
    assert result.fund_date is not None
    assert len(result.wells) == 7
    assert {well.source_category for well in result.wells} == {
        FundCategory.PRODUCING,
        FundCategory.STOPPED_IN_REPORTING_MONTH,
        FundCategory.INACTIVE_CURRENT_YEAR,
        FundCategory.INACTIVE_PREVIOUS_YEARS,
    }
    assert result.wells[0].raw_display == "011"
    assert result.wells[0].base_well_id == "11"
    assert result.wells[0].object_plast_hint == "ПК1"
    assert not any(well.raw_display == "018" for well in result.wells)
    assert not any(issue.code == "FUND_COUNT_MISMATCH" for issue in result.issues)
    j_issue = next(issue for issue in result.issues if issue.code == "FUND_CATEGORY_J_IGNORED")
    assert j_issue.severity is IssueSeverity.INFO
    assert j_issue.message == (
        "Колонка J (освоение/ожидание освоения) не включена в фонд рейтингования: "
        "по исходному README рейтингованию подлежит действующий + бездействующий фонд."
    )
    assert j_issue.locations == (SourceLocation(str(path.resolve()), "02.02.2025", "J12"),)


def test_fund_count_mismatch_is_diagnostic(tmp_path: Path) -> None:
    path = tmp_path / "fund.xlsx"
    _make_fund(path)
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("01.01.2026")
    sheet["C1"] = datetime(2026, 1, 1)
    sheet["A11"] = "Поле"
    sheet["B11"] = "Поле (ПК 1)"
    sheet["C11"] = 2
    sheet["C12"] = "1"
    workbook.save(path)

    result = read_fund(path)

    assert any(issue.code == "FUND_COUNT_MISMATCH" for issue in result.issues)


def test_fund_list_parser_handles_numeric_newlines_and_leading_zeroes() -> None:
    assert parse_well_list_cell("011, 012\n013") == ["011", "012", "013"]
    assert parse_well_list_cell("153 162") == ["153", "162"]
    assert parse_well_list_cell(11) == ["11"]


def test_fund_missing_separator_is_info_with_source_cell(tmp_path: Path) -> None:
    path = tmp_path / "fund.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "01.08.2026"
    worksheet["C1"] = datetime(2026, 8, 1)
    worksheet["A11"] = "Береговое"
    worksheet["B11"] = "Береговое (ПК-1)"
    worksheet["C11"] = 2
    worksheet["C12"] = "153 162"
    workbook.save(path)

    result = read_fund(path)

    issues = [issue for issue in result.issues if issue.code == "FUND_MISSING_SEPARATOR_NORMALIZED"]
    assert len(issues) == 1
    assert issues[0].severity is IssueSeverity.INFO
    assert "два числовых идентификатора" in issues[0].message
    assert "разделено на две скважины" in issues[0].message
    assert issues[0].locations == (SourceLocation(str(path.resolve()), "01.08.2026", "C12"),)
    assert [well.raw_display for well in result.wells] == ["153", "162"]
