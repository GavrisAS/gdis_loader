from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from gdis_loader.domain.models import DatePrecision, PartialDate, WellEventKind
from gdis_loader.readers.well_history import (
    FIELD_SHEETS,
    find_well_history_schema,
    read_well_history,
)


def _write_history(path: Path, *, valid_schema: bool = True) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "БЛУ"
    sheet.append(["№", "№", "Дата", None, None, None, None, *([None] * 24)])
    second = ["куста", "скв.", "ВНС", "Пласт", "информация", "Заканчивание", "Примечание"]
    for block_number in range(1, 5):
        second.extend(
            [f"ГТМ{block_number}", "Дата", "Пласт", "Отчет", "Заканчивание", "Примечание"]
        )
    if not valid_schema:
        second[9] = "Дата события"
    sheet.append(second)
    sheet.append([1, "051", 3.2002, "ПК1", "Ссылка", None, None])
    sheet["E3"].hyperlink = "https://invalid.example/private"
    sheet["H3"] = "ЗБС"
    sheet["I3"] = 7.202
    sheet["I3"].number_format = "0.0000"
    sheet["J3"] = "ПК1"
    sheet["N3"] = "ПВЛГ"
    sheet["O3"] = datetime(2025, 9, 16)
    sheet["P3"] = "ПК2"
    sheet["T3"] = "ПВР"
    sheet["U3"] = 2015
    sheet["Z3"] = "КЛК"
    sheet["AA3"] = 13.2020
    sheet.append([1, "012", datetime(2008, 2, 1), "ПК1"])
    workbook.save(path)


def test_two_row_blu_reader_finds_dynamic_gtm_blocks_and_keeps_provenance(tmp_path: Path) -> None:
    path = tmp_path / "База по скважинам.xlsx"
    _write_history(path)

    result = read_well_history(path)

    assert FIELD_SHEETS == {
        "БЛУ": "Береговое",
        "ХЛУ": "Хадырьяхинское",
        "ПЛУ": "Пырейное",
    }
    assert len(result.histories) == 2
    first = result.histories[0]
    assert first.physical_root == "51"
    assert first.source_display_well == "051"
    assert first.events[0].event_date is not None
    assert first.events[0].event_date.precision is DatePrecision.MONTH
    assert first.events[0].source_location.cell == "C3"
    assert first.events[0].raw_block[4] == "Ссылка"
    assert all("invalid.example" not in str(value) for value in first.events[0].raw_block)
    assert [event.event_kind for event in first.events[1:]] == [
        WellEventKind.ZBS,
        WellEventKind.PVLG,
        WellEventKind.PVR,
        WellEventKind.KLK,
    ]
    assert first.events[1].event_date == parse_expected_month_date()
    assert [stage.event_kind for stage in first.stages] == [
        WellEventKind.INITIAL,
        WellEventKind.ZBS,
        WellEventKind.PVLG,
    ]
    assert result.event_counts == {"ЗБС": 1, "КЛК": 1, "ПВЛГ": 1, "ПВР": 1}
    assert result.date_precision_counts == {"DAY": 2, "MONTH": 2, "YEAR": 1, "INVALID": 1}
    invalid = next(issue for issue in result.issues if issue.code == "WELL_HISTORY_INVALID_DATE")
    assert invalid.locations[0].cell == "AA3"


def parse_expected_month_date() -> PartialDate:
    return PartialDate(2020, 7, None, DatePrecision.MONTH)


def test_schema_detects_two_row_header_and_rejects_broken_dynamic_block(tmp_path: Path) -> None:
    valid = tmp_path / "valid.xlsx"
    invalid = tmp_path / "invalid.xlsx"
    _write_history(valid)
    _write_history(invalid, valid_schema=False)

    from openpyxl import load_workbook

    valid_workbook = load_workbook(valid, read_only=True, data_only=True)
    invalid_workbook = load_workbook(invalid, read_only=True, data_only=True)
    try:
        assert find_well_history_schema(valid_workbook["БЛУ"]) is not None
        assert find_well_history_schema(invalid_workbook["БЛУ"]) is None
        from gdis_loader.domain.models import DiagnosticIssue
        from gdis_loader.validation.sources import _check_well_history

        issues: list[DiagnosticIssue] = []
        _check_well_history(invalid_workbook, invalid, issues)
        assert issues[0].code == "WELL_HISTORY_SCHEMA_INVALID"
        assert issues[0].locations[0].file == str(invalid.resolve())
        assert issues[0].locations[0].sheet == "БЛУ"
    finally:
        valid_workbook.close()
        invalid_workbook.close()
