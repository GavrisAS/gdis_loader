from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import Rule
from openpyxl.worksheet.table import Table, TableStyleInfo

from gdis_loader.domain.models import TARGET_FIELDS, TargetRow
from gdis_loader.excel.writer import TargetWorkbookError, write_target_rows


def _make_calculator(
    path: Path,
    *,
    reordered_headers: bool = False,
    outside_row: int | None = None,
    conditional_formatting: bool = False,
) -> None:
    workbook = Workbook()
    data = workbook.active
    data.title = "Данные"
    data["A1"] = "Title"
    data["A4"] = "Input"
    headers = list(TARGET_FIELDS)
    if reordered_headers:
        headers = headers[10:] + headers[:10]
    for column, header in enumerate(headers, 1):
        data.cell(6, column, header)
    for row in range(7, 9):
        for column in range(1, len(TARGET_FIELDS) + 1):
            data.cell(row, column, f"stale-{row}-{column}")
    table = Table(displayName="tblData", ref="A6:Y8")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    data.add_table(table)
    if outside_row is not None:
        data.cell(outside_row, 1, "outside-table-note")
    if conditional_formatting:
        data.conditional_formatting.add("A7:A8", Rule(type="duplicateValues", priority=1))
    other = workbook.create_sheet("Other")
    other["A1"] = "not targeted"
    other["B2"] = "=1+1"
    formula = workbook.create_sheet("Formula")
    formula["A1"] = "=Other!B2"
    workbook.save(path)


def _target_entries(path: Path) -> dict[str, bytes]:
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _table_ref(path: Path) -> str:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        return workbook["Данные"].tables["tblData"].ref
    finally:
        workbook.close()


def test_writer_uses_named_headers_clears_old_rows_and_preserves_package(tmp_path: Path) -> None:
    source = tmp_path / "calculator.xlsx"
    output = tmp_path / "result.xlsx"
    _make_calculator(source, reordered_headers=True)
    source_hash = sha256(source.read_bytes()).hexdigest()
    before = _target_entries(source)

    rows = (
        TargetRow(well="11_ЗБС", plast="ПК-1", cl_current=12.5, vgf_current=0),
        TargetRow(well="12", plast="ПК-1"),
    )
    result = write_target_rows(source, output, rows)

    assert result.rows_written == 2
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    assert _table_ref(output) == "A6:Y8"
    after = _target_entries(output)
    changed = {name for name in before if before[name] != after[name]}
    assert set(result.changed_entries) == changed
    assert all(before[name] == after[name] for name in before if name not in changed)

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        worksheet = workbook["Данные"]
        header = next(
            worksheet.iter_rows(min_row=6, max_row=6, max_col=len(TARGET_FIELDS), values_only=True)
        )
        positions = {name: index for index, name in enumerate(header)}
        first = next(
            worksheet.iter_rows(min_row=7, max_row=7, max_col=len(TARGET_FIELDS), values_only=True)
        )
        second = next(
            worksheet.iter_rows(min_row=8, max_row=8, max_col=len(TARGET_FIELDS), values_only=True)
        )
        assert first[positions["Скважина"]] == "11_ЗБС"
        assert first[positions["Cl текущий"]] == 12.5
        assert first[positions["ВГФ текущий"]] == 0
        assert second[positions["Скважина"]] == "12"
        assert second[positions["Cl текущий"]] is None
    finally:
        workbook.close()


def test_writer_expands_table_over_100_rows_only_on_target_sheet(tmp_path: Path) -> None:
    source = tmp_path / "calculator.xlsx"
    output = tmp_path / "large-result.xlsx"
    _make_calculator(source)
    before = _target_entries(source)
    rows = tuple(TargetRow(well=str(index)) for index in range(101))

    result = write_target_rows(source, output, rows)

    assert result.rows_written == 101
    assert "CALCULATOR_FORMULA_CAPACITY_UNKNOWN" in result.warnings
    assert _table_ref(output) == "A6:Y107"
    after = _target_entries(output)
    changed = set(result.changed_entries)
    assert all(before[name] == after[name] for name in before if name not in changed)
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        worksheet = workbook["Данные"]
        assert worksheet.max_row == 107
        values = list(worksheet.iter_rows(min_row=7, max_row=107, max_col=1, values_only=True))
        assert values[0][0] == "0"
        assert values[-1][0] == "100"
    finally:
        workbook.close()


def test_writer_preserves_rows_outside_old_table_range(tmp_path: Path) -> None:
    source = tmp_path / "calculator.xlsx"
    output = tmp_path / "result.xlsx"
    _make_calculator(source, outside_row=20)

    write_target_rows(source, output, (TargetRow(well="11"),))

    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        worksheet = workbook["Данные"]
        assert worksheet["A20"].value == "outside-table-note"
        assert worksheet["A7"].value == "11"
        assert all(
            worksheet.cell(8, column).value is None for column in range(1, len(TARGET_FIELDS) + 1)
        )
    finally:
        workbook.close()


def test_writer_rejects_expansion_over_non_table_row(tmp_path: Path) -> None:
    source = tmp_path / "calculator.xlsx"
    output = tmp_path / "result.xlsx"
    _make_calculator(source, outside_row=20)
    output.write_bytes(b"existing-output")

    with pytest.raises(TargetWorkbookError, match="TARGET_TABLE_EXPANSION_CONFLICT"):
        write_target_rows(source, output, tuple(TargetRow(well=str(i)) for i in range(14)))

    assert output.read_bytes() == b"existing-output"


def test_writer_expands_exact_duplicate_conditional_formatting_range(tmp_path: Path) -> None:
    source = tmp_path / "calculator.xlsx"
    output = tmp_path / "result.xlsx"
    _make_calculator(source, conditional_formatting=True)

    write_target_rows(source, output, tuple(TargetRow(well=str(i)) for i in range(5)))

    with ZipFile(output) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert 'sqref="A7:A11"' in sheet_xml


def test_writer_preserves_ignorable_namespace_declaration(tmp_path: Path) -> None:
    source = tmp_path / "calculator.xlsx"
    output = tmp_path / "result.xlsx"
    _make_calculator(source)
    namespace = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"
    with ZipFile(source, "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    sheet_name = "xl/worksheets/sheet1.xml"
    sheet_xml = entries[sheet_name].decode("utf-8")
    sheet_xml = sheet_xml.replace(
        "<worksheet ",
        (
            '<worksheet xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'mc:Ignorable="x14ac" '
            f'xmlns:x14ac="{namespace}" '
        ),
        1,
    )
    entries[sheet_name] = sheet_xml.encode("utf-8")
    with ZipFile(source, "w") as archive:
        for name, value in entries.items():
            archive.writestr(name, value)

    write_target_rows(source, output, (TargetRow(well="11"),))

    with ZipFile(output) as archive:
        result_xml = archive.read(sheet_name)
    ET.fromstring(result_xml)
    assert b'mc:Ignorable="x14ac"' in result_xml
    assert f'xmlns:x14ac="{namespace}"'.encode() in result_xml
    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.sheetnames
    finally:
        workbook.close()
