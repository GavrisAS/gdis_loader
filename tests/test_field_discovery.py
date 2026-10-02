from __future__ import annotations

from dataclasses import replace
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from gdis_loader.domain.models import TARGET_FIELDS, SourcePaths
from gdis_loader.validation.sources import validate_inputs, validate_output_path


def _make_calculator(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Данные"
    for column, header in enumerate(TARGET_FIELDS, 1):
        worksheet.cell(6, column, header)
    worksheet.add_table(Table(displayName="tblData", ref="A6:Z106"))
    workbook.save(path)


def _make_fund(path: Path, *, include_unknown: bool = False) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "01.08.2026"
    worksheet["C1"] = datetime(2026, 8, 1)
    objects = [
        ("Береговой", "Береговое (ПК 1)", "011"),
        ("Береговое", "Береговое (ПК-1)", "012"),
        ("Пырейный", "Пырейное (ПК 1)", "013"),
    ]
    if include_unknown:
        objects.append(("Новое ЛУ", "Новое ЛУ (ПК 1)", "014"))
    for index, (field, object_name, well) in enumerate(objects):
        header_row = 11 + index * 2
        worksheet.cell(header_row, 1, field)
        worksheet.cell(header_row, 2, object_name)
        worksheet.cell(header_row, 3, 1)
        worksheet.cell(header_row + 1, 3, well)
    workbook.save(path)


def _make_gdi(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ГДИС_Сибнефтегаз в целом"
    worksheet.append(
        [
            "Месторождение",
            "Скважина",
            "Пласт",
            "Вид исследования",
            "Дата",
            "Выход на РФП (да/нет)",
            "Достоверность согласно ЛНД",
            "Дата проведения ГТМ",
            "Глубина записи, м",
        ]
    )
    worksheet.append(["Береговое", 11, "ПК1", "ИД", datetime(2026, 1, 1), "да", "Высокая", None, 1])
    worksheet.append(
        ["Хадырьяхинское", 91, "ПК1", "ИД", datetime(2026, 1, 1), "да", "Высокая", None, 1]
    )
    workbook.save(path)


def _make_lha(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "БД Сибнефтегаз"
    worksheet.append(
        [
            "Месторождение",
            "Номер скважины",
            "Пласт",
            "Дата и время исследования",
            "Данные ЛХА, CL-, мг/дм3",
            "ВГФ, г/м3",
        ]
    )
    worksheet.append(["Береговое", 11, "ПК1", datetime(2026, 1, 1), 0, 0])
    workbook.save(path)


def _make_daily(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Добыча телеметрия"
    worksheet.append(["ГСК", "СКВ."])
    workbook.save(path)


def _make_history(path: Path, *, include_pyr_sheet: bool = False) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "БЛУ"
    headers = ("№ куста", "№ скв.", "Дата ВНС", "Пласт", "Информация", "Заканчивание", "Примечание")
    for column, value in enumerate(headers, 1):
        worksheet.cell(1, column, value)
    for column, value in enumerate(
        ("ГТМ1", "Дата", "Пласт", "Отчет", "Заканчивание", "Примечание"), 8
    ):
        worksheet.cell(2, column, value)
    worksheet["B3"] = "011"
    worksheet["C3"] = datetime(2020, 1, 1)
    worksheet["D3"] = "ПК1"
    if include_pyr_sheet:
        pyr = workbook.create_sheet("ПЛУ")
        for column, value in enumerate(headers, 1):
            pyr.cell(1, column, value)
        for column, value in enumerate(
            ("ГТМ1", "Дата", "Пласт", "Отчет", "Заканчивание", "Примечание"), 8
        ):
            pyr.cell(2, column, value)
    workbook.save(path)


def _make_status(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "АР_таблица_БЛУ"
    for column, value in (
        (1, "Скважина"),
        (2, "Пласт"),
        (3, "Исследования"),
        (4, "Дата исследования"),
        (10, "Дельта"),
    ):
        worksheet.cell(4, column, "Параметры")
        worksheet.cell(5, column, value)
    worksheet.cell(6, 1, 11)
    worksheet.cell(6, 2, "ПК1")
    worksheet.cell(6, 3, "ИД")
    worksheet.cell(6, 4, datetime(2026, 1, 1))
    worksheet.cell(6, 10, 0)
    workbook.save(path)


def _paths(tmp_path: Path, *, include_unknown: bool = False) -> SourcePaths:
    paths = SourcePaths(
        calculator=tmp_path / "calculator.xlsx",
        fund=tmp_path / "fund.xlsx",
        gdi=tmp_path / "gdi.xlsx",
        lha=tmp_path / "lha.xlsm",
        daily_report=tmp_path / "daily.xlsm",
        well_history=tmp_path / "history.xlsx",
        research_status=tmp_path / "status.xlsx",
        output=None,
    )
    _make_calculator(paths.calculator)
    _make_fund(paths.fund, include_unknown=include_unknown)
    _make_gdi(paths.gdi)
    _make_lha(paths.lha)
    _make_daily(paths.daily_report)
    _make_history(paths.well_history)
    _make_status(paths.research_status)
    return paths


def test_validation_discovers_canonical_fields_from_fund_and_reports_coverage(
    tmp_path: Path,
) -> None:
    result = validate_inputs(_paths(tmp_path))

    assert result.valid
    assert result.available_fields == ("Береговое", "Пырейное")
    assert result.field_coverage["Береговое"]["fund_rows"] == 2
    assert result.field_coverage["Пырейное"]["fund_rows"] == 1
    assert result.field_coverage["Береговое"]["gdi_rows"] == 1
    assert result.field_coverage["Пырейное"]["gdi_rows"] == 0
    assert "Хадырьяхинское" not in result.available_fields
    assert any(
        issue.code == "FIELD_SOURCE_SHEET_MISSING"
        and issue.source == "well_history"
        and issue.locations[0].sheet == "ПЛУ"
        for issue in result.issues
    )


def test_validation_keeps_fund_field_without_mapping_and_reports_it(tmp_path: Path) -> None:
    result = validate_inputs(_paths(tmp_path, include_unknown=True))

    assert "Новое ЛУ" in result.available_fields
    assert any(
        issue.code == "FIELD_SOURCE_SHEET_MAPPING_MISSING" and "Новое ЛУ" in issue.message
        for issue in result.issues
    )


def test_output_cannot_overwrite_any_source_workbook(tmp_path: Path) -> None:
    paths = _paths(tmp_path)
    result = validate_inputs(replace(paths, output=paths.fund))

    assert result.valid
    assert any(
        issue.code == "OUTPUT_EQUALS_INPUT" for issue in validate_output_path(paths.fund, paths)
    )
