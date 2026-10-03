from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from gdis_loader.domain.models import IssueSeverity, SourcePaths
from gdis_loader.matching.normalize import normalize_plast, normalize_well_id
from gdis_loader.readers.lha import LhaIndex, read_lha
from gdis_loader.service import run_update, run_validate

DATA_ROOT = Path(__file__).parents[1] / "data"
REAL_FILES = (
    DATA_ROOT / "GDIS_Rating_Calculator_FIXED.xlsx",
    DATA_ROOT / "Фонд 2026 на 01.08.2026.xlsx",
    DATA_ROOT / "Сводная таблица ГДИ.xlsx",
    DATA_ROOT / "База ЛХА_Итог!.xlsm",
    DATA_ROOT / "07.2026 Cуточный рапорт БГП_полный.xlsm",
    DATA_ROOT / "База по скважинам.xlsx",
    DATA_ROOT / "Статус исследований_2026.xlsx",
)


@pytest.mark.skipif(
    not all(path.is_file() for path in REAL_FILES), reason="real data files are not present"
)
def test_real_bundle_update(tmp_path: Path) -> None:
    output = tmp_path / "GDIS_Rating_Calculator_UPDATED.xlsx"
    paths = SourcePaths(
        calculator=REAL_FILES[0],
        fund=REAL_FILES[1],
        gdi=REAL_FILES[2],
        lha=REAL_FILES[3],
        daily_report=REAL_FILES[4],
        well_history=REAL_FILES[5],
        research_status=REAL_FILES[6],
        output=output,
    )

    result = run_update(
        paths,
        tmp_path / "loader.log",
        target_field="Береговое",
        calculation_date=date(2026, 10, 2),
    )

    assert result.report.calculation_date == date(2026, 10, 2)
    assert result.report.fund_total_source == 129
    assert result.report.fund_total == 129
    assert result.report.target_field == "Береговое"
    assert result.report.target_field_fund_total == 109
    assert result.report.rating_eligible == 78
    assert result.report.rating_not_eligible == 31
    assert result.report.rows_prepared == 78
    assert result.report.rows_written == result.report.rows_prepared
    assert result.report.gdi_kind_counts == {
        "EXCLUDED": 208,
        "ID": 32,
        "ID_KVD": 866,
        "KVD": 701,
        "UNKNOWN": 23,
    }
    assert result.report.gdi_raw_kind_counts == {
        "EXCLUDED": 209,
        "ID": 32,
        "ID_KVD": 867,
        "KVD": 701,
        "UNKNOWN": 23,
    }
    assert result.report.gdi_raw_record_count == 1832
    assert result.report.gdi_deduplicated_record_count == 1830
    assert result.report.gdi_exact_duplicates_collapsed == 2
    assert result.report.gdi_pre_launch_excluded == 16
    assert result.report.matched == 71
    assert result.report.ambiguous == 3
    assert result.report.no_gdi == 4
    assert result.report.debit_filled == 63
    assert result.report.debit_blank == 15
    assert result.report.debit_valid_day_count == 1505
    assert result.report.gdm_filled_yes == 5
    assert result.report.gdm_filled_no == 45
    assert result.report.gdm_blank == 28
    assert result.report.gdm_fallback_selected == 12
    assert result.report.gdm_stale_fallback_rejected == 17
    assert result.report.gdm_target_candidate_rows == 61
    expected_criteria = {
        "insufficient_gdi_min_count": 2,
        "productivity_min_increase": 0.0,
        "chlorine_current_min_mg_dm3": 2000.0,
        "chlorine_delta_min_mg_dm3": 1000.0,
        "downhole_depth_threshold_m": 0.0,
        "gdm_abs_delta_threshold": 5.0,
        "gdm_fallback_max_age_days": 365,
        "gas_rate_working_day_min_rate": 0.0,
    }
    assert result.report.criteria_snapshot == {
        "config_source": None,
        "values": expected_criteria,
        "defaults": expected_criteria,
        "value_origin": {key: "default" for key in expected_criteria},
    }
    assert result.report.well_history_records > 0
    assert result.report.well_history_gtm_events > 0
    assert result.report.gdi_kind_counts
    assert not any(issue.severity is IssueSeverity.ERROR for issue in result.report.issues)
    assert not any(
        issue.code == "LHA_DUPLICATE_TIMESTAMP_CONFLICT" for issue in result.report.issues
    )
    assert not any(issue.code == "UNCONFIRMED_WELL_VARIANT_MZS" for issue in result.report.issues)
    assert all(
        issue.severity is IssueSeverity.INFO and issue.locations
        for issue in result.report.issues
        if issue.code == "LHA_DUPLICATE_TIMESTAMP_RESOLVED_MAX"
    )
    assert any(
        issue.code == "GDI_EXACT_DUPLICATE_COLLAPSED"
        and issue.severity is IssueSeverity.INFO
        and issue.locations
        for issue in result.report.issues
    )
    assert not any(
        issue.code == "LHA_INVALID_VALUE" and issue.severity.value == "warning"
        for issue in result.report.issues
    )
    assert output.is_file()
    report_path = output.with_suffix(".load-report.json")
    assert report_path.is_file()
    serialized_report = json.loads(report_path.read_text(encoding="utf-8"))
    assert serialized_report["criteria_snapshot"] == result.report.criteria_snapshot
    workbook = load_workbook(output, read_only=True, data_only=True, keep_links=False)
    try:
        worksheet = workbook["Данные"]
        headers = tuple(cell.value for cell in next(worksheet.iter_rows(min_row=6, max_row=6)))
        well_index = headers.index("Скважина")
        count_index = headers.index("Количество ГДИ")
        launch_index = headers.index("Дата запуска")
        insufficient_index = headers.index("Недостаточно ГДИС")
        no_gdi_index = headers.index("Продолжительность без ГДИ")
        no_id_index = headers.index("Продолжительность без ИД")
        no_kvd_index = headers.index("Продолжительность без КВД")
        specific_index = headers.index("Удельное количество ГДИ")
        lnd_index = headers.index("ЛНД")
        downhole_index = headers.index("Глубинные датчики")
        productivity_index = headers.index("Снижение продуктивности")
        cl_current_index = headers.index("Cl текущий")
        cl_previous_index = headers.index("Cl предыдущий")
        gdm_index = headers.index("Проблемы ГДМ")
        debit_index = headers.index("Дебит")
        rows_by_well = {
            str(row[well_index]): row
            for row in worksheet.iter_rows(min_row=7, values_only=True)
            if row[well_index] is not None
        }
        assert len(rows_by_well) == 78
        expected_counts = {"12": 20, "13": 21, "21": 19, "22": 19}
        for well, expected_count in expected_counts.items():
            assert rows_by_well[well][count_index] == expected_count
        assert rows_by_well["12"][launch_index].date() == date(2008, 2, 1)
        assert rows_by_well["13"][launch_index].date() == date(2002, 5, 1)
        assert rows_by_well["21"][launch_index].date() == date(2004, 3, 1)
        assert rows_by_well["22"][launch_index].date() == date(2004, 3, 1)
        assert rows_by_well["10302_ЗБС"][headers.index("Пласт")] == "ПК16"
        assert rows_by_well["10601_ЗБС"][headers.index("Пласт")] == "ПК19"
        controlled = rows_by_well["10201_ЗБС3"]
        assert controlled[count_index] == 1
        assert controlled[launch_index].date() == date(2026, 1, 22)
        assert controlled[lnd_index] == "Да"
        assert controlled[downhole_index] == "Да"
        assert controlled[productivity_index] is None
        assert controlled[insufficient_index] == "Да"
        assert controlled[no_kvd_index] == (date(2026, 10, 2) - date(2026, 4, 5)).days / 365
        assert controlled[no_gdi_index] == controlled[no_kvd_index]
        assert rows_by_well["112_ЗБС"][gdm_index] == "Нет"
        assert rows_by_well["10301_ЗБС"][count_index] == 10
        assert rows_by_well["10301_ЗБС"][insufficient_index] == "Нет"
        assert rows_by_well["10301_ЗБС"][lnd_index] == "Да"
        assert rows_by_well["10301_ЗБС"][downhole_index] == "Нет"
        assert rows_by_well["10301_ЗБС"][productivity_index] == "Да"
        assert rows_by_well["10401_ЗБС"][count_index] == 1
        assert rows_by_well["10401_ЗБС"][lnd_index] == "Да"
        assert rows_by_well["10401_ЗБС"][downhole_index] == "Да"
        assert rows_by_well["10405"][headers.index("Пласт")] == "ПК19+20"
        assert rows_by_well["10405"][count_index] == 0
        assert rows_by_well["10306"][count_index] == 0
        assert any(
            issue.code == "GDI_UNCONFIRMED_LIFECYCLE_RECORDS" and issue.well_id == "10306"
            for issue in result.report.issues
        )
        zero_gdi_rows = [row for row in rows_by_well.values() if row[count_index] == 0]
        assert len(zero_gdi_rows) == 4
        for row in zero_gdi_rows:
            launch = row[launch_index].date()
            expected_duration = (date(2026, 10, 2) - launch).days / 365
            assert row[insufficient_index] == "Да"
            assert row[no_gdi_index] == expected_duration
            assert row[no_id_index] == expected_duration
            assert row[no_kvd_index] == expected_duration
            assert row[specific_index] is None
            assert row[lnd_index] is None
            assert row[downhole_index] is None
            assert row[productivity_index] is None
        assert sum(row[cl_current_index] is not None for row in zero_gdi_rows) == 2
        assert sum(row[cl_previous_index] is not None for row in zero_gdi_rows) == 1
        assert sum(row[gdm_index] is not None for row in zero_gdi_rows) == 1
        assert sum(row[debit_index] is not None for row in zero_gdi_rows) == 4
    finally:
        workbook.close()
    with ZipFile(output) as archive:
        table_xml = archive.read("xl/tables/table2.xml").decode("utf-8")
        sheet_xml = archive.read("xl/worksheets/sheet3.xml").decode("utf-8")
    table_end = 6 + result.report.rows_written
    assert f'ref="A6:Z{table_end}"' in table_xml
    assert f'sqref="A7:A{table_end}"' in sheet_xml
    assert 'mc:Ignorable="x14ac"' in sheet_xml
    assert "xmlns:x14ac=" in sheet_xml
    sheet_root = ElementTree.fromstring(sheet_xml)
    sheet_data = sheet_root.find(
        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheetData"
    )
    assert sheet_data is not None
    row_numbers = [int(row.attrib["r"]) for row in sheet_data]
    assert max(row_numbers) == table_end

    with ZipFile(REAL_FILES[0]) as source_archive, ZipFile(output) as output_archive:
        assert set(source_archive.namelist()) == set(output_archive.namelist())
        changed_entries = {
            name
            for name in source_archive.namelist()
            if source_archive.read(name) != output_archive.read(name)
        }
        non_target_entries_identical = all(
            source_archive.read(name) == output_archive.read(name)
            for name in source_archive.namelist()
            if name not in {"xl/worksheets/sheet3.xml", "xl/tables/table2.xml"}
        )
    assert changed_entries == {"xl/worksheets/sheet3.xml", "xl/tables/table2.xml"}
    assert non_target_entries_identical


@pytest.mark.skipif(
    not all(path.is_file() for path in REAL_FILES), reason="real data files are not present"
)
def test_real_validation_discovers_fields_from_fund_and_reports_coverage(tmp_path: Path) -> None:
    paths = SourcePaths(
        calculator=REAL_FILES[0],
        fund=REAL_FILES[1],
        gdi=REAL_FILES[2],
        lha=REAL_FILES[3],
        daily_report=REAL_FILES[4],
        well_history=REAL_FILES[5],
        research_status=REAL_FILES[6],
        output=None,
    )

    result = run_validate(paths, tmp_path / "validate.log", calculation_date=date(2026, 10, 2))

    assert set(result.report.available_fields) == {
        "Береговое",
        "Хадырьяхинское",
        "Пырейное",
    }
    assert set(result.report.field_coverage) == set(result.report.available_fields)
    assert all(
        isinstance(entry["fund_rows"], int) and entry["fund_rows"] > 0
        for entry in result.report.field_coverage.values()
    )
    assert result.report.output_path is None
    assert "UPDATED_" not in (result.report.input_paths.get("output") or "")


@pytest.mark.skipif(
    not all(path.is_file() for path in REAL_FILES), reason="real data files are not present"
)
def test_real_lha_duplicate_timestamps_choose_max_independent_of_source_row_order() -> None:
    records, _ = read_lha(REAL_FILES[3])
    index = LhaIndex(records)
    cases = (
        ("212", "Береговое", "ПК-1", "cl_value", "2021-10-16 00:00"),
        ("11301", "Береговое", "ПК-15", "cl_value", "2022-07-20 00:00"),
        ("021", "Пырейное", "ПК-1", "cl_value", "2024-05-17 15:00"),
        ("323", "Хадырьяхинское", "ПК-1", "cl_value", "2024-03-24 05:40"),
        ("113", "Хадырьяхинское", "ПК-20-3", "cl_value", "2024-04-18 21:00"),
        ("212", "Береговое", "ПК-1", "vgf_value", "2021-10-16 00:00"),
        ("331", "Береговое", "ПК-1", "vgf_value", "2021-12-02 21:00"),
        ("323", "Хадырьяхинское", "ПК-1", "vgf_value", "2024-03-24 05:40"),
        ("113", "Хадырьяхинское", "ПК-20-3", "vgf_value", "2024-04-18 21:00"),
    )
    for raw_well, field, plast, metric, timestamp in cases:
        base_well_id = normalize_well_id(raw_well).base_well_id
        normalized_field = field.casefold()
        normalized_plast = normalize_plast(plast)
        when = datetime.fromisoformat(timestamp)
        matching = [
            record
            for record in records
            if record.base_well_id == base_well_id
            and record.normalized_field == normalized_field
            and record.normalized_plast == normalized_plast
            and record.measured_at == when
            and getattr(record, metric) is not None
        ]
        assert matching, (raw_well, field, plast, metric, timestamp)
        target_well_id = normalize_well_id(raw_well).target_well_id
        selection = index.select(
            base_well_id=base_well_id or "",
            normalized_field=normalized_field,
            normalized_plast=normalized_plast or "",
            cutoff=datetime.max,
            target_well_id=target_well_id,
        )
        reversed_selection = LhaIndex(reversed(records)).select(
            base_well_id=base_well_id or "",
            normalized_field=normalized_field,
            normalized_plast=normalized_plast or "",
            cutoff=datetime.max,
            target_well_id=target_well_id,
        )
        assert selection.values == reversed_selection.values
        assert not any(
            issue.code == "LHA_DUPLICATE_TIMESTAMP_CONFLICT" for issue in selection.issues
        )
        resolved = [
            issue
            for issue in selection.issues
            if issue.code == "LHA_DUPLICATE_TIMESTAMP_RESOLVED_MAX" and timestamp in issue.message
        ]
        if resolved:
            assert all(
                issue.severity is IssueSeverity.INFO and issue.locations for issue in resolved
            )


@pytest.mark.skipif(
    not all(path.is_file() for path in REAL_FILES), reason="real data files are not present"
)
@pytest.mark.parametrize("field", ["Береговое", "Пырейное", "Хадырьяхинское"])
def test_real_all_fields_update_and_ooxml(tmp_path: Path, field: str) -> None:
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "real_reconcile", DATA_ROOT.parent / "scripts/reconcile_gdi.py"
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    reconcile = module.reconcile

    output = tmp_path / f"{field}.xlsx"
    paths = SourcePaths(
        calculator=REAL_FILES[0],
        fund=REAL_FILES[1],
        gdi=REAL_FILES[2],
        lha=REAL_FILES[3],
        daily_report=REAL_FILES[4],
        well_history=REAL_FILES[5],
        research_status=REAL_FILES[6],
        output=output,
    )
    result = run_update(
        paths,
        tmp_path / f"{field}.log",
        target_field=field,
        calculation_date=date(2026, 10, 3),
    )
    report = result.report
    assert field in report.available_fields
    assert report.rows_written == report.rating_eligible > 0
    assert report.rows_written == report.matched + report.no_gdi + report.ambiguous
    reconciliation = reconcile(paths, field, output, date(2026, 10, 3))
    assert reconciliation["metrics"].get("application_count_mismatches", 0) == 0
    assert not result.has_errors
    with ZipFile(paths.calculator) as original, ZipFile(output) as generated:
        assert set(original.namelist()) == set(generated.namelist())
        changed = {
            name for name in original.namelist() if original.read(name) != generated.read(name)
        }
        assert changed == {"xl/worksheets/sheet3.xml", "xl/tables/table2.xml"}
        sheet_xml = generated.read("xl/worksheets/sheet3.xml").decode()
        table_xml = generated.read("xl/tables/table2.xml").decode()
    end = 6 + report.rows_written
    assert f'ref="A6:Z{end}"' in table_xml
    assert f'sqref="A7:A{end}"' in sheet_xml
    assert 'mc:Ignorable="x14ac"' in sheet_xml
    assert "xmlns:x14ac=" in sheet_xml
    root = ElementTree.fromstring(sheet_xml)
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    sheet_data = root.find(namespace + "sheetData")
    assert sheet_data is not None
    assert max(int(row.attrib["r"]) for row in sheet_data) == end
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        sheet = workbook["Данные"]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=6, max_row=6))]
        rows = {
            str(row[headers.index("Скважина")]): dict(zip(headers, row, strict=True))
            for row in sheet.iter_rows(min_row=7, max_row=end, values_only=True)
        }
        assert len(rows) == report.rows_written
        assert (
            sum(row["Количество ГДИ"] or 0 for row in rows.values())
            == reconciliation["metrics"]["deduplicated_valid_rows_after_launch"]
        )
        # Exact customer counts are pinned only to the unchanged production bundle.
        import re

        baseline = (DATA_ROOT.parent / "docs/REAL_DATA_DELTA_v0.18.md").read_text()
        expected_hashes = re.findall(r"`([a-f0-9]{64})`", baseline)
        actual_hashes = [
            entry["sha256"]
            for entry in report.source_manifest["sources"]
            if entry["role"] != "generated_calculator"
        ]
        if set(expected_hashes) == set(actual_hashes) and field == "Пырейное":
            assert report.rows_written == report.matched == 5
            assert report.ambiguous == report.no_gdi == 0
            assert rows["22"]["Количество ГДИ"] == 11
            assert rows["62"]["Количество ГДИ"] == 4
            assert rows["22"]["Пласт"] == rows["62"]["Пласт"] == "ПК1"
        if set(expected_hashes) == set(actual_hashes) and field == "Береговое":
            assert report.rows_written == 78
            for well, count in {"12": 20, "13": 21, "21": 19, "22": 19}.items():
                assert rows[well]["Количество ГДИ"] == count
    finally:
        workbook.close()
