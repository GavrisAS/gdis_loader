from __future__ import annotations

from dataclasses import replace
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from gdis_loader.domain.models import (
    DatePrecision,
    IssueSeverity,
    LhaRecord,
    PartialDate,
    SourceLocation,
    WellEventKind,
    WellHistory,
    WellStage,
)
from gdis_loader.matching.normalize import normalize_field_name, normalize_plast, normalize_well_id
from gdis_loader.readers.common import parse_number
from gdis_loader.readers.lha import LhaIndex, read_lha


def _record(
    when: datetime,
    *,
    plast: str = "ПК₁",
    cl: float | None = None,
    vgf: float | None = None,
    row: int = 2,
    raw_cl: str | int | float | None = None,
    raw_vgf: str | int | float | None = None,
) -> LhaRecord:
    return LhaRecord(
        raw_field_name="Береговое",
        normalized_field=normalize_field_name("Береговое"),
        raw_well_name="11",
        base_well_id=normalize_well_id("11").base_well_id,
        raw_plast=plast,
        normalized_plast=normalize_plast(plast),
        measured_at=when,
        cl_value=cl,
        vgf_value=vgf,
        source_row=row,
        raw_cl_value=raw_cl,
        raw_vgf_value=raw_vgf,
    )


def test_lha_selects_two_latest_points_independently_and_preserves_zero() -> None:
    index = LhaIndex(
        [
            _record(datetime(2024, 1, 1), cl=10, vgf=0, row=2),
            _record(datetime(2025, 1, 1), cl=20, vgf=None, row=3),
            _record(datetime(2026, 1, 1), cl=30, vgf=4, row=4),
            _record(datetime(2026, 2, 1), plast="ПК-19-2", cl=999, vgf=999, row=5),
        ]
    )

    result = index.select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert result.cl_current == 30
    assert result.cl_previous == 20
    assert result.vgf_current == 4
    assert result.vgf_previous == 0


def test_lha_single_valid_point_leaves_previous_empty() -> None:
    index = LhaIndex([_record(datetime(2026, 1, 1), cl=0, vgf=None)])

    result = index.select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert result.cl_current == 0
    assert result.cl_previous is None


def test_lha_errors_and_non_numeric_values_are_missing() -> None:
    assert parse_number("#DIV/0!") is None
    assert parse_number("#REF!") is None
    assert parse_number("") is None
    assert parse_number("-") is None
    assert parse_number(0) == 0


def test_lha_duplicate_timestamp_equal_values_is_one_point_with_info() -> None:
    latest = datetime(2026, 1, 2)
    index = LhaIndex(
        [
            _record(latest, cl=10, row=2),
            _record(latest, cl=10.0, row=3),
            _record(datetime(2026, 1, 1), cl=5, row=4),
        ]
    )

    result = index.select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert result.cl_current == 10
    assert result.cl_previous == 5
    assert [issue.code for issue in result.issues] == ["LHA_DUPLICATE_TIMESTAMP_RESOLVED_MAX"]
    assert result.issues[0].severity is IssueSeverity.INFO


def test_lha_two_distinct_values_choose_max() -> None:
    latest = datetime(2026, 1, 2)
    index = LhaIndex(
        [
            _record(latest, cl=10, row=2),
            _record(latest, cl=11, row=3),
            _record(datetime(2026, 1, 1), cl=5, row=4),
        ]
    )

    result = index.select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert result.cl_current == 11
    assert result.cl_previous == 5
    assert [issue.code for issue in result.issues] == ["LHA_DUPLICATE_TIMESTAMP_RESOLVED_MAX"]
    assert result.issues[0].severity is IssueSeverity.INFO
    assert "максимальное валидное значение" in result.issues[0].message


def test_lha_previous_duplicate_timestamp_uses_max() -> None:
    index = LhaIndex(
        [
            _record(datetime(2026, 1, 3), cl=20, row=2),
            _record(datetime(2026, 1, 2), cl=10, row=3),
            _record(datetime(2026, 1, 2), cl=11, row=4),
            _record(datetime(2026, 1, 1), cl=5, row=5),
        ]
    )

    result = index.select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert result.cl_current == 20
    assert result.cl_previous == 11
    assert len(result.issues) == 1
    assert "максимальное валидное значение" in result.issues[0].message


def test_lha_older_duplicate_timestamp_is_also_resolved_and_reported() -> None:
    index = LhaIndex(
        [
            _record(datetime(2026, 1, 3), cl=30, row=2),
            _record(datetime(2026, 1, 2), cl=20, row=3),
            _record(datetime(2026, 1, 1), cl=10, row=4),
            _record(datetime(2026, 1, 1), cl=11, row=5),
        ]
    )

    result = index.select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert result.cl_current == 30
    assert result.cl_previous == 20
    assert len(result.issues) == 1
    assert "2026-01-01 00:00:00" in result.issues[0].message


def test_lha_current_and_previous_duplicates_each_choose_max() -> None:
    index = LhaIndex(
        [
            _record(datetime(2026, 1, 3), cl=30, row=2),
            _record(datetime(2026, 1, 3), cl=31, row=3),
            _record(datetime(2026, 1, 2), cl=20, row=4),
            _record(datetime(2026, 1, 2), cl=21, row=5),
            _record(datetime(2026, 1, 1), cl=10, row=6),
            _record(datetime(2026, 1, 1), cl=11, row=7),
        ]
    )

    result = index.select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert result.cl_current == 31
    assert result.cl_previous == 21
    assert len(result.issues) == 3


def test_lha_duplicate_resolution_is_independent_of_row_order() -> None:
    records = [
        _record(datetime(2026, 1, 2), cl=10, row=2),
        _record(datetime(2026, 1, 2), cl=15, row=3),
        _record(datetime(2026, 1, 1), cl=5, row=4),
    ]
    first = LhaIndex(records).select(
        base_well_id="11", normalized_field="береговое", normalized_plast="ПК1"
    )
    reversed_result = LhaIndex(reversed(records)).select(
        base_well_id="11", normalized_field="береговое", normalized_plast="ПК1"
    )

    assert first.values == reversed_result.values
    assert first.cl_current == 15


def test_lha_numeric_value_wins_over_missing_or_invalid_duplicates() -> None:
    latest = datetime(2026, 1, 2)
    numeric = _record(latest, cl=10, row=2)
    missing = _record(latest, row=3, raw_cl="")
    invalid = _record(latest, row=4, raw_cl="bad value")

    result = LhaIndex([numeric, missing, invalid]).select(
        base_well_id="11", normalized_field="береговое", normalized_plast="ПК1"
    )

    assert result.cl_current == 10
    assert len(result.issues) == 1
    assert result.issues[0].severity is IssueSeverity.INFO
    assert "из 3 строк" in result.issues[0].message


def test_lha_all_invalid_duplicate_timestamp_has_no_selected_value_or_resolution_info() -> None:
    latest = datetime(2026, 1, 2)
    first = _record(latest, row=2, raw_cl="bad")
    second = _record(latest, row=3, raw_cl="-")

    result = LhaIndex([first, second]).select(
        base_well_id="11", normalized_field="береговое", normalized_plast="ПК1"
    )

    assert result.cl_current is None
    assert result.cl_previous is None
    assert result.issues == ()


def test_lha_reader_aggregates_expected_missing_values_and_notes(tmp_path: Path) -> None:
    path = tmp_path / "lha.xlsm"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "БД Сибнефтегаз"
    worksheet.append(
        [
            "Месторождение",
            "Номер скважины",
            "Пласт",
            "Дата и время исследования",
            "Данные ЛХА, CL-, мг/дм3",
            "ВГФ, г/м3",
        ]
    )
    worksheet.append(["Береговое", 11, "ПК-1", datetime(2026, 1, 1), "#DIV/0!", "-"])
    worksheet.append(["Береговое", 11, "ПК-1", datetime(2026, 1, 2), "Проба потеряна", ""])
    workbook.save(path)

    records, issues = read_lha(path)

    assert records[0].cl_value is None
    assert records[0].vgf_value is None
    assert all(issue.severity is not IssueSeverity.WARNING for issue in issues)
    assert any(issue.code == "LHA_MISSING_VALUE_SUMMARY" for issue in issues)
    assert any(
        issue.code == "LHA_INVALID_VALUE" and "Проба потеряна" in issue.message for issue in issues
    )


def test_lha_reader_resolves_numeric_duplicate_and_keeps_invalid_value_diagnostic(
    tmp_path: Path,
) -> None:
    path = tmp_path / "lha.xlsm"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "БД Сибнефтегаз"
    worksheet.append(
        [
            "Месторождение",
            "Номер скважины",
            "Пласт",
            "Дата и время исследования",
            "Данные ЛХА, CL-, мг/дм3",
            "ВГФ, г/м3",
        ]
    )
    timestamp = datetime(2026, 1, 2)
    worksheet.append(["Береговое", 11, "ПК1", timestamp, 10, None])
    worksheet.append(["Береговое", 11, "ПК1", timestamp, "ошибка", None])
    workbook.save(path)

    records, read_issues = read_lha(path)
    selected = LhaIndex(records).select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert selected.cl_current == 10
    assert any(issue.code == "LHA_INVALID_VALUE" for issue in read_issues)
    resolved = next(
        issue for issue in selected.issues if issue.code == "LHA_DUPLICATE_TIMESTAMP_RESOLVED_MAX"
    )
    assert resolved.locations[0].cell == "E2"
    assert resolved.locations[1].cell == "E3"


def test_lha_ignores_row_with_only_non_identity_value(tmp_path: Path) -> None:
    path = tmp_path / "lha.xlsm"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "БД Сибнефтегаз"
    worksheet.append(
        [
            "Месторождение",
            "Номер скважины",
            "Пласт",
            "Дата и время исследования",
            "Данные ЛХА, CL-, мг/дм3",
            "ВГФ, г/м3",
        ]
    )
    worksheet.append([None, None, None, None, 42, None])
    workbook.save(path)

    records, issues = read_lha(path)

    assert records == ()
    assert not any(issue.code.startswith("LHA_INVALID_") for issue in issues)


def test_lha_partially_filled_identity_row_keeps_diagnostics(tmp_path: Path) -> None:
    path = tmp_path / "lha.xlsm"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "БД Сибнефтегаз"
    worksheet.append(
        [
            "Месторождение",
            "Номер скважины",
            "Пласт",
            "Дата и время исследования",
            "Данные ЛХА, CL-, мг/дм3",
            "ВГФ, г/м3",
        ]
    )
    worksheet.append(["Береговое", 11, None, None, 42, None])
    workbook.save(path)

    records, issues = read_lha(path)
    codes = {issue.code for issue in issues}

    assert len(records) == 1
    assert {"LHA_INVALID_DATE", "LHA_INVALID_PLAST"} <= codes


def test_lha_metrics_collapse_duplicates_independently() -> None:
    latest = datetime(2026, 1, 3)
    index = LhaIndex(
        [
            _record(latest, cl=10, vgf=100, row=2),
            _record(latest, cl=11, vgf=100, row=3),
            _record(datetime(2026, 1, 2), cl=5, vgf=50, row=4),
            _record(datetime(2026, 1, 1), vgf=40, row=5),
        ]
    )

    result = index.select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert result.cl_current == 11
    assert result.cl_previous == 5
    assert result.vgf_current == 100
    assert result.vgf_previous == 50
    assert len(result.issues) == 2


def test_suffixless_lha_rows_attach_to_zbs_stage_only_by_time_history() -> None:
    initial = WellStage(
        field_name="Береговое",
        physical_root="11",
        source_display_well="011",
        target_well_id="11",
        plast="ПК1",
        event_kind=WellEventKind.INITIAL,
        started_at=PartialDate(2010, 1, 1, DatePrecision.DAY),
        stage_sequence=0,
        source_location=SourceLocation("history.xlsx", "БЛУ", "C3"),
        raw_block=(),
    )
    zbs = replace(
        initial,
        event_kind=WellEventKind.ZBS,
        started_at=PartialDate(2025, 1, 10, DatePrecision.DAY),
        stage_sequence=1,
        source_location=SourceLocation("history.xlsx", "БЛУ", "I8"),
    )
    history = WellHistory("Береговое", "11", "011", (), (initial, zbs))
    before = replace(
        _record(datetime(2024, 10, 1), cl=4, row=2),
        raw_well_name="11",
        source_file="lha.xlsm",
        source_sheet="БД Сибнефтегаз",
        date_cell="D2",
        cl_cell="E2",
    )
    after = replace(
        _record(datetime(2025, 2, 1), cl=12, row=3),
        raw_well_name="11",
        source_file="lha.xlsm",
        source_sheet="БД Сибнефтегаз",
        date_cell="D3",
        cl_cell="E3",
    )
    target_id = normalize_well_id("11_ЗБС").target_well_id

    result = LhaIndex((before, after)).select(
        base_well_id="11",
        physical_root="11",
        target_well_id=target_id,
        normalized_field="береговое",
        normalized_plast="ПК1",
        histories=(history,),
        lifecycle_stage=zbs,
    )

    assert result.cl_current == 12
    assert result.cl_previous is None
    assert not result.issues


def test_lha_three_values_at_one_timestamp_choose_max() -> None:
    latest = datetime(2026, 1, 3)
    index = LhaIndex(
        [
            _record(latest, cl=1, row=2),
            _record(latest, cl=2, row=3),
            _record(latest, cl=3, row=4),
        ]
    )

    result = index.select(
        base_well_id="11",
        normalized_field="береговое",
        normalized_plast="ПК1",
    )

    assert result.cl_current == 3
    assert result.cl_previous is None
    assert "максимальное валидное значение" in result.issues[0].message
