from __future__ import annotations

from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from gdis_loader.readers.research_status import read_research_status
from gdis_loader.transforms.research_status import gdm_problem


def _make_status(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "АР_таблица_БЛУ"
    worksheet["A4"] = "Скважина"
    worksheet["B4"] = "Пласт"
    worksheet["C4"] = "Исследования"
    worksheet["D4"] = "Дата исследования"
    worksheet["J4"] = "Дельта"
    worksheet["A5"] = "Объект"
    worksheet["J5"] = "мг/дм3"
    worksheet.append([])
    rows: tuple[tuple[object, ...], ...] = (
        (11, "ПК1", "ИД", datetime(2026, 1, 1), 5),
        (12, "ПК-1", "КВД", datetime(2026, 1, 2), 5.01),
        (13, "ПК1", "Рпл", datetime(2026, 1, 3), -5.1),
        (14, "ПК1", "ИД", datetime(2026, 1, 4), "-"),
        (15, "ПК1", "ИД", datetime(2026, 1, 5), None),
        (16, "ПК1", "ИД", datetime(2026, 1, 6), "#VALUE!"),
        (17, "ПК1", "ИД", datetime(2026, 1, 7), "=2+5"),
    )
    for row_number, (well, plast, kind, when, delta) in enumerate(rows, 6):
        worksheet.cell(row_number, 1, well)
        worksheet.cell(row_number, 2, plast)
        worksheet.cell(row_number, 3, kind)
        worksheet.cell(row_number, 4, when)
        worksheet.cell(row_number, 10, delta)
    workbook.save(path)


def _set_formula_cache(path: Path, cell_reference: str, cached_value: str) -> None:
    with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as stream:
        temporary_path = Path(stream.name)
    try:
        with ZipFile(path) as source, ZipFile(temporary_path, "w", ZIP_DEFLATED) as output:
            for entry in source.infolist():
                data = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    cell = next(
                        item
                        for item in root.iter()
                        if item.tag.endswith("}c") and item.attrib.get("r") == cell_reference
                    )
                    value = next(item for item in cell if item.tag.endswith("}v"))
                    value.text = cached_value
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                output.writestr(entry, data)
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def test_reader_maps_multirow_headers_preserves_delta_states_and_provenance(
    tmp_path: Path,
) -> None:
    path = tmp_path / "Статус исследований_2026.xlsx"
    _make_status(path)
    _set_formula_cache(path, "J12", "7")

    records, issues = read_research_status(path, "БЛУ")

    assert len(records) == 7
    assert [item.delta_value for item in records] == [5, 5.01, -5.1, None, None, None, 7]
    assert [
        gdm_problem(item.delta_value, 5.0) if item.delta_value is not None else None
        for item in records
    ] == [
        False,
        True,
        True,
        None,
        None,
        None,
        True,
    ]
    assert records[0].source_sheet == "АР_таблица_БЛУ"
    assert records[0].well_cell == "A6"
    assert records[0].plast_cell == "B6"
    assert records[0].research_type_cell == "C6"
    assert records[0].date_cell == "D6"
    assert records[0].delta_cell == "J6"
    assert records[-1].delta_cell == "J12"
    assert records[-1].delta_formula_without_cache is False
    assert any(issue.code == "GDM_DELTA_INVALID" for issue in issues)
    assert any(issue.code == "GDM_DELTA_MISSING" for issue in issues)


def test_reader_retains_formula_without_cached_value_as_unknown(tmp_path: Path) -> None:
    path = tmp_path / "status.xlsx"
    _make_status(path)

    records, issues = read_research_status(path, "БЛУ")

    formula_record = records[-1]
    assert formula_record.delta_value is None
    assert formula_record.delta_formula_without_cache
    assert formula_record.delta_value is None
    assert any(issue.code == "GDM_DELTA_FORMULA_CACHE_MISSING" for issue in issues)
